"""Read-only workbook profiler for Release Guard integration tests."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

# 分类id (产品分类) and 产品属性 (属性 JSON) are now probed so the guard can
# validate category to attribute.templatePid consistency. Without these the
# 2026-08-06 L071/L072/L081/L088 cid corruption was invisible to every gate.
HEADERS = ("产品标题", "产品货号", "变种属性值一", "SKU货号", "预览图", "轮播图", "产品素材图", "SKC属性", "分类id", "产品属性")


def _parse_template_pids(value: Any) -> list[str]:
    """Return the templatePid list from a 产品属性 JSON cell, or marker on error."""
    if value is None or str(value).strip() == "":
        return []
    try:
        obj = json.loads(str(value))
        if isinstance(obj, list):
            out = []
            for x in obj:
                if isinstance(x, dict) and x.get("templatePid") is not None:
                    out.append(str(x["templatePid"]))
            return out
    except (ValueError, TypeError):
        return ["__PARSE_ERROR__"]
    return []


def _dimension(ws) -> str:
    try:
        return ws.calculate_dimension(force=True)
    except Exception:
        return ""


def _rows_of(probe: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for sheet in probe["sheets"]:
        rows.extend(sheet["physical_rows"])
    return rows


def probe_xlsx(path: str | Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl is required for workbook probing") from e
    p = Path(path).resolve()
    if not p.is_file():
        raise FileNotFoundError(p)
    digest = hashlib.sha256(p.read_bytes()).hexdigest()
    wb = load_workbook(p, read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        headers = [str(x).strip() if x is not None else "" for x in (values[0] if values else ())]
        index = {h: i for i, h in enumerate(headers)}
        missing = [h for h in HEADERS if h not in index]
        physical_rows = []
        for row_number, row in enumerate(values[1:], start=2):
            d = row[index["产品货号"]] if "产品货号" in index and index["产品货号"] < len(row) else None
            if d is None or str(d).strip() == "":
                continue
            get = lambda h: row[index[h]] if h in index and index[h] < len(row) else None
            cid_val = str(get("分类id") or "")
            pattr_val = get("产品属性")
            physical_rows.append({"row": row_number, "D": str(d), "G": str(get("变种属性值一") or ""),
                                  "SKU": str(get("SKU货号") or ""), "J": str(get("预览图") or ""),
                                  "T": str(get("轮播图") or ""), "U": str(get("产品素材图") or ""),
                                  "cid": cid_val, "templatePids": _parse_template_pids(pattr_val)})
        sheets.append({"sheet": ws.title, "dimension": _dimension(ws), "header_count": len(headers),
                       "headers": headers, "missing_required_headers": missing,
                       "physical_row_count": len(physical_rows), "physical_rows": physical_rows})
    return {"schema": "temu-release-workbook-probe/v1", "path": str(p), "sha256": digest,
            "sheet_count": len(sheets), "sheets": sheets}


def audit_category_attributes(path: str | Path, reference_path: str | Path | None = None, category_registry_path: str | Path | None = None) -> dict[str, Any]:
    """Audit 分类id to 产品属性.templatePid consistency for every physical row.

    Checks: non-empty 分类id, valid 产品属性 JSON, no templatePid shared across
    different cids (cross-contamination), and per-L0xx cid match against a
    certified reference workbook when one is supplied. Returns a dict consumed
    by ReleaseGuard.verify_category_attributes.
    """
    probe = probe_xlsx(path)
    rows = _rows_of(probe)
    ref_lxx_cid: dict[str, str] = {}
    if reference_path:
        for r in _rows_of(probe_xlsx(reference_path)):
            m = re.match(r"(L\d{3})", r["D"])
            if m:
                ref_lxx_cid[m.group(1)] = r["cid"]
    if category_registry_path:
        _reg = json.loads(Path(category_registry_path).read_text(encoding="utf-8"))
        for _lxx, _cid in _reg.get("L0xx_cid", {}).items():
            ref_lxx_cid.setdefault(_lxx, str(_cid))
    empty_cid, empty_attributes, parse_errors, tp_mismatches, ref_mismatches = [], [], [], [], []
    tp_to_cid: dict[str, str] = {}
    for r in rows:
        if not r["cid"]:
            empty_cid.append({"row": r["row"], "D": r["D"]})
        if not r["templatePids"]:
            empty_attributes.append({"row": r["row"], "D": r["D"]})
        if "__PARSE_ERROR__" in r["templatePids"]:
            parse_errors.append({"row": r["row"], "D": r["D"]})
        for tp in r["templatePids"]:
            if tp and tp != "__PARSE_ERROR__":
                if tp in tp_to_cid and tp_to_cid[tp] != r["cid"]:
                    tp_mismatches.append({"row": r["row"], "D": r["D"], "templatePid": tp,
                                          "cid": r["cid"], "other_cid": tp_to_cid[tp]})
                else:
                    tp_to_cid.setdefault(tp, r["cid"])
        m = re.match(r"(L\d{3})", r["D"])
        if m and m.group(1) in ref_lxx_cid and ref_lxx_cid[m.group(1)] != r["cid"]:
            ref_mismatches.append({"row": r["row"], "D": r["D"], "L0xx": m.group(1),
                                   "cid": r["cid"], "reference_cid": ref_lxx_cid[m.group(1)]})
    failures = empty_cid + empty_attributes + parse_errors + tp_mismatches + ref_mismatches
    return {"schema": "temu-category-attribute-audit/v1", "workbook_path": probe["path"],
            "workbook_sha256": probe["sha256"], "reference_path": str(reference_path) if reference_path else None,
            "row_count": len(rows), "empty_cid_rows": empty_cid, "empty_attribute_rows": empty_attributes,
            "parse_error_rows": parse_errors,
            "cid_templatepid_mismatches": tp_mismatches, "cid_vs_reference_mismatches": ref_mismatches,
            "failure_count": len(failures)}


def write_probe(path: str | Path, output: str | Path) -> dict[str, Any]:
    data = probe_xlsx(path)
    out = Path(output); out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data
