"""Build the second full 1999-D Agnes product-only T1 generation plan."""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(r"D:\Desktop\jit\HJXYmall\YeahF_1999D_0808新作_执行资料_20260810")
ROOT = BASE / "t1_cangyuan_current_workbook_t1"
WORKBOOK = Path(r"D:\Desktop\jit\HJXYmall\YeahF_1999D_0808新作_店小秘上传预备版_申报价格按核价2.5倍修正_20260813.xlsx")
J_MANIFEST = BASE / "t1_original_j_reference_manifest.jsonl"
PRIOR_SELECTION = ROOT / "finalization_20260914" / "YeahF_1999D_final_unbadged_T1_selection_manifest.json"
OUT = ROOT / "agnes_set2_full1999_20260915"
PLAN = OUT / "YeahF_1999D_Agnes第二套1999计划_20260915.json"


CAMERAS = (
    "low three-quarter front-left camera",
    "eye-level three-quarter front-right camera",
    "slightly high oblique left camera",
    "slightly high oblique right camera",
    "near-product wide-angle side camera",
    "deeper telephoto view through a foreground layer",
    "low side camera with strong floor or counter depth",
    "centered eye-level camera with asymmetric surroundings",
)
PALETTES = (
    "cool white daylight with pale stone and soft gray accents",
    "blue-gray premium daylight with restrained contrast",
    "warm oak and cream natural window light without yellow cast",
    "dark walnut and charcoal with clean directional daylight",
    "fresh green and light stone natural daylight",
    "black-white modern contrast with soft overcast light",
    "soft ivory and muted sage with bright side light",
    "neutral linen and travertine with realistic morning light",
)
CASTS = (
    "a young adult woman in her mid twenties",
    "a young adult man in his late twenties",
    "an adult woman in her thirties",
    "an adult man in his thirties",
    "a middle-aged woman in her forties",
    "a middle-aged man in his forties",
    "an adult woman in her fifties",
    "an adult man in his fifties",
    "no person, but an immediately use-ready arrangement that clearly explains the function",
    "two adults, one in their thirties and one in their forties",
)
POSITIONS = ("lower-left", "lower-right", "center-left", "center-right", "central foreground", "off-center midground")
FOREGROUNDS = (
    "a softly blurred linen edge",
    "a softly blurred basket edge",
    "a softly blurred plant leaf",
    "a softly blurred plain ceramic object",
    "a softly blurred work cloth",
    "a softly blurred furniture edge",
)
BACKGROUNDS = (
    "a window and doorway forming deep layers",
    "open shelving and a distant room forming depth",
    "a long counter and soft architectural depth",
    "glass doors and a restrained outdoor view",
    "a premium cabinet wall and side window",
    "a clean corridor and layered furniture behind",
)


FAMILY = {
    "L042": {
        "name": "flexible PE lawn and flower-bed edging with anchoring spikes",
        "lock": "Preserve the exact edging strip profile, colour, flexibility, spike shape, spike count/specification and installed boundary logic.",
        "settings": ("landscaped front lawn", "curved flower-bed border", "modern courtyard garden", "garden path edge", "suburban backyard lawn", "villa entrance planting bed"),
        "tasks": ("installing spikes along a smooth curve", "aligning the border beside a flower bed", "pressing one anchor into soil", "checking a completed lawn boundary", "laying out the strip before installation"),
    },
    "L043": {
        "name": "white garment folding-board set",
        "lock": "Preserve the exact perforated folding-board outline, hinge/panel geometry, holes, colour and the variant pack quantity; every board must remain identical.",
        "settings": ("bright utility-room folding counter", "walk-in wardrobe island", "hotel laundry worktable", "bedroom dressing bench", "dry-cleaning preparation table", "sunlit linen room"),
        "tasks": ("folding a shirt with one board while the remaining boards are countable", "sorting the full board set into neat stacks", "demonstrating the folding sequence", "placing folded garments beside the complete board set"),
    },
    "L047": {
        "name": "freestanding metal garden arch trellis",
        "lock": "Preserve the exact arch silhouette, adjustable frame, rods, connectors, feet, colour and freestanding ground contact.",
        "settings": ("garden pathway entrance", "courtyard gate", "backyard lawn aisle", "villa flower garden", "patio garden threshold", "outdoor ceremony garden without text"),
        "tasks": ("training an unbranded vine along the frame", "securing a plant tie to the arch", "checking the arch alignment", "arranging plain greenery around the base"),
    },
    "L048": {
        "name": "artificial green privacy foliage fence panel",
        "lock": "Preserve the exact panel/grid structure, leaf density, green colour, edge shape, fastening points and panel count.",
        "settings": ("apartment balcony railing", "garden privacy wall", "courtyard fence", "patio divider", "terrace side screen", "villa boundary fence"),
        "tasks": ("fastening the panel to a railing", "aligning two panel edges", "checking privacy coverage", "smoothing the foliage after installation"),
    },
    "L058": {
        "name": "spin-mop bucket system with mop handle and heads",
        "lock": "Preserve the exact bucket body, colour blocks, spin/wring mechanism, handle, wheels/feet if present, mop pole and included head count.",
        "settings": ("premium mudroom", "sunlit tiled kitchen", "large entry hall", "bright utility room", "covered patio floor", "modern laundry room"),
        "tasks": ("wringing the mop in the exact bucket", "mopping beside the bucket", "replacing the mop head", "moving the complete system between floor zones"),
    },
    "L063": {
        "name": "foldable multi-function core and Pilates training board",
        "lock": "Preserve the exact board outline, colour, handles, tracks, pads, elastic components, fold joints and accessory count.",
        "settings": ("bright home fitness studio", "calm bedroom exercise corner", "sunroom workout area", "modern apartment gym nook", "wood-floor living room", "quiet wellness studio"),
        "tasks": ("performing a controlled core exercise", "setting the handles for a workout", "demonstrating a safe kneeling movement", "folding the board after exercise"),
    },
    "L068": {
        "name": "large covered multi-level dish drying and storage rack",
        "lock": "Preserve the exact cover, rack silhouette, shelves, dish slots, side units, rods, feet, drainage parts, colour and width-height-depth proportions.",
        "settings": ("cool marble kitchen counter", "warm walnut kitchen", "bright island side counter", "dark stone kitchen", "white-tile sink-side counter", "large pantry preparation counter"),
        "tasks": ("placing a plate into the rack", "arranging cups and bowls", "removing a dry plate", "checking the drainage tray"),
    },
    "L071": {
        "name": "manual adjustable mobile bedside or sofa table",
        "lock": "Preserve the exact tabletop, height-adjustment column, base, wheels/casters, controls, colour and stable floor contact.",
        "settings": ("bedroom bedside workspace", "sofa reading corner", "bright study nook", "guest-room writing area", "sunroom chair side", "compact apartment living room"),
        "tasks": ("adjusting the table height", "writing in a notebook at the table", "rolling the table beside a chair", "placing a book and cup on the stable top"),
    },
    "L072": {
        "name": "covered multi-layer shoe storage cabinet or rack",
        "lock": "Preserve the exact layer/drawer count, front openings, handles, rods, connector rings, top handles, feet, cover colour and proportions.",
        "settings": ("entryway shoe-storage wall", "walk-in wardrobe", "bedroom dressing corner", "mudroom cabinet bay", "apartment hallway", "guest-room closet"),
        "tasks": ("placing shoes into one layer", "opening the correct front section", "organising pairs by shelf", "closing the dust cover after storage"),
    },
    "L074": {
        "name": "large three-tier kitchen dish drying rack",
        "lock": "Preserve exactly three tiers, frame, plate slots, cup/accessory zones, drip trays, drainage parts, feet and colour.",
        "settings": ("bright sink-side counter", "large island cleanup zone", "warm wood kitchen", "cool gray kitchen", "white-tile kitchen", "dark premium kitchen"),
        "tasks": ("placing dishes on the correct tiers", "removing a dry bowl", "emptying the drip tray", "sorting cups and plates"),
    },
    "L075": {
        "name": "countertop dish drying rack with drainboard",
        "lock": "Preserve the exact tier count, frame, plate slots, utensil holders, drainboard/drain spout, feet, colour and proportions.",
        "settings": ("marble sink-side counter", "oak kitchen island", "compact apartment kitchen", "farmhouse preparation counter", "blue-gray kitchen", "black-white kitchen"),
        "tasks": ("loading washed plates", "arranging cutlery", "checking the drainage outlet", "removing dry tableware"),
    },
    "L076": {
        "name": "rectangular long-pile pet bed mat",
        "lock": "Preserve the exact rectangular outline, pile texture, thickness, edge construction, colour and variant dimensions.",
        "settings": ("living-room pet corner", "bedroom window side", "home-office pet area", "sunroom floor", "quiet hallway nook", "family-room sofa side"),
        "tasks": ("smoothing the mat while a dog settles", "guiding a pet onto the mat", "placing the mat beside a sofa", "showing a relaxed pet resting without covering the outline"),
    },
    "L077": {
        "name": "compact wooden corrugated cat house and scratcher",
        "lock": "Preserve the exact compact body proportions, curved rear panel, front opening, round hole, slots, fasteners and corrugated side scratching panels; keep realistic cat scale.",
        "settings": ("living-room cat corner", "home-office window nook", "bedroom pet area", "reading-room floor", "sunroom cat zone", "quiet hallway pet space"),
        "tasks": ("playing with a cat through the opening", "placing a toy inside without covering the product", "encouraging a cat to enter", "sitting beside the cat house while the cat uses it"),
    },
    "L081": {
        "name": "countertop organiser with wood top and pull-out wire basket",
        "lock": "Preserve the exact wood top, black frame, short legs, side rails, one pull-out wire basket, colour and proportions; all feet must stay on one counter with no extra stand.",
        "settings": ("apartment coffee counter", "hotel-suite beverage station", "co-working lounge counter", "bright breakfast island", "home-office refreshment bar", "sunroom beverage console"),
        "tasks": ("placing coffee packets into the basket", "pulling the basket out on its original rails", "organising cups under the top", "returning the basket into the frame"),
    },
    "L082": {
        "name": "expandable or pull-out cabinet dish organiser",
        "lock": "Preserve the exact body, left-right telescoping mechanism, support surfaces, slots, front-facing horizontal axis, colour and proportions; add no drawer rails or duplicate tray.",
        "settings": ("lower kitchen cabinet", "premium pantry base cabinet", "large island cabinet", "utility cabinet", "modern scullery cabinet", "bright dish-storage cupboard"),
        "tasks": ("extending the organiser horizontally", "placing dishes on the supported tray", "adjusting the width inside the cabinet", "pulling the exact tray forward without lifting it"),
    },
    "L083": {
        "name": "freestanding two-level storage organiser with upper wire basket and lower caddy",
        "lock": "Preserve the exact full bottom base, lower solid caddy, handle-slot side panels, rear rods, upper wire basket, connectors and proportions; the base must remain fully supported during the task.",
        "settings": ("home-office supply table", "premium utility counter", "pantry organisation counter", "wardrobe linen shelf", "bathroom vanity counter", "gold-toned storage console"),
        "tasks": ("sorting supplies between upper basket and lower caddy", "placing bottles into the lower tray", "arranging small items in the upper basket", "loading folded linen while keeping the base planted"),
    },
    "L085": {
        "name": "three-piece stainless wall-repair scraper tool set",
        "lock": "Preserve exactly three scraper sizes, stainless blades, handle shapes, straight edges and set proportions.",
        "settings": ("wall-repair workbench", "paint-preparation table", "interior maintenance area", "drop-cloth renovation corner", "utility tool bench", "plaster patching station"),
        "tasks": ("selecting one scraper while all three remain visible", "spreading safe plain wall filler", "cleaning the three tools after use", "arranging the complete set by size"),
    },
    "L086": {
        "name": "black kitchen appliance rack with two drawer or basket units",
        "lock": "Preserve the black-only frame, exact two drawer/basket units, front faces, top board, vertical supports, side frame, short legs and proportions; no white product variant.",
        "settings": ("kitchen appliance counter", "pantry equipment shelf", "coffee preparation sideboard", "large island appliance zone", "dark walnut kitchen", "bright utility kitchen"),
        "tasks": ("organising tools in the two drawers", "placing containers on the shelves", "opening one original drawer", "arranging supplies around an unbranded appliance with its screen off"),
    },
    "L087": {
        "name": "two-tier pull-out metal under-sink rack",
        "lock": "Preserve the exact two-tier frame, basket/shelf structure, rods, supports, pull-out direction, colour and width-height proportions.",
        "settings": ("sink-base cabinet", "vanity lower cabinet", "utility cabinet", "pantry lower bay", "laundry cabinet", "island cleaning-supply cabinet"),
        "tasks": ("pulling out the lower tier", "placing cleaning bottles on the correct shelf", "organising supplies around plumbing without obstruction", "returning the basket into the cabinet"),
    },
    "L088": {
        "name": "offset stepped multi-basket fruit and vegetable rack",
        "lock": "Preserve the offset stepped basket layout, long bottom basket, upper baskets, central and side support rods, feet, connectors, colour and outline; never straighten it into a tower.",
        "settings": ("bright kitchen wall", "pantry produce corner", "breakfast-room sideboard", "market-style home pantry", "sunroom food-storage area", "large island end"),
        "tasks": ("sorting fruit across the stepped baskets", "placing vegetables into the long bottom basket", "restocking bread and produce", "removing one fruit while every basket stays visible"),
    },
    "L089": {
        "name": "two-tier metal countertop storage shelf",
        "lock": "Preserve exactly two tiers, frame, shelf surfaces, supports, short feet, colour and proportions.",
        "settings": ("spice preparation counter", "coffee and tea station", "pantry shelf", "dining-room sideboard", "small apartment kitchen", "sunlit breakfast counter"),
        "tasks": ("arranging spice jars on both tiers", "organising cups and tea supplies", "placing a container on the upper shelf", "removing one item from the lower shelf"),
    },
    "L091": {
        "name": "white multi-drawer storage organiser with transparent fronts",
        "lock": "Preserve the exact drawer count, white frame, transparent fronts, black handles, top groove/recess grid pattern, upper edge and front-facing proportions.",
        "settings": ("warm walnut walk-in closet", "cool white wardrobe shelf", "dark premium closet", "entryway cabinet", "home-office storage cubby", "bedroom dressing shelf"),
        "tasks": ("opening one exact drawer", "placing folded items into the correct drawer", "organising accessories while the top remains visible", "closing the transparent-front drawer"),
    },
    "L092": {
        "name": "three-piece stainless cutting-board set",
        "lock": "Preserve exactly three boards, their three sizes, stainless surfaces, edge geometry, hole count and hole positions; all three must remain visible and countable.",
        "settings": ("cool gray marble kitchen", "warm wood kitchen island", "dark stone preparation counter", "white-tile sink-side counter", "bright chef-style island", "sunlit pantry worktop"),
        "tasks": ("using the smallest board while all three remain visible", "arranging the three boards by size", "preparing vegetables on one board beside the other two", "drying the complete three-board set"),
    },
    "L095": {
        "name": "two- or three-section raised garden planter",
        "lock": "Preserve the exact variant section count, planter body, partitions, legs/support, colour, depth and stable ground or deck contact.",
        "settings": ("rooftop terrace garden", "apartment balcony herb garden", "sunny courtyard", "greenhouse bench area", "backyard planting corner", "modern patio garden"),
        "tasks": ("planting herbs in separate sections", "adding soil to one compartment", "watering seedlings without hiding the structure", "arranging labels with no readable text while every section stays visible"),
    },
    "L096": {
        "name": "folding portable outdoor barbecue grill",
        "lock": "Preserve the exact grill body, grate, folding frame, legs, hinges, locks, vents, panels, colour and heat-safe ground contact; no ordinary table placement and no fire.",
        "settings": ("wild campsite ground", "backyard patio pavers", "lawn picnic area", "RV campsite gravel", "courtyard deck floor", "park picnic preparation zone"),
        "tasks": ("preparing unlit skewers beside the grill", "unfolding the stable legs", "placing vegetables on the cold grate", "arranging safe grill tools beside the unlit product"),
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def compact(value: object) -> str:
    return " ".join(str(value or "").split())


def quantity_lock(l0xx: str, g_value: str) -> str:
    if l0xx == "L043":
        found = re.search(r"(15|20|30)\s*(?:片|个|件)", g_value)
        return f" Exactly {found.group(1)} identical folding boards must be visibly represented and countable." if found else " Preserve the exact pack quantity shown by the reference and variant."
    if l0xx == "L092":
        return " Exactly three cutting boards must be visible and countable."
    if l0xx == "L095":
        found = re.search(r"([23])\s*(?:联|格|层)", g_value)
        return f" Exactly {found.group(1)} connected planter sections must be visible and countable." if found else " Preserve the exact two- or three-section variant shown in the reference."
    return ""


def main() -> None:
    prior = json.loads(PRIOR_SELECTION.read_text(encoding="utf-8-sig"))
    prior_by_d = {row["D"]: row for row in prior["records"]}
    if len(prior_by_d) != 1999:
        raise RuntimeError("Prior final T1 selection is not 1999/1999")

    j_rows = []
    for line in J_MANIFEST.read_text(encoding="utf-8-sig").splitlines():
        if line.strip():
            j_rows.append(json.loads(line))
    j_by_d = {str(row["target_D"]): row for row in j_rows}
    if len(j_by_d) != 1999 or any(row.get("status") != "validated_original_direct_sku_reference" for row in j_rows):
        raise RuntimeError("Exact-D product-material manifest is incomplete")

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    ws = workbook.active
    headers = [compact(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    ix = {name: headers.index(name) for name in ("产品货号", "产品标题", "变种属性值一", "SKU货号")}
    workbook_rows = {}
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        d_value = compact(row[ix["产品货号"]])
        if d_value and d_value not in workbook_rows:
            workbook_rows[d_value] = {
                "row": row_num,
                "title": compact(row[ix["产品标题"]]),
                "G": compact(row[ix["变种属性值一"]]),
                "SKU": compact(row[ix["SKU货号"]]),
            }
    workbook.close()
    if len(workbook_rows) != 1999 or set(workbook_rows) != set(j_by_d):
        raise RuntimeError("Workbook exact-D coverage does not match product materials")

    family_index = defaultdict(int)
    recipes_seen = defaultdict(set)
    records = []
    for d_value in sorted(j_by_d):
        l0xx = d_value[:4]
        spec = FAMILY.get(l0xx)
        if not spec:
            raise RuntimeError(f"No family scene/product lock for {l0xx}")
        i = family_index[l0xx]
        family_index[l0xx] += 1
        settings, tasks = spec["settings"], spec["tasks"]
        setting = settings[i % len(settings)]
        task = tasks[(i // len(settings)) % len(tasks)]
        camera = CAMERAS[(i // (len(settings) * len(tasks))) % len(CAMERAS)]
        palette = PALETTES[(i * 3 + i // 7) % len(PALETTES)]
        cast = CASTS[(i * 7 + i // 9) % len(CASTS)]
        position = POSITIONS[(i * 5 + i // 11) % len(POSITIONS)]
        foreground = FOREGROUNDS[(i * 3 + i // 13) % len(FOREGROUNDS)]
        background = BACKGROUNDS[(i * 5 + i // 17) % len(BACKGROUNDS)]
        scene_recipe = (
            f"setting={setting}; task={task}; cast={cast}; camera={camera}; palette={palette}; "
            f"product_position={position}; foreground={foreground}; background={background}"
        )
        if scene_recipe in recipes_seen[l0xx]:
            scene_recipe += f"; unique_spatial_variant={i + 1}"
        recipes_seen[l0xx].add(scene_recipe)

        j = j_by_d[d_value]
        product = Path(j["local_path"])
        if not product.is_file() or sha256(product).lower() != str(j["sha256"]).lower():
            raise RuntimeError(f"Product material missing/hash drift: {d_value}")
        previous = Path(prior_by_d[d_value]["local_path"])
        if not previous.is_file() or sha256(previous) != prior_by_d[d_value]["sha256"]:
            raise RuntimeError(f"Prior final candidate missing/hash drift: {d_value}")
        w = workbook_rows[d_value]
        q_lock = quantity_lock(l0xx, w["G"])
        person_clause = (
            "Show the stated person performing the product-related task with believable hands and posture."
            if not cast.startswith("no person")
            else "Show no person; preserve an unmistakable use-ready state without inventing product parts."
        )
        prompt = (
            "MODE: product-only high-differentiation lifestyle reconstruction for a second independent T1 set. "
            "The single input image is the absolute and only visual authority for the sale product. "
            f"Product family: {spec['name']}. Variant specification (do not render as text): {w['G']}. "
            f"PRODUCT LOCK: {spec['lock']}{q_lock} "
            f"Create this exact recipe: {scene_recipe}. {person_clause} "
            "Make a genuinely new photograph that differs from the prior T1 set in setting, layout, person identity/age/clothing, "
            "camera lane, product position/scale, foreground/background, palette and lighting. Preserve exact product body, colour, "
            "count, hardware, proportions, support/contact, gravity and real function. The entire product must remain inspectable and "
            "physically supported with contact shadows. Exactly one sale-product set unless the quantity lock says otherwise. "
            "No readable text, measurements, arrows, badges, inset circles, logos, watermarks, duplicate/fused products, extra legs, "
            "extra rails, changed count, warped hardware, floating, impossible use or furniture-edge overhang. Square 1:1 polished "
            "realistic premium ecommerce photograph."
        )
        records.append({
            "D": d_value,
            "L0xx": l0xx,
            "source_row": int(j["source_row"]),
            "G": w["G"],
            "SKU": w["SKU"],
            "submitted_images": [{
                "order": 1,
                "role": "absolute_exact_D_product_identity",
                "path": str(product),
                "sha256": sha256(product),
            }],
            "prior_final_candidate_not_submitted": {
                "path": str(previous),
                "sha256": sha256(previous),
            },
            "old_T1_submitted": False,
            "prior_generated_T1_submitted": False,
            "scene_recipe": scene_recipe,
            "product_lock": spec["lock"] + q_lock,
            "prompt": prompt,
            "status": "planned_user_authorized_20260914",
            "writeback": "separate_alternative_set_not_for_current_workbook",
        })

    if len(records) != 1999 or len({row["D"] for row in records}) != 1999:
        raise RuntimeError("Plan coverage/uniqueness failure")
    duplicate_recipes = {key: len(value) for key, value in recipes_seen.items() if len(value) != family_index[key]}
    if duplicate_recipes:
        raise RuntimeError(f"Scene recipe uniqueness failure: {duplicate_recipes}")
    OUT.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": "yeahf-1999d-agnes-second-full-set-product-only-plan/v1",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "authorization": "user explicitly authorized another full 1999 Agnes set on 2026-09-14",
        "workbook": str(WORKBOOK),
        "workbook_sha256": sha256(WORKBOOK),
        "record_count": len(records),
        "by_L0xx": dict(sorted(Counter(row["L0xx"] for row in records).items())),
        "all_exact_D_product_references_hash_verified": True,
        "all_old_and_prior_generated_T1_not_submitted": True,
        "scene_recipe_unique_within_L0xx": True,
        "cast_policy": "20s/30s/40s/50s/no-person/couple rotation; no elderly-dominant lane",
        "adaptive_concurrency_policy": "ramp 1..8; first error/timeout stops ramp; cooldown; use last fully successful level",
        "badge_applied": False,
        "oss_upload": False,
        "workbook_writeback": False,
        "records": records,
    }
    PLAN.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(PLAN), "records": len(records), "by_L0xx": payload["by_L0xx"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
