#!/usr/bin/env python3
"""Validate the L042/L043 J correction writeback workbook."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ALLOWED_CHANGED_HEADERS = {"预览图", "SKC属性"}
EXPECTED_PREFIX = "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/dxxmall-0616-2/192-set1-j-corrections-l042-l043/"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def save_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def headers(ws) -> list[str]:
    return [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]


def require_col(header: list[str], name: str) -> int:
    if name not in header:
        raise RuntimeError(f"missing required column: {name}")
    return header.index(name) + 1


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--combined-json", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    args = parser.parse_args()

    combined = load_json(args.combined_json)["records"]
    corrected = {
        int(r["row"]): r
        for r in combined
        if r.get("status") in {"corrected_l042_final_cutout", "corrected_l043_quantity"}
    }

    src_wb = load_workbook(args.source, read_only=False, data_only=False)
    out_wb = load_workbook(args.output, read_only=False, data_only=False)
    src_ws = src_wb.active
    out_ws = out_wb.active
    src_header = headers(src_ws)
    out_header = headers(out_ws)
    if src_header != out_header:
        raise RuntimeError("header mismatch")

    c_d = require_col(out_header, "产品货号")
    c_g = require_col(out_header, "变种属性值一")
    c_j = require_col(out_header, "预览图")
    c_sku = require_col(out_header, "SKU货号")
    c_t = require_col(out_header, "轮播图")
    c_u = require_col(out_header, "产品素材图")
    c_skc = require_col(out_header, "SKC属性")

    changed_cells: list[dict[str, Any]] = []
    protected_changes: list[dict[str, Any]] = []
    corrected_rows_ok: list[dict[str, Any]] = []
    corrected_rows_bad: list[dict[str, Any]] = []
    empty_j: list[dict[str, Any]] = []
    row_identity_mismatch: list[dict[str, Any]] = []
    t_changed: list[dict[str, Any]] = []
    u_changed: list[dict[str, Any]] = []

    max_row = max(src_ws.max_row, out_ws.max_row)
    max_col = max(src_ws.max_column, out_ws.max_column)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            old = src_ws.cell(row, col).value
            new = out_ws.cell(row, col).value
            if old != new:
                column_name = out_header[col - 1] if col <= len(out_header) else f"col{col}"
                item = {"row": row, "column": column_name, "old": old, "new": new}
                changed_cells.append(item)
                if column_name not in ALLOWED_CHANGED_HEADERS:
                    protected_changes.append(item)

    for row in range(2, out_ws.max_row + 1):
        d_value = str(out_ws.cell(row, c_d).value or "").strip()
        if not d_value:
            continue
        j_value = str(out_ws.cell(row, c_j).value or "").strip()
        if not j_value:
            empty_j.append({"row": row, "D": d_value})
        if src_ws.cell(row, c_t).value != out_ws.cell(row, c_t).value:
            t_changed.append({"row": row, "D": d_value})
        if src_ws.cell(row, c_u).value != out_ws.cell(row, c_u).value:
            u_changed.append({"row": row, "D": d_value})
        for col, name in ((c_d, "产品货号"), (c_g, "变种属性值一"), (c_sku, "SKU货号")):
            if src_ws.cell(row, col).value != out_ws.cell(row, col).value:
                row_identity_mismatch.append({"row": row, "column": name})
        if row in corrected:
            record = corrected[row]
            expected_d = str(record.get("D") or "").strip()
            expected_g = str(record.get("G") or "").strip()
            expected_sku = str(record.get("SKU") or "").strip()
            row_errors = []
            if d_value != expected_d:
                row_errors.append("D_mismatch")
            if str(out_ws.cell(row, c_g).value or "").strip() != expected_g:
                row_errors.append("G_mismatch")
            if str(out_ws.cell(row, c_sku).value or "").strip() != expected_sku:
                row_errors.append("SKU_mismatch")
            if not j_value.startswith(EXPECTED_PREFIX):
                row_errors.append("J_url_prefix_not_corrected")
            skc_text = str(out_ws.cell(row, c_skc).value or "")
            if j_value and j_value not in skc_text:
                row_errors.append("SKC_not_synced_to_J")
            target = corrected_rows_bad if row_errors else corrected_rows_ok
            target.append({"row": row, "D": d_value, "status": record.get("status"), "J": j_value, "errors": row_errors})

    source_rows = src_ws.max_row
    output_rows = out_ws.max_row
    source_cols = src_ws.max_column
    output_cols = out_ws.max_column

    changed_counter = Counter(item["column"] for item in changed_cells)
    corrected_status_counts = Counter(r.get("status") for r in corrected.values())
    validation = {
        "source": str(args.source),
        "output": str(args.output),
        "combined_review": str(args.combined_json),
        "source_rows": source_rows,
        "output_rows": output_rows,
        "source_cols": source_cols,
        "output_cols": output_cols,
        "corrected_expected_count": len(corrected),
        "corrected_ok_count": len(corrected_rows_ok),
        "corrected_bad_count": len(corrected_rows_bad),
        "corrected_status_counts": dict(corrected_status_counts),
        "changed_cell_count": len(changed_cells),
        "changed_columns": dict(changed_counter),
        "empty_j_count": len(empty_j),
        "protected_cell_changed_count": len(protected_changes),
        "row_identity_mismatch_count": len(row_identity_mismatch),
        "t_changed_count": len(t_changed),
        "u_changed_count": len(u_changed),
        "pass_hard_checks": not any([corrected_rows_bad, empty_j, protected_changes, row_identity_mismatch, t_changed, u_changed]),
    }
    report = {
        "validation": validation,
        "corrected_rows_ok": corrected_rows_ok,
        "corrected_rows_bad": corrected_rows_bad,
        "changed_cells": changed_cells,
        "issues": {
            "empty_j": empty_j[:20],
            "protected_changes": protected_changes[:20],
            "row_identity_mismatch": row_identity_mismatch[:20],
            "t_changed": t_changed[:20],
            "u_changed": u_changed[:20],
        },
    }
    args.out_dir.mkdir(parents=True, exist_ok=True)
    save_json(args.out_dir / "j_corrections_validation_report.json", validation)
    save_json(args.out_dir / "j_corrections_writeback_report.json", report)
    save_json(args.out_dir / "j_corrections_cell_diff_summary.json", changed_cells)
    src_wb.close()
    out_wb.close()
    print(json.dumps(validation, ensure_ascii=False, indent=2))
    return 0 if validation["pass_hard_checks"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
