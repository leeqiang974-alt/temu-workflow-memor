"""Zero-trust global audit for the active YeahF 398-D draft.

This is deliberately read-only.  It distinguishes spreadsheet consistency,
J-asset lineage, and visual approval.  A synchronized JSON preview is not
accepted as evidence that the preview depicts the claimed row variant.
"""

from __future__ import annotations

import csv
import html
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Administrator\Documents\temu自动化")
OUT = ROOT / "outputs" / "yeahf_title_dedup_20260709" / "yeahf_400d_refill_20260718"
WORKBOOK = OUT / "YeahF_398D_剔除T1不可用L085两组_待J_T1回填.xlsx"
J_ASSETS = Path(r"D:\temu素材库\T首图已付费库\YeahF_400D_20260716_rebuild\YeahF_400D_resolved_assets_oss_urls_20260717.json")
LOCKS = OUT / "j_historical_rejection_locks_20260718.json"
REPORT = OUT / "YeahF_398D_零信任全局审计_20260718.json"
HTML = OUT / "YeahF_398D_零信任全局审查_20260718.html"


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def file_url(path: str | Path) -> str:
    return "file:///" + quote(str(path).replace("\\", "/"), safe="/:._-")


def esc(value: object) -> str:
    return html.escape(text(value), quote=True)


def parse_list(value: object, name: str, row: int, issues: list[dict]) -> list[dict]:
    raw = text(value)
    if not raw:
        return []
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        issues.append({"row": row, "layer": "data", "field": name, "kind": "invalid-json", "detail": str(exc)})
        return []
    if not isinstance(parsed, list):
        issues.append({"row": row, "layer": "data", "field": name, "kind": "json-not-list", "detail": type(parsed).__name__})
        return []
    return [item for item in parsed if isinstance(item, dict)]


def title_fingerprint(title: str) -> str:
    match = re.search(r"(?:^|\s)([A-Z0-9]{4})\s*$", title)
    return match.group(1) if match else ""


def count_claim(value: str) -> str:
    compact = value.replace(" ", "")
    match = re.search(r"(?:[2-9](?:格|联|个|件|套)|[一二三四五六七八九十]+(?:格|联|个|件|套)|(?:15|30)(?:pcs|PCS|个|件|套))", compact)
    return match.group(0) if match else ""


def render(groups: dict[str, list[dict]], report: dict) -> None:
    cards: list[str] = []
    for d_value, rows in groups.items():
        level = "blocked" if any(row["hard"] for row in rows) else "review"
        row_html = []
        for row in rows:
            source = row["source_path"]
            source_view = (
                f"<button class='zoom' data-src='{esc(file_url(source))}' data-label='{esc(d_value)} · 行 {row['row']} 源图'><img loading='lazy' src='{esc(file_url(source))}'></button>"
                if source and Path(source).is_file() else "<div class='none'>无可验证源图</div>"
            )
            issue_text = "<br>".join(esc(item) for item in row["notes"]) or "数据/资产链路无结构异常；仍待视觉确认。"
            row_html.append(
                f"<article class='row {('hard' if row['hard'] else 'pending')}' data-search='{esc((d_value+' '+row['G']+' '+row['SKU']+' '+row['status']+' '+ ' '.join(row['notes'])).lower())}'>"
                f"<header><b>行 {row['row']} · {esc(row['G'])}</b><span>{esc(row['SKU'])} · {esc(row['status'])}</span></header>"
                f"<div class='pics'><figure>{source_view}<figcaption>J 源图：{esc(source)}</figcaption></figure>"
                f"<figure><button class='zoom' data-src='{esc(row['J'])}' data-label='{esc(d_value)} · 行 {row['row']} 当前 J'><img loading='lazy' src='{esc(row['J'])}'></button><figcaption>当前 J（OSS）</figcaption></figure></div>"
                f"<p>{issue_text}</p><div class='decision'><select><option value='pending'>待视觉审</option><option value='approved'>视觉通过</option><option value='redo'>重做</option><option value='discard'>丢弃</option></select><textarea placeholder='审核意见'></textarea></div></article>"
            )
        cards.append(
            f"<section class='group {level}' data-search='{esc((d_value+' '+ ' '.join(r['G']+' '+r['SKU']+' '+ ' '.join(r['notes']) for r in rows)).lower())}'>"
            f"<h2>{esc(d_value)} <small>{esc(rows[0]['title'][-24:])}</small></h2><div class='rows'>{''.join(row_html)}</div></section>"
        )
    counts = report["summary"]
    doc = f"""<!doctype html><html lang='zh-CN'><meta charset='utf-8'><title>YeahF 398D 零信任全局审查</title>
<style>*{{box-sizing:border-box}}body{{margin:0;background:#f2f5f7;font:14px Arial,'Microsoft YaHei',sans-serif;color:#17202c}}.top{{position:sticky;top:0;z-index:9;padding:14px 20px;background:#fff;border-bottom:2px solid #c9d5df}}h1{{font-size:20px;margin:0 0 6px}}.top p{{margin:4px 0;color:#506071}}input,button,select,textarea{{font:inherit;border:1px solid #b9c7d3;border-radius:4px;padding:7px;background:#fff}}input{{width:min(680px,95vw)}}main{{padding:16px;max-width:2200px;margin:auto}}.notice{{padding:12px;background:#fff2db;border:1px solid #e5bd74;border-radius:5px;line-height:1.5}}.summary{{display:flex;gap:9px;flex-wrap:wrap;margin:12px 0}}.badge{{padding:6px 9px;background:#fff;border:1px solid #cbd6df;border-radius:4px}}.group{{background:#fff;border:1px solid #ced9e2;border-radius:6px;margin:14px 0;padding:11px}}.group.blocked{{border:2px solid #d45e5e;background:#fff8f8}}h2{{margin:0 0 9px;font-size:17px;color:#263f59}}h2 small{{font-weight:normal;color:#697a89;margin-left:8px}}.rows{{display:grid;grid-template-columns:repeat(auto-fit,minmax(390px,1fr));gap:9px}}.row{{border:1px solid #d7e0e7;border-radius:5px;padding:8px;background:#fff}}.row.hard{{border:2px solid #d76565;background:#fff7f7}}.row header{{display:flex;justify-content:space-between;gap:7px;color:#2c4d66}}.pics{{display:grid;grid-template-columns:1fr 1fr;gap:7px;margin-top:7px}}figure{{margin:0}}.zoom{{padding:0;border:0;width:100%;cursor:zoom-in;background:#edf1f4}}.zoom img{{width:100%;display:block;aspect-ratio:1;object-fit:contain;background:#edf1f4}}figcaption{{font-size:10px;color:#647483;overflow-wrap:anywhere;padding-top:3px}}.none{{height:155px;display:grid;place-items:center;background:#f3f5f6;color:#a43f3f}}.row p{{font-size:12px;line-height:1.45;color:#4c5c69}}.decision{{display:grid;grid-template-columns:100px 1fr;gap:6px}}textarea{{width:100%;height:40px}}.hidden{{display:none!important}}#modal{{display:none;position:fixed;z-index:20;inset:0;background:#000c;place-items:center;padding:20px}}#modal.open{{display:grid}}#modal img{{max-width:94vw;max-height:88vh}}#modal button{{position:absolute;right:20px;top:20px}}</style>
<body><div class='top'><h1>YeahF 398D · 零信任全局审查</h1><p>702 物理行 / 398 个 D。绿色不存在：所有未人工确认的行均为“待视觉审”，不会自动变成通过。</p><input id='q' placeholder='搜索 D / G / SKU / 异常'><button id='export'>导出审核意见</button></div><main><div class='notice'><b>三层结论分离：</b>数据联动与 J 资产血缘只证明字段一致/来源可追溯；不证明视觉正确。红框是硬阻断。当前 L095 `-01` 的 3格错误会整组显示红框。</div><div class='summary'><span class='badge'>数据行：{counts['rows']}</span><span class='badge'>精确 D：{counts['d_groups']}</span><span class='badge'>数据硬错：{counts['data_hard']}</span><span class='badge'>J 血缘硬错：{counts['j_hard']}</span><span class='badge'>T/U/T4 硬错：{counts['carousel_hard']}</span><span class='badge'>待视觉审：{counts['visual_pending']}</span></div>{''.join(cards)}</main><div id='modal'><button>关闭</button><img></div><script>const s=JSON.parse(localStorage.getItem('yeahf-zero-trust-20260718')||'{{}}');document.querySelectorAll('.row').forEach(x=>{{let k=x.querySelector('header').innerText,sel=x.querySelector('select'),ta=x.querySelector('textarea'),v=s[k]||{{d:'pending',n:''}};sel.value=v.d;ta.value=v.n;let f=()=>{{s[k]={{d:sel.value,n:ta.value,at:new Date().toISOString()}};localStorage.setItem('yeahf-zero-trust-20260718',JSON.stringify(s))}};sel.onchange=f;ta.oninput=f}});q.oninput=()=>{{let v=q.value.toLowerCase();document.querySelectorAll('[data-search]').forEach(x=>x.classList.toggle('hidden',!x.dataset.search.includes(v)))}};let m=document.querySelector('#modal');document.querySelectorAll('.zoom').forEach(x=>x.onclick=()=>{{m.querySelector('img').src=x.dataset.src;m.classList.add('open')}});m.onclick=e=>{{if(e.target===m)m.classList.remove('open')}});m.querySelector('button').onclick=()=>m.classList.remove('open');export.onclick=()=>{{let a=document.createElement('a'),b=new Blob([JSON.stringify(s,null,2)],{{type:'application/json'}});a.href=URL.createObjectURL(b);a.download='yeahf_398d_zero_trust_review.json';a.click();URL.revokeObjectURL(a.href)}};</script></body></html>"""
    HTML.write_text(doc, encoding="utf-8")


def main() -> None:
    assets = json.loads(J_ASSETS.read_text(encoding="utf-8"))["j"]
    asset_index = {(item["D"], item["G"], item["SKU"]): item for item in assets}
    locks = {item["D"] for item in json.loads(LOCKS.read_text(encoding="utf-8"))["locks"]}
    ws = load_workbook(WORKBOOK, read_only=True, data_only=False).active
    headers = {text(cell.value): index + 1 for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
    groups: dict[str, list[dict]] = defaultdict(list)
    issues: list[dict] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        def v(name: str) -> str:
            return text(values[headers[name] - 1])
        d_value, g_value, sku = v("产品货号"), v("变种属性值一"), v("SKU货号")
        if not d_value:
            continue
        notes: list[str] = []
        hard: list[str] = []
        title, j_value, carousel, u_value = v("产品标题"), v("预览图"), v("轮播图"), v("产品素材图")
        fingerprint = title_fingerprint(title)
        if not fingerprint:
            hard.append("标题缺少有效末尾四位指纹")
        model = f"{d_value}-1"
        if v("变种属性名称二") != "型号" or v("变种属性值二") != model:
            hard.append("H/I 不等于 型号 / D-1")
        if v("变种名称") != f"{g_value}，{model}":
            hard.append("E 不等于 G，I")
        sku_attrs = parse_list(values[headers["SKU属性"] - 1], "AD SKU属性", row_number, issues)
        if len(sku_attrs) != 2 or (sku_attrs and (text(sku_attrs[0].get("specName")) != g_value or text(sku_attrs[0].get("parentSpecName")) != v("变种属性名称一") or text(sku_attrs[-1].get("specName")) != model or text(sku_attrs[-1].get("parentSpecName")) != "型号")):
            hard.append("AD SKU属性 未与 F/G/H/I 同步")
        skc_attrs = parse_list(values[headers["SKC属性"] - 1], "AC SKC属性", row_number, issues)
        if not skc_attrs:
            hard.append("AC SKC属性为空或无效")
        else:
            skc = skc_attrs[0]
            if text(skc.get("extCode")) != d_value:
                hard.append("AC extCode 不等于 D")
            if text(skc.get("previewImgUrls")) != j_value:
                hard.append("AC previewImgUrls 不等于 J")
        parse_list(values[headers["产品属性"] - 1], "AA 产品属性", row_number, issues)
        parse_list(values[headers["SPU属性"] - 1], "AB SPU属性", row_number, issues)
        # Xuanxshop stores carousel URLs one per line, not comma-separated.
        urls = [part.strip() for part in re.split(r"[\r\n,]+", carousel) if part.strip()]
        if not (4 <= len(urls) <= 10):
            hard.append("T 轮播图数量不在 4-10")
        if urls and u_value != urls[0]:
            hard.append("U 不等于 T 首图")
        if len(urls) >= 4 and not urls[3].startswith(("http://", "https://")):
            hard.append("T4 不是有效 URL")
        if d_value[:4] == "L058":
            if v("重量") != "3500": hard.append("L058 重量不等于 3500")
        elif v("重量") != "150":
            hard.append("非 L058 重量不等于 150")
        try:
            if float(v("申报价格")) <= 0: hard.append("申报价格非正数")
            for name in ("长", "宽", "高"):
                if float(v(name)) <= 0: hard.append(f"{name} 非正数")
        except ValueError:
            hard.append("申报价格或长宽高不是数值")
        asset = asset_index.get((d_value, g_value, sku))
        source_path = text(asset.get("source_path")) if asset else ""
        if not j_value.startswith(("http://", "https://")):
            hard.append("J 缺少长期 URL")
        if not asset:
            hard.append("J 没有精确 D+G+SKU 资产血缘")
        elif text(asset.get("j_value")) != j_value:
            hard.append("当前 J 与精确资产血缘 URL 不一致")
        if d_value in locks:
            hard.append("用户确认的历史错误 D：禁止复用其 J 证据")
        if d_value[:4] == "L095":
            expected = r"\2\2.png" if sku == "L095-00" else (r"\3\3.png" if sku == "L095-01" else "")
            if expected and expected not in source_path:
                hard.append(f"L095 {sku} 数量源错误：必须使用 {expected}")
            if count_claim(g_value):
                notes.append(f"数量声明：{count_claim(g_value)}；禁止唯一 SKU 兜底")
        if hard:
            notes.extend(hard)
        groups[d_value].append({"row": row_number, "D": d_value, "G": g_value, "SKU": sku, "J": j_value, "title": title, "source_path": source_path, "hard": hard, "notes": notes, "status": "硬阻断" if hard else "待视觉审"})
        for item in hard:
            issues.append({"row": row_number, "D": d_value, "G": g_value, "SKU": sku, "layer": "global", "kind": item})
    for d_value, rows in groups.items():
        if len({row["title"] for row in rows}) != 1:
            for row in rows: row["hard"].append("同 D 标题不一致")
        if len({"|".join([row["J"], row["G"], row["SKU"]]) for row in rows}) != len(rows):
            for row in rows: row["hard"].append("同 D 变体行存在重复 J+G+SKU 组合")
        if len({row["J"] for row in rows}) != len(rows):
            for row in rows: row["hard"].append("同 D 不同变体共用 J")
        if len({row["J"] for row in rows}) != len(rows):
            for row in rows: row["notes"].append("同 D J 重复")
    categories = Counter()
    for issue in issues:
        categories[issue["kind"]] += 1
    summary = {
        "rows": sum(len(rows) for rows in groups.values()), "d_groups": len(groups),
        "data_hard": sum(1 for issue in issues if issue["kind"] in {"H/I 不等于 型号 / D-1", "E 不等于 G，I", "AD SKU属性 未与 F/G/H/I 同步", "AC extCode 不等于 D", "AC previewImgUrls 不等于 J", "标题缺少有效末尾四位指纹"}),
        "j_hard": sum(1 for issue in issues if "J" in issue["kind"] or "L095" in issue["kind"] or "历史错误" in issue["kind"]),
        "carousel_hard": sum(1 for issue in issues if issue["kind"].startswith(("T ", "U "))),
        "visual_pending": sum(1 for rows in groups.values() for row in rows if not row["hard"]),
    }
    report = {"workbook": str(WORKBOOK), "scope": "read-only 54-column + J lineage audit", "summary": summary, "issue_counts": dict(categories), "issues": issues, "rules": {"L095-00": r"...\2\2.png", "L095-01": r"...\3\3.png", "visual_gate": "No row is visually approved by this script; source lineage is not visual proof."}}
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    render(groups, report)
    print(json.dumps({"report": str(REPORT), "html": str(HTML), "summary": summary, "issue_counts": dict(categories)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
