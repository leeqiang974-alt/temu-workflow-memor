"""Fail-closed OSS upload and workbook writeback for the reviewed 398-D batch."""
from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import sys
import time
from pathlib import Path

import oss2
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Administrator\Documents\temu自动化")
OUT = ROOT / r"outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718"
SOURCE = OUT / "YeahF_398D_按L0xx顺位归组_待J_T1回填_20260719.xlsx"
T_MANIFEST = OUT / r"t1_badged_local_review_20260719\YeahF_398D_T1_角标本地候选_manifest_20260719.json"
J_MANIFEST = OUT / r"j_full_candidate_review_20260719\YeahF_398D_J_702行候选_manifest_20260719.json"
L095_REPORT = OUT / "YeahF_398D_L095三格J修正_写回报告_20260718.json"
UPLOAD_STATE = OUT / "YeahF_398D_T1_角标OSS上传状态_20260719.json"
FINAL = OUT / "YeahF_398D_按L0xx顺位归组_最终OSS回填_待提交_20260719.xlsx"
REPORT = OUT / "YeahF_398D_最终OSS回填报告_20260719.json"
OSS_PREFIX = "temu-jit/yeahf-398d-final/t1-badged-20260719"
J_OSS_PREFIX = "temu-jit/yeahf-398d-final/j-reviewed-20260719"
J_UPLOAD_STATE = OUT / "YeahF_398D_J_OSS上传状态_20260719.json"

BASE = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui")
sys.path.insert(0, str(BASE / "work"))
import generate_apply_ali_tfirst_0608 as ali_t  # noqa: E402


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def approval_map(path: Path, expected: set[str], label: str) -> dict[str, dict]:
    payload = load_json(path)
    decisions = payload.get("decisions", {})
    missing = sorted(expected - set(decisions))
    extra = sorted(set(decisions) - expected)
    bad = sorted(key for key in expected if decisions.get(key, {}).get("decision") != "approve")
    if missing or extra or bad:
        raise RuntimeError(f"{label} approval rejected: missing={len(missing)}, extra={len(extra)}, non_approve={len(bad)}; examples={bad[:8] or missing[:8]}")
    return decisions


def validate_manifests() -> tuple[dict, dict]:
    t_manifest = load_json(T_MANIFEST)
    j_manifest = load_json(J_MANIFEST)
    source_hash = sha256(SOURCE)
    for label, manifest in (("T1", t_manifest), ("J", j_manifest)):
        if manifest["workbook_sha256"] != source_hash:
            raise RuntimeError(f"{label} manifest workbook hash changed")
    if t_manifest["rendered"] != 398 or t_manifest["missing"]:
        raise RuntimeError("T1 manifest is incomplete")
    if len(j_manifest["records"]) != 702 or j_manifest["status_counts"].get("blocked", 0):
        raise RuntimeError("J manifest is incomplete or blocked")
    for record in t_manifest["records"]:
        rendered = Path(record["rendered_path"])
        if not rendered.is_file() or sha256(rendered) != record["rendered_sha256"]:
            raise RuntimeError(f"T1 candidate changed: {record['D']}")
    for record in j_manifest["records"]:
        if record["candidate_kind"].startswith("local-"):
            local = Path(record["candidate_J"])
            if not local.is_file() or sha256(local) != record["candidate_sha256"]:
                raise RuntimeError(f"J candidate changed: row {record['row']}")
    return t_manifest, j_manifest


def bucket() -> oss2.Bucket:
    config = ali_t.ali.read_oss_config()
    return oss2.Bucket(oss2.Auth(config["access_key_id"], config["access_key_secret"]), f"https://{config['endpoint']}", config["bucket"])


def upload_t1(records: list[dict]) -> dict[str, str]:
    client = bucket()
    state = load_json(UPLOAD_STATE) if UPLOAD_STATE.is_file() else {"uploaded": {}}
    result = {}
    for position, record in enumerate(records, 1):
        d_value = record["D"]
        path = Path(record["rendered_path"])
        digest = record["rendered_sha256"]
        identity = f"sha256:{digest}"
        if identity in state["uploaded"]:
            url = state["uploaded"][identity]["url"]
        else:
            key = f"{OSS_PREFIX}/{d_value[:4]}/{d_value}_{digest[:24]}.png"
            client.put_object(key, path.read_bytes(), headers={"Content-Type": mimetypes.guess_type(path.name)[0] or "image/png"})
            endpoint = client.endpoint.removeprefix("https://").removeprefix("http://")
            url = f"https://{client.bucket_name}.{endpoint}/{key}"
            state["uploaded"][identity] = {"D": d_value, "url": url, "object_key": key, "sha256": digest, "uploaded_at": time.strftime("%Y-%m-%dT%H:%M:%S")}
            UPLOAD_STATE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        result[d_value] = url
        if position % 25 == 0:
            print(json.dumps({"uploaded_or_reused": position, "total": len(records)}, ensure_ascii=False), flush=True)
    return result


def final_j_map(j_manifest: dict) -> dict[int, str]:
    l095_urls = {(record["D"], record["G"], record["SKU"]): record["new_j"] for record in load_json(L095_REPORT)["records"]}
    client = None
    state = load_json(J_UPLOAD_STATE) if J_UPLOAD_STATE.is_file() else {"uploaded": {}}
    result = {}
    for record in j_manifest["records"]:
        row_number = int(record["row"])
        if record["candidate_kind"] == "local-unapproved-candidate":
            identity = (record["D"], record["G"], record["SKU"])
            if identity not in l095_urls:
                raise RuntimeError(f"L095 approved candidate has no uploaded URL: row {row_number}")
            result[row_number] = l095_urls[identity]
        elif record["candidate_kind"].startswith("local-"):
            local = Path(record["candidate_J"])
            digest = record["candidate_sha256"]
            identity = f"sha256:{digest}"
            if identity in state["uploaded"]:
                result[row_number] = state["uploaded"][identity]["url"]
            else:
                if client is None:
                    client = bucket()
                key = f"{J_OSS_PREFIX}/{record['L0xx']}/r{row_number}_{record['D']}_{digest[:24]}.jpg"
                client.put_object(key, local.read_bytes(), headers={"Content-Type": mimetypes.guess_type(local.name)[0] or "image/jpeg"})
                endpoint = client.endpoint.removeprefix("https://").removeprefix("http://")
                url = f"https://{client.bucket_name}.{endpoint}/{key}"
                state["uploaded"][identity] = {"row": row_number, "D": record["D"], "url": url, "object_key": key, "sha256": digest, "uploaded_at": time.strftime("%Y-%m-%dT%H:%M:%S")}
                J_UPLOAD_STATE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
                result[row_number] = url
        else:
            result[row_number] = record["candidate_J"]
    return result


def update_skc(raw: object, d_value: str, j_url: str) -> str:
    value = json.loads(str(raw or "[]"))
    objects = value if isinstance(value, list) else [value]
    if not objects:
        raise ValueError("empty SKC JSON")
    for item in objects:
        if not isinstance(item, dict):
            raise ValueError("SKC JSON member is not an object")
        item["previewImgUrls"] = j_url
        item["extCode"] = d_value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def writeback(t1_urls: dict[str, str], j_urls: dict[int, str]) -> dict:
    workbook = load_workbook(SOURCE)
    ws = workbook.active
    headers = [cell.value for cell in ws[1]]
    index = {name: headers.index(name) + 1 for name in headers}
    changed = {"J": 0, "T": 0, "U": 0, "SKC": 0}
    for row_number in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_number, index["产品货号"]).value or "").strip()
        if not d_value:
            continue
        if d_value not in t1_urls or row_number not in j_urls:
            raise RuntimeError(f"missing approved writeback asset at row {row_number}, D={d_value}")
        t_cell = ws.cell(row_number, index["轮播图"])
        old_t = [item.strip() for item in str(t_cell.value or "").splitlines() if item.strip()]
        if len(old_t) < 4:
            raise RuntimeError(f"row {row_number} has no protected T4 slot")
        new_t = [t1_urls[d_value], *old_t[1:]]
        if len(new_t) > 10:
            raise RuntimeError(f"row {row_number} carousel exceeds 10")
        j_url = j_urls[row_number]
        ws.cell(row_number, index["预览图"]).value = j_url
        t_cell.value = "\n".join(new_t)
        ws.cell(row_number, index["产品素材图"]).value = t1_urls[d_value]
        skc_cell = ws.cell(row_number, index["SKC属性"])
        skc_cell.value = update_skc(skc_cell.value, d_value, j_url)
        changed["J"] += 1
        changed["T"] += 1
        changed["U"] += 1
        changed["SKC"] += 1
    workbook.save(FINAL)
    return changed


def verify_output(t1_urls: dict[str, str], j_urls: dict[int, str]) -> dict:
    ws = load_workbook(FINAL, read_only=True, data_only=False).active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {name: headers.index(name) for name in headers}
    groups: dict[str, list[tuple]] = {}
    issues = []
    effective = 0
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        d_value = str(row[index["产品货号"]] or "").strip()
        if not d_value:
            continue
        effective += 1
        t_urls = [item.strip() for item in str(row[index["轮播图"]] or "").splitlines() if item.strip()]
        j_url = str(row[index["预览图"]] or "")
        u_url = str(row[index["产品素材图"]] or "")
        skc = json.loads(str(row[index["SKC属性"]] or "[]"))
        objects = skc if isinstance(skc, list) else [skc]
        if not t_urls or t_urls[0] != t1_urls[d_value] or u_url != t1_urls[d_value]:
            issues.append(f"row {row_number}: T1/U mismatch")
        if len(t_urls) < 4 or len(t_urls) > 10:
            issues.append(f"row {row_number}: invalid T count {len(t_urls)}")
        if j_url != j_urls[row_number]:
            issues.append(f"row {row_number}: J mismatch")
        if any(item.get("previewImgUrls") != j_url or item.get("extCode") != d_value for item in objects):
            issues.append(f"row {row_number}: SKC linkage mismatch")
        groups.setdefault(d_value, []).append((row[index["产品标题"]], tuple(t_urls), u_url))
    for d_value, values in groups.items():
        if len(set(values)) != 1:
            issues.append(f"{d_value}: same-D title/T/U mismatch")
    if effective != 702 or len(groups) != 398 or issues:
        raise RuntimeError(f"output verification failed: rows={effective}, D={len(groups)}, issues={issues[:12]}")
    return {"effective_rows": effective, "exact_d": len(groups), "issues": issues, "output_sha256": sha256(FINAL)}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--t-decisions", type=Path)
    parser.add_argument("--j-decisions", type=Path)
    parser.add_argument("--preflight-only", action="store_true")
    args = parser.parse_args()
    t_manifest, j_manifest = validate_manifests()
    preflight = {"source": str(SOURCE), "source_sha256": sha256(SOURCE), "t_candidates": len(t_manifest["records"]), "j_candidates": len(j_manifest["records"]), "preflight": "pass"}
    if args.preflight_only:
        print(json.dumps(preflight, ensure_ascii=False, indent=2))
        return
    if not args.t_decisions or not args.j_decisions:
        raise RuntimeError("both --t-decisions and --j-decisions are required")
    approval_map(args.t_decisions, {record["D"] for record in t_manifest["records"]}, "T1")
    approval_map(args.j_decisions, {record["decision_key"] for record in j_manifest["records"]}, "J")
    t1_urls = upload_t1(t_manifest["records"])
    j_urls = final_j_map(j_manifest)
    changed = writeback(t1_urls, j_urls)
    verification = verify_output(t1_urls, j_urls)
    payload = {**preflight, "t_decisions": str(args.t_decisions), "j_decisions": str(args.j_decisions), "changed": changed, "verification": verification, "output": str(FINAL)}
    REPORT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
