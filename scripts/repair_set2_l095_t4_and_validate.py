#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


COMFYUI_BASE = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui")
if str(COMFYUI_BASE / "work") not in sys.path:
    sys.path.insert(0, str(COMFYUI_BASE / "work"))

try:
    from encode_workbook_image_urls import encode_text, scan_url_safety
except Exception:  # pragma: no cover
    encode_text = None
    scan_url_safety = None


URL_RE = re.compile(r"https?://.*?(?=(?:[,;，；]?\s*https?://)|$)", re.S)
SIZE_RE = re.compile(r"(尺寸|尺码|size|cm|inch|length|height|width|宽|长|高)", re.I)
BAD_L095_T4_HASH = "85ef73d852584a65aaf05a1c5615dac7_L095"
L095_T4_FALLBACK = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "f6d7d52f6a3c405096b1cb687e5d0256_L095_%E5%B0%BA%E5%AF%B8_2.jpg"
)
FORBIDDEN_SOURCE_WORDS = ("九宫格", "9grid", "output", "背景素材")
FORBIDDEN_SOURCE_DIRS = {"out", "output", "outputs"}


def headers(ws) -> list[str]:
    return [str(cell.value or "").strip() for cell in ws[1]]


def col(header_values: list[str], name: str) -> int:
    if name not in header_values:
        raise RuntimeError(f"missing required column: {name}")
    return header_values.index(name) + 1


def split_urls(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    urls: list[str] = []
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line:
            continue
        matches = [match.group(0).strip() for match in URL_RE.finditer(line)]
        urls.extend(matches or [line])
    return urls


def dedupe(values: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            out.append(value)
            seen.add(value)
    return out


def is_size_url(url: str) -> bool:
    text = str(url or "")
    if not text:
        return False
    parts = urllib.parse.urlsplit(text)
    decoded_path = urllib.parse.unquote(parts.path or text)
    basename = decoded_path.rsplit("/", 1)[-1]
    return bool(SIZE_RE.search(basename) or "/carousel-ocr-size/" in decoded_path.lower())


def check_url(url: str, timeout: int = 10) -> dict[str, Any]:
    try:
        req = urllib.request.Request(url, method="HEAD", headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return {"ok": 200 <= resp.status < 400, "status": resp.status}
    except Exception as exc:
        return {"ok": False, "error": str(exc)}


def has_forbidden_source_path(path: str) -> bool:
    lower = path.lower()
    if any(word.lower() in lower for word in FORBIDDEN_SOURCE_WORDS):
        return True
    return any(part.lower() in FORBIDDEN_SOURCE_DIRS for part in Path(path).parts)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--j-audit-json", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)

    wb = load_workbook(args.output)
    ws = wb.active
    h = headers(ws)
    col_title = col(h, "产品标题")
    col_d = col(h, "产品货号")
    col_j = col(h, "预览图")
    col_t = col(h, "轮播图")
    col_u = col(h, "产品素材图")
    col_skc = col(h, "SKC属性")

    repaired_rows: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        urls = split_urls(ws.cell(row_idx, col_t).value)
        if d_value == "L095060505" and urls:
            old_urls = list(urls)
            urls = [L095_T4_FALLBACK if BAD_L095_T4_HASH in url else url for url in urls]
            if len(urls) >= 4:
                urls[3] = L095_T4_FALLBACK
            urls = dedupe(urls)
            if len(urls) > 10:
                urls = urls[:10]
            if urls != old_urls:
                ws.cell(row_idx, col_t).value = "\n".join(urls)
                ws.cell(row_idx, col_u).value = urls[0]
                repaired_rows.append(
                    {
                        "row": row_idx,
                        "D": d_value,
                        "old_t4": old_urls[3] if len(old_urls) >= 4 else "",
                        "new_t4": urls[3] if len(urls) >= 4 else "",
                    }
                )

    if encode_text:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "http" in cell.value:
                        cell.value = encode_text(cell.value)[0]

    wb.save(args.output)
    wb.close()

    wb2 = load_workbook(args.output, read_only=True, data_only=False)
    ws2 = wb2.active
    h2 = headers(ws2)
    c_title = col(h2, "产品标题")
    c_d = col(h2, "产品货号")
    c_j = col(h2, "预览图")
    c_t = col(h2, "轮播图")
    c_u = col(h2, "产品素材图")
    c_skc = col(h2, "SKC属性")

    issue_examples: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_d_title: dict[str, set[str]] = defaultdict(set)
    by_d_t: dict[str, set[str]] = defaultdict(set)
    title_code_by_d: dict[str, str] = {}
    effective_rows = 0
    unique_t4: dict[str, dict[str, Any]] = {}

    for row_idx, row in enumerate(ws2.iter_rows(min_row=2, values_only=True), start=2):
        d_value = str(row[c_d - 1] or "").strip()
        if not d_value:
            continue
        effective_rows += 1
        title = str(row[c_title - 1] or "")
        j_value = str(row[c_j - 1] or "").strip()
        t_urls = split_urls(row[c_t - 1])
        u_value = str(row[c_u - 1] or "").strip()
        skc_value = str(row[c_skc - 1] or "")
        by_d_title[d_value].add(title)
        by_d_t[d_value].add("\n".join(t_urls))
        if not j_value:
            issue_examples["empty_j"].append({"row": row_idx, "D": d_value})
        if not t_urls:
            issue_examples["empty_t"].append({"row": row_idx, "D": d_value})
        if len(t_urls) > 10:
            issue_examples["t_over_10"].append({"row": row_idx, "D": d_value, "count": len(t_urls)})
        if len(t_urls) < 6:
            issue_examples["t_below_6"].append({"row": row_idx, "D": d_value, "count": len(t_urls)})
        if t_urls and u_value != t_urls[0]:
            issue_examples["u_not_t1"].append({"row": row_idx, "D": d_value})
        if len(t_urls) < 4 or not is_size_url(t_urls[3]):
            issue_examples["t4_not_size"].append({"row": row_idx, "D": d_value, "T4": t_urls[3] if len(t_urls) >= 4 else ""})
        if t_urls and BAD_L095_T4_HASH in "\n".join(t_urls):
            issue_examples["bad_l095_t4_residual"].append({"row": row_idx, "D": d_value})
        if "previewImgUrls" in skc_value and j_value and j_value not in skc_value:
            issue_examples["skc_preview_not_synced_to_j"].append({"row": row_idx, "D": d_value})
        code_match = re.search(r"\s([A-Z][0-9][A-Z])$", title)
        if not code_match:
            issue_examples["title_code_missing"].append({"row": row_idx, "D": d_value, "title": title})
        else:
            code = code_match.group(1)
            old = title_code_by_d.setdefault(d_value, code)
            if old != code:
                issue_examples["title_code_mismatch"].append({"row": row_idx, "D": d_value, "old": old, "code": code})
        if len(t_urls) >= 4:
            unique_t4.setdefault(t_urls[3], {"url": t_urls[3], "rows": []})["rows"].append(row_idx)

    for d_value, titles in by_d_title.items():
        if len(titles) != 1:
            issue_examples["same_d_title_mismatch"].append({"D": d_value, "count": len(titles)})
    for d_value, values in by_d_t.items():
        if len(values) != 1:
            issue_examples["same_d_t_mismatch"].append({"D": d_value, "count": len(values)})

    t4_unreachable: list[dict[str, Any]] = []
    for item in unique_t4.values():
        check = check_url(item["url"])
        if not check.get("ok"):
            t4_unreachable.append({"rows": item["rows"][:5], "url": item["url"], "check": check})

    j_payload = json.loads(args.j_audit_json.read_text(encoding="utf-8-sig"))
    j_records = j_payload.get("records", j_payload) if isinstance(j_payload, dict) else j_payload
    j_missing_source = [r for r in j_records if not r.get("sku_source")]
    j_missing_file = [r for r in j_records if r.get("sku_source") and not Path(str(r["sku_source"])).exists()]
    j_forbidden_source = [r for r in j_records if has_forbidden_source_path(str(r.get("sku_source", "")))]

    issue_counts = {f"{key}_count": len(value) for key, value in sorted(issue_examples.items())}
    required_zero_keys = [
        "empty_j_count",
        "empty_t_count",
        "t_over_10_count",
        "t_below_6_count",
        "u_not_t1_count",
        "t4_not_size_count",
        "bad_l095_t4_residual_count",
        "skc_preview_not_synced_to_j_count",
        "title_code_missing_count",
        "title_code_mismatch_count",
        "same_d_title_mismatch_count",
        "same_d_t_mismatch_count",
    ]
    for key in required_zero_keys:
        issue_counts.setdefault(key, 0)
    issue_counts.update(
        {
            "t4_unreachable_count": len(t4_unreachable),
            "j_missing_source_count": len(j_missing_source),
            "j_missing_file_count": len(j_missing_file),
            "j_forbidden_source_count": len(j_forbidden_source),
        }
    )
    url_safety = scan_url_safety(args.output) if scan_url_safety else {"skipped": True}
    if isinstance(url_safety, dict) and "unsafe_url_occurrences_after" in url_safety:
        issue_counts["unsafe_url_occurrences_after"] = url_safety["unsafe_url_occurrences_after"]

    pass_hard_checks = all(value == 0 for value in issue_counts.values() if isinstance(value, int))
    report = {
        "source": str(args.source),
        "output": str(args.output),
        "effective_rows": effective_rows,
        "unique_d": len(by_d_t),
        "repaired_rows": repaired_rows,
        "l095_t4_fallback": L095_T4_FALLBACK,
        "issue_counts": issue_counts,
        "pass_hard_checks": pass_hard_checks,
        "url_safety": url_safety,
        "t4_unreachable_examples": t4_unreachable[:20],
        "j_audit_summary": {
            "records": len(j_records),
            "missing_source": len(j_missing_source),
            "missing_file": len(j_missing_file),
            "forbidden_source": len(j_forbidden_source),
        },
        "examples": {key: value[:20] for key, value in issue_examples.items()},
    }
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    wb2.close()
    return 0 if pass_hard_checks else 2


if __name__ == "__main__":
    raise SystemExit(main())
