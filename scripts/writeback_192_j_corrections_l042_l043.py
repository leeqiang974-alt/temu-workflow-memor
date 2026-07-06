#!/usr/bin/env python3
"""Write approved L042/L043 J corrections into a copied workbook.

Scope:
- Source workbook is copied, never modified in place.
- Only approved corrected J rows from the combined review are written.
- `预览图` and matching `SKC属性.previewImgUrls` may change.
- D/G/SKU/title/T/U and other protected columns must remain unchanged.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
import time
import uuid
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import oss2
from openpyxl import load_workbook
from PIL import Image, ImageOps


COMFYUI_BASE = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui")
if str(COMFYUI_BASE / "work") not in sys.path:
    sys.path.insert(0, str(COMFYUI_BASE / "work"))

import generate_apply_ali_tfirst_0608 as ali_t


OSS_PREFIX = "temu-jit/dxxmall-0616-2/192-set1-j-corrections-l042-l043"
TARGET_SIZE = 800
ALLOWED_CHANGED_HEADERS = {"预览图", "SKC属性"}


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def save_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def headers(ws) -> list[str]:
    return [str(ws.cell(1, column).value or "").strip() for column in range(1, ws.max_column + 1)]


def require_col(header: list[str], name: str) -> int:
    if name not in header:
        raise RuntimeError(f"missing required column: {name}")
    return header.index(name) + 1


def snapshot(ws) -> dict[tuple[int, int], Any]:
    return {
        (row, col): ws.cell(row, col).value
        for row in range(1, ws.max_row + 1)
        for col in range(1, ws.max_column + 1)
    }


def load_manifest(path: Path) -> dict[str, Any]:
    if path.exists():
        return load_json(path)
    return {"uploaded": {}}


def upload_once(bucket: oss2.Bucket, manifest_path: Path, manifest: dict[str, Any], row_key: str, image_path: Path) -> str:
    cache_key = f"{row_key}|{image_path.resolve()}"
    uploaded = manifest.setdefault("uploaded", {})
    if cache_key in uploaded:
        return uploaded[cache_key]
    object_key = f"{OSS_PREFIX}/{time.strftime('%Y%m%d')}/{uuid.uuid4().hex}_{row_key}_j_800.jpg"
    bucket.put_object_from_file(object_key, str(image_path), headers={"Content-Type": "image/jpeg"})
    endpoint = bucket.endpoint.replace("https://", "").replace("http://", "")
    url = f"https://{bucket.bucket_name}.{endpoint}/{object_key}"
    uploaded[cache_key] = url
    save_json(manifest_path, manifest)
    return url


def update_skc_preview(value: object, preview_url: str) -> tuple[str, bool]:
    text = str(value or "").strip()
    if not text:
        return text, False
    try:
        payload = json.loads(text)
    except Exception:
        return text, False
    changed = False

    def visit(node: Any) -> None:
        nonlocal changed
        if isinstance(node, dict):
            if "previewImgUrls" in node:
                current = node.get("previewImgUrls")
                replacement: Any = [preview_url] if isinstance(current, list) else preview_url
                if current != replacement:
                    node["previewImgUrls"] = replacement
                    changed = True
            for child in node.values():
                visit(child)
        elif isinstance(node, list):
            for child in node:
                visit(child)

    visit(payload)
    if not changed:
        return text, False
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":")), True


def image_check(path: Path) -> dict[str, Any]:
    image = ImageOps.exif_transpose(Image.open(path))
    return {"path": str(path), "width": image.width, "height": image.height, "mode": image.mode, "bytes": path.stat().st_size}


def build_correction_map(combined_json: Path) -> dict[int, dict[str, Any]]:
    records = load_json(combined_json)["records"]
    corrections = {
        int(r["row"]): r
        for r in records
        if r.get("status") in {"corrected_l042_final_cutout", "corrected_l043_quantity"}
    }
    if not corrections:
        raise RuntimeError("no approved corrected J records found")
    return corrections


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--combined-json", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--output-name", required=True)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    output_path = args.out_dir / args.output_name
    shutil.copy2(args.source, output_path)

    corrections = build_correction_map(args.combined_json)
    image_issues: list[dict[str, Any]] = []
    for row, record in sorted(corrections.items()):
        path = Path(str(record.get("review_image") or ""))
        if not path.exists():
            image_issues.append({"row": row, "D": record.get("D"), "issue": "missing_review_image", "path": str(path)})
            continue
        info = image_check(path)
        if info["width"] != TARGET_SIZE or info["height"] != TARGET_SIZE:
            image_issues.append({"row": row, "D": record.get("D"), "issue": "wrong_image_size", **info})
    if image_issues:
        raise RuntimeError(json.dumps({"image_issues": image_issues[:20]}, ensure_ascii=False, indent=2))

    oss_config = ali_t.ali.read_oss_config()
    bucket = oss2.Bucket(
        oss2.Auth(oss_config["access_key_id"], oss_config["access_key_secret"]),
        f"https://{oss_config['endpoint']}",
        oss_config["bucket"],
    )
    upload_manifest_path = args.out_dir / "j_corrections_upload_manifest.json"
    upload_manifest = load_manifest(upload_manifest_path)

    wb = load_workbook(output_path)
    ws = wb.active
    header = headers(ws)
    c_d = require_col(header, "产品货号")
    c_g = require_col(header, "变种属性值一")
    c_j = require_col(header, "预览图")
    c_sku = require_col(header, "SKU货号")
    c_skc = require_col(header, "SKC属性")
    c_t = require_col(header, "轮播图")
    c_u = require_col(header, "产品素材图")
    before = snapshot(ws)

    changed_cells: list[dict[str, Any]] = []
    write_records: list[dict[str, Any]] = []
    uploaded_by_row: dict[int, str] = {}
    for row_idx, record in sorted(corrections.items()):
        d_value = str(ws.cell(row_idx, c_d).value or "").strip()
        g_value = str(ws.cell(row_idx, c_g).value or "").strip()
        sku_value = str(ws.cell(row_idx, c_sku).value or "").strip()
        if d_value != str(record.get("D") or "").strip():
            raise RuntimeError(f"row {row_idx} D mismatch: workbook={d_value} review={record.get('D')}")
        if g_value != str(record.get("G") or "").strip():
            raise RuntimeError(f"row {row_idx} G mismatch: workbook={g_value} review={record.get('G')}")
        if sku_value != str(record.get("SKU") or "").strip():
            raise RuntimeError(f"row {row_idx} SKU mismatch: workbook={sku_value} review={record.get('SKU')}")

        local_image = Path(str(record["review_image"]))
        url = upload_once(bucket, upload_manifest_path, upload_manifest, f"{d_value}_r{row_idx}_{record['status']}", local_image)
        uploaded_by_row[row_idx] = url

        old_j = ws.cell(row_idx, c_j).value
        ws.cell(row_idx, c_j).value = url
        changed_cells.append({"row": row_idx, "D": d_value, "column": "预览图", "old": old_j, "new": url})

        old_skc = ws.cell(row_idx, c_skc).value
        new_skc, skc_changed = update_skc_preview(old_skc, url)
        if skc_changed:
            ws.cell(row_idx, c_skc).value = new_skc
            changed_cells.append({"row": row_idx, "D": d_value, "column": "SKC属性", "old": old_skc, "new": new_skc})

        write_records.append({
            "row": row_idx,
            "D": d_value,
            "G": g_value,
            "SKU": sku_value,
            "status": record["status"],
            "sku_source": record.get("sku_source"),
            "local_image": str(local_image),
            "uploaded_url": url,
            "skc_synced": skc_changed,
        })

    wb.save(output_path)
    wb.close()

    wb2 = load_workbook(output_path, read_only=False, data_only=False)
    ws2 = wb2.active
    after_header = headers(ws2)
    protected_changes: list[dict[str, Any]] = []
    actual_changes: list[dict[str, Any]] = []
    empty_j: list[dict[str, Any]] = []
    t_changed: list[dict[str, Any]] = []
    u_changed: list[dict[str, Any]] = []
    row_identity_mismatch: list[dict[str, Any]] = []
    for row in range(1, ws2.max_row + 1):
        for col in range(1, ws2.max_column + 1):
            old = before.get((row, col))
            new = ws2.cell(row, col).value
            if old != new:
                column_name = after_header[col - 1] if col <= len(after_header) else f"col{col}"
                actual_changes.append({"row": row, "column": column_name, "old": old, "new": new})
                if column_name not in ALLOWED_CHANGED_HEADERS:
                    protected_changes.append({"row": row, "column": column_name, "old": old, "new": new})
    for row in range(2, ws2.max_row + 1):
        d_value = str(ws2.cell(row, c_d).value or "").strip()
        if not d_value:
            continue
        if not str(ws2.cell(row, c_j).value or "").strip():
            empty_j.append({"row": row, "D": d_value})
        if before[(row, c_t)] != ws2.cell(row, c_t).value:
            t_changed.append({"row": row, "D": d_value})
        if before[(row, c_u)] != ws2.cell(row, c_u).value:
            u_changed.append({"row": row, "D": d_value})
        for col, name in ((c_d, "产品货号"), (c_g, "变种属性值一"), (c_sku, "SKU货号")):
            if before[(row, col)] != ws2.cell(row, col).value:
                row_identity_mismatch.append({"row": row, "column": name})
    wb2.close()

    changed_counter = Counter(item["column"] for item in actual_changes)
    status_counter = Counter(item["status"] for item in write_records)
    prefix_counter = Counter(item["D"][:4] for item in write_records)
    validation = {
        "source": str(args.source),
        "output": str(output_path),
        "combined_review": str(args.combined_json),
        "effective_rows": ws.max_row - 1,
        "changed_rows": len(write_records),
        "changed_cell_count": len(actual_changes),
        "changed_columns": dict(changed_counter),
        "status_counts": dict(status_counter),
        "prefix_counts": dict(prefix_counter),
        "empty_j_count": len(empty_j),
        "protected_cell_changed_count": len(protected_changes),
        "row_identity_mismatch_count": len(row_identity_mismatch),
        "t_changed_count": len(t_changed),
        "u_changed_count": len(u_changed),
        "uploaded_url_count": len(uploaded_by_row),
        "pass_hard_checks": not any([empty_j, protected_changes, row_identity_mismatch, t_changed, u_changed]),
    }
    report = {
        "validation": validation,
        "write_records": write_records,
        "changed_cells": actual_changes,
        "issues": {
            "empty_j": empty_j[:20],
            "protected_changes": protected_changes[:20],
            "row_identity_mismatch": row_identity_mismatch[:20],
            "t_changed": t_changed[:20],
            "u_changed": u_changed[:20],
        },
    }
    save_json(args.out_dir / "j_corrections_writeback_report.json", report)
    save_json(args.out_dir / "j_corrections_validation_report.json", validation)
    save_json(args.out_dir / "j_corrections_cell_diff_summary.json", actual_changes)
    print(json.dumps(validation, ensure_ascii=False, indent=2))
    return 0 if validation["pass_hard_checks"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
