#!/usr/bin/env python3
"""Finalize the 20260704 195-D workbook after fixed set2 T writeback.

This repairs the known L095 T4 truncated OSS URL, validates workbook-level T/J/U
contracts, reruns URL/T4 display checks, and writes a compact HTML audit page.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from encode_workbook_image_urls import SIZE_RE, scan_url_safety, validate_t4


BROKEN_L095_T4 = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "85ef73d852584a65aaf05a1c5615dac7_L095_%E5%B0%BA%E5%AF%B8_jimeng-2026-05-08-1613-"
    "%E8%BF%99%E5%BC%A0%E5%9B%BE%E7%89%87"
)

FULL_L095_T4 = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "85ef73d852584a65aaf05a1c5615dac7_L095_%E5%B0%BA%E5%AF%B8_jimeng-2026-05-08-1613-"
    "%E8%BF%99%E5%BC%A0%E5%9B%BE%E7%89%87%EF%BC%8C%E4%BF%9D%E6%8C%81%E4%BA%A7%E5%93%81"
    "%E5%92%8C%E6%A0%87%E5%B0%BA%E3%80%81%E6%95%B0%E5%AD%97%E3%80%81%E6%96%87%E5%AD%97"
    "%E4%BF%A1%E6%81%AF%E4%B8%8D%E5%81%9A.jpg"
)

URL_RE = re.compile(r"https?://[^\r\n\"'<>]+")


def headers(ws) -> list[str]:
    return [str(ws.cell(1, column).value or "").strip() for column in range(1, ws.max_column + 1)]


def col(header_values: list[str], name: str) -> int:
    if name not in header_values:
        raise RuntimeError(f"missing required column: {name}")
    return header_values.index(name) + 1


def split_lines(value: object) -> list[str]:
    return [
        line.strip()
        for line in str(value or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
        if line.strip()
    ]


def repair_l095_t4(input_path: Path, output_path: Path) -> list[dict[str, Any]]:
    shutil.copy2(input_path, output_path)
    wb = load_workbook(output_path)
    repairs: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or BROKEN_L095_T4 not in cell.value:
                    continue
                old = cell.value
                cell.value = old.replace(BROKEN_L095_T4, FULL_L095_T4)
                repairs.append(
                    {
                        "sheet": ws.title,
                        "cell": cell.coordinate,
                        "old_contains": BROKEN_L095_T4,
                        "new_contains": FULL_L095_T4,
                    }
                )
    wb.save(output_path)
    wb.close()
    return repairs


def validate_workbook(workbook_path: Path, writeback_report_path: Path) -> dict[str, Any]:
    writeback_report = json.loads(writeback_report_path.read_text(encoding="utf-8-sig"))
    expected_first_by_d = {item["d"]: item["new_t1"] for item in writeback_report["changes"]}
    expected_set_by_d = {item["d"]: int(item["chosen_set"]) for item in writeback_report["changes"]}

    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    ws = wb.active
    h = headers(ws)
    col_d = col(h, "产品货号")
    col_title = col(h, "产品标题")
    col_j = col(h, "预览图")
    col_t = col(h, "轮播图")
    col_u = col(h, "产品素材图")

    by_d: dict[str, list[dict[str, Any]]] = defaultdict(list)
    effective_rows = 0
    empty_j = []
    empty_t = []
    t_over_10 = []
    t_below_6 = []
    u_not_t1 = []
    t4_missing = []
    t4_not_size = []
    first_not_set2 = []
    first_not_expected = []
    title_code_re = re.compile(r"[A-Z][0-9][A-Z]")

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d_value = str(row[col_d - 1] or "").strip()
        if not d_value:
            continue
        effective_rows += 1
        title = str(row[col_title - 1] or "").strip()
        j_value = str(row[col_j - 1] or "").strip()
        t_urls = split_lines(row[col_t - 1])
        u_value = str(row[col_u - 1] or "").strip()
        t1 = t_urls[0] if t_urls else ""
        t4 = t_urls[3] if len(t_urls) >= 4 else ""
        record = {"row": row_idx, "D": d_value, "title": title, "J": j_value, "T": t_urls, "U": u_value}
        by_d[d_value].append(record)

        if not j_value:
            empty_j.append({"row": row_idx, "D": d_value})
        if not t_urls:
            empty_t.append({"row": row_idx, "D": d_value})
        if len(t_urls) > 10:
            t_over_10.append({"row": row_idx, "D": d_value, "count": len(t_urls)})
        if len(t_urls) < 6:
            t_below_6.append({"row": row_idx, "D": d_value, "count": len(t_urls)})
        if t1 and u_value != t1:
            u_not_t1.append({"row": row_idx, "D": d_value, "U": u_value, "T1": t1})
        if not t4:
            t4_missing.append({"row": row_idx, "D": d_value})
        elif not SIZE_RE.search(t4):
            t4_not_size.append({"row": row_idx, "D": d_value, "T4": t4})
        if expected_set_by_d.get(d_value) != 2:
            first_not_set2.append({"row": row_idx, "D": d_value, "chosen_set": expected_set_by_d.get(d_value)})
        if expected_first_by_d.get(d_value) and t1 != expected_first_by_d[d_value]:
            first_not_expected.append({"row": row_idx, "D": d_value, "T1": t1, "expected": expected_first_by_d[d_value]})

    title_mismatch = []
    t_mismatch = []
    title_code_missing = []
    title_code_mismatch = []
    for d_value, rows in by_d.items():
        titles = {row["title"] for row in rows}
        t_values = {"\n".join(row["T"]) for row in rows}
        if len(titles) > 1:
            title_mismatch.append({"D": d_value, "rows": [row["row"] for row in rows]})
        if len(t_values) > 1:
            t_mismatch.append({"D": d_value, "rows": [row["row"] for row in rows]})
        codes = {match.group(0) for row in rows for match in [title_code_re.search(row["title"])] if match}
        if not codes:
            title_code_missing.append({"D": d_value, "rows": [row["row"] for row in rows]})
        elif len(codes) > 1:
            title_code_mismatch.append({"D": d_value, "codes": sorted(codes), "rows": [row["row"] for row in rows]})

    t4_validation = validate_t4(workbook_path)
    url_safety = scan_url_safety(workbook_path)

    issue_counts = {
        "empty_j_count": len(empty_j),
        "empty_t_count": len(empty_t),
        "t_over_10_count": len(t_over_10),
        "t_below_6_count": len(t_below_6),
        "u_not_t1_count": len(u_not_t1),
        "same_d_title_mismatch_count": len(title_mismatch),
        "same_d_t_mismatch_count": len(t_mismatch),
        "t4_missing_count": len(t4_missing),
        "t4_not_size_count": len(t4_not_size),
        "first_not_set2_count": len(first_not_set2),
        "first_not_expected_count": len(first_not_expected),
        "title_code_missing_count": len(title_code_missing),
        "title_code_mismatch_count": len(title_code_mismatch),
        "writeback_validation_issues_count": len(writeback_report.get("validation_issues") or []),
        "t4_unreachable_count": t4_validation["t4_unreachable_count"],
        "unsafe_url_occurrences_after": url_safety["unsafe_url_occurrences_after"],
    }
    pass_hard_checks = all(value == 0 for value in issue_counts.values())

    return {
        "workbook": str(workbook_path),
        "writeback_report": str(writeback_report_path),
        "effective_rows": effective_rows,
        "unique_d": len(by_d),
        "chosen_set_dist": writeback_report.get("chosen_set_dist", {}),
        "fixed_first_set": writeback_report.get("fixed_first_set"),
        "t_count_dist": dict(Counter(len(rows[0]["T"]) for rows in by_d.values())),
        "issue_counts": issue_counts,
        "pass_hard_checks": pass_hard_checks,
        "examples": {
            "empty_j": empty_j[:20],
            "empty_t": empty_t[:20],
            "t_over_10": t_over_10[:20],
            "t_below_6": t_below_6[:20],
            "u_not_t1": u_not_t1[:20],
            "same_d_title_mismatch": title_mismatch[:20],
            "same_d_t_mismatch": t_mismatch[:20],
            "t4_missing": t4_missing[:20],
            "t4_not_size": t4_not_size[:20],
            "first_not_set2": first_not_set2[:20],
            "first_not_expected": first_not_expected[:20],
            "title_code_missing": title_code_missing[:20],
            "title_code_mismatch": title_code_mismatch[:20],
        },
        "t4_validation": t4_validation,
        "url_safety": url_safety,
    }


def build_audit_html(workbook_path: Path, report_path: Path, validation: dict[str, Any], output_path: Path) -> None:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    ws = wb.active
    h = headers(ws)
    col_d = col(h, "产品货号")
    col_title = col(h, "产品标题")
    col_t = col(h, "轮播图")
    seen: set[str] = set()
    sections = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d_value = str(row[col_d - 1] or "").strip()
        if not d_value or d_value in seen:
            continue
        seen.add(d_value)
        title = str(row[col_title - 1] or "").strip()
        urls = split_lines(row[col_t - 1])
        imgs = "".join(
            "<figure>"
            f"<figcaption>T{i + 1}{' 尺寸图' if i == 3 else ''}</figcaption>"
            f"<img loading='lazy' src='{html.escape(url)}'>"
            f"<div class='url'>{html.escape(url)}</div>"
            "</figure>"
            for i, url in enumerate(urls[:6])
        )
        sections.append(
            f"<section><h2>{html.escape(d_value)} <span>row {row_idx}</span></h2>"
            f"<p class='title'>{html.escape(title)}</p><div class='grid'>{imgs}</div></section>"
        )
    wb.close()

    output_path.write_text(
        f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>195 set2 T writeback final audit</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;background:#f5f6f8;color:#172033;margin:20px}}
.summary,section{{background:#fff;border:1px solid #d8dee9;border-radius:8px;padding:14px;margin:12px 0}}
h1{{font-size:22px;margin:0 0 10px}} h2{{font-size:17px;margin:0 0 6px}} h2 span{{font-size:12px;color:#667085}}
.ok{{color:#047857;font-weight:700}} .bad{{color:#b42318;font-weight:700}}
.meta,.title{{word-break:break-all;font-size:13px;color:#374151}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:10px}}
figure{{margin:0;border:1px solid #e5e7eb;border-radius:6px;padding:8px;background:#fafafa}}
figcaption{{font-size:12px;font-weight:700;margin-bottom:6px}} img{{width:100%;aspect-ratio:1/1;object-fit:cover;background:white;border:1px solid #edf0f5}}
.url{{font-size:10px;color:#667085;word-break:break-all;margin-top:5px;max-height:42px;overflow:auto}}
pre{{white-space:pre-wrap;background:#f3f4f6;padding:10px;border-radius:6px;font-size:12px}}
</style></head><body>
<div class="summary">
<h1>195 表 197x3 第二套 T 首图回填最终审核</h1>
<p class="meta">workbook: {html.escape(str(workbook_path))}</p>
<p class="meta">validation: {html.escape(str(report_path))}</p>
<p>hard checks: <span class="{'ok' if validation['pass_hard_checks'] else 'bad'}">{validation['pass_hard_checks']}</span></p>
<pre>{html.escape(json.dumps(validation['issue_counts'], ensure_ascii=False, indent=2))}</pre>
</div>
{''.join(sections)}
</body></html>""",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--writeback-report", required=True, type=Path)
    parser.add_argument("--repair-report", required=True, type=Path)
    parser.add_argument("--validation-report", required=True, type=Path)
    parser.add_argument("--audit-html", required=True, type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    repairs = repair_l095_t4(args.input, args.output)
    validation = validate_workbook(args.output, args.writeback_report)
    repair_report = {
        "input": str(args.input),
        "output": str(args.output),
        "broken_url": BROKEN_L095_T4,
        "full_url": FULL_L095_T4,
        "repair_count": len(repairs),
        "repairs": repairs,
    }
    args.repair_report.write_text(json.dumps(repair_report, ensure_ascii=False, indent=2), encoding="utf-8")
    args.validation_report.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    build_audit_html(args.output, args.validation_report, validation, args.audit_html)
    print(
        json.dumps(
            {
                "output": str(args.output),
                "repair_count": len(repairs),
                "validation_report": str(args.validation_report),
                "audit_html": str(args.audit_html),
                "pass_hard_checks": validation["pass_hard_checks"],
                "issue_counts": validation["issue_counts"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if validation["pass_hard_checks"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
