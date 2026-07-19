"""Read-only 54-column and linkage audit for the post-T1-exclusion YeahF draft.

This audit distinguishes real cross-cell data defects from deliberately pending T1/U
writeback and historical/system metadata. It never edits the workbook.
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Administrator\Documents\temu自动化\outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718")
SOURCE = ROOT / "YeahF_400D_结构底稿_剔除本轮17D_补17D_待J_T1.xlsx"
BOOK = ROOT / "YeahF_398D_剔除T1不可用L085两组_待J_T1回填.xlsx"
OUT = ROOT / "YeahF_398D_54列与联动审计_20260718.json"

REMOVED = {"L085070902", "L085070903"}
FP_RE = re.compile(r"(?=.*[A-Z])(?=.*\d)[A-Z0-9]{4}$")
D_RE = re.compile(r"^L\d{3}0709\d{2}$")


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def is_url(value: object) -> bool:
    parsed = urlparse(clean(value))
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def urls(value: object) -> list[str]:
    return [item.strip() for item in re.split(r"[\r\n,]+", clean(value)) if item.strip()]


def load_records(path: Path):
    ws = load_workbook(path, read_only=True, data_only=False).active
    values = ws.iter_rows(values_only=True)
    headers = [clean(value) for value in next(values)]
    rows = []
    for no, raw in enumerate(values, 2):
        row = dict(zip(headers, raw))
        if clean(row["产品货号"]):
            row["_row"] = no
            row["_key"] = "|".join(clean(row[x]) for x in ("产品货号", "变种属性值一", "SKU货号"))
            rows.append(row)
    return headers, rows


def parse_skc(value: object):
    try:
        data = json.loads(clean(value))
    except Exception:
        return [], []
    ext, preview = [], []
    def walk(node):
        if isinstance(node, list):
            for item in node: walk(item)
        elif isinstance(node, dict):
            if "extCode" in node: ext.append(clean(node["extCode"]))
            if "previewImgUrls" in node: preview.append(clean(node["previewImgUrls"]))
            for item in node.values(): walk(item)
    walk(data)
    return ext, preview


def parse_sku(value: object):
    try:
        data = json.loads(clean(value))
        return {clean(item.get("parentSpecName")): clean(item.get("specName")) for item in data if isinstance(item, dict)}
    except Exception:
        return None


def main() -> None:
    headers, source_rows = load_records(SOURCE)
    target_headers, rows = load_records(BOOK)
    fatal, pending, metadata = [], [], []
    if len(headers) != 54 or headers != target_headers:
        fatal.append({"kind": "54_header_contract", "detail": {"source": headers, "target": target_headers}})

    source_by_key = {row["_key"]: row for row in source_rows}
    target_by_key = {row["_key"]: row for row in rows}
    if len(source_by_key) != len(source_rows) or len(target_by_key) != len(rows):
        fatal.append({"kind": "duplicate_D_G_SKU_key"})
    unexpected = sorted(set(target_by_key) - set(source_by_key))
    expected_remaining = {key for key, row in source_by_key.items() if clean(row["产品货号"]) not in REMOVED}
    if unexpected or set(target_by_key) != expected_remaining:
        fatal.append({"kind": "delete_diff_not_exact", "unexpected": unexpected[:20], "missing": sorted(expected_remaining-set(target_by_key))[:20]})
    # The deletion copy must be cell-identical for every remaining physical row.
    cell_changes = []
    for key, row in target_by_key.items():
        old = source_by_key[key]
        changed = [header for header in headers if clean(row[header]) != clean(old[header])]
        if changed: cell_changes.append({"row": row["_row"], "key": key, "changed_columns": changed})
    if cell_changes:
        fatal.append({"kind": "unexpected_remaining_cell_change", "count": len(cell_changes), "samples": cell_changes[:20]})

    groups = defaultdict(list)
    fp_to_d = defaultdict(set)
    column_stats = []
    for header in headers:
        vals = [clean(row[header]) for row in rows if clean(row[header])]
        column_stats.append({"column": header, "nonempty": len(vals), "distinct_nonempty": len(set(vals)), "status": "checked"})

    for row in rows:
        d = clean(row["产品货号"]); groups[d].append(row)
        title, g, f, h, i = (clean(row[x]) for x in ("产品标题", "变种属性值一", "变种属性名称一", "变种属性名称二", "变种属性值二"))
        fp = title[-4:].upper()
        if not D_RE.fullmatch(d): fatal.append({"row": row["_row"], "D": d, "kind": "D_pattern"})
        if not FP_RE.fullmatch(fp) or d.lower() in title.lower(): fatal.append({"row": row["_row"], "D": d, "kind": "title_fingerprint"})
        else: fp_to_d[fp].add(d)
        if not f or not g or h != "型号" or i != f"{d}-1" or clean(row["变种名称"]) != f"{g}，{i}":
            fatal.append({"row": row["_row"], "D": d, "kind": "E_to_I_variant_link"})
        sku = clean(row["SKU货号"])
        if not sku.startswith(d[:4] + "-"):
            fatal.append({"row": row["_row"], "D": d, "kind": "SKU_prefix"})
        skc_ext, previews = parse_skc(row["SKC属性"])
        if not skc_ext or any(x != d for x in skc_ext): fatal.append({"row": row["_row"], "D": d, "kind": "SKC_extCode_to_D"})
        if not previews or any(x != clean(row["预览图"]) for x in previews): fatal.append({"row": row["_row"], "D": d, "kind": "SKC_preview_to_J"})
        sku_specs = parse_sku(row["SKU属性"])
        if sku_specs is None or sku_specs.get(f) != g or sku_specs.get("型号") != i:
            fatal.append({"row": row["_row"], "D": d, "kind": "SKU_JSON_to_FGHI"})
        if not is_url(row["预览图"]): fatal.append({"row": row["_row"], "D": d, "kind": "J_url"})
        for field in ("申报价格", "长", "宽", "高", "重量"):
            try: good = float(row[field]) > 0
            except (ValueError, TypeError): good = False
            if not good: fatal.append({"row": row["_row"], "D": d, "kind": f"invalid_{field}"})
        if float(row["重量"]) != (3500 if d.startswith("L058") else 150):
            fatal.append({"row": row["_row"], "D": d, "kind": "weight_contract"})
        carousel = urls(row["轮播图"])
        if not (1 <= len(carousel) <= 10) or not all(is_url(url) for url in carousel):
            fatal.append({"row": row["_row"], "D": d, "kind": "T_url_or_count"})
        elif clean(row["产品素材图"]) != carousel[0]:
            pending.append({"row": row["_row"], "D": d, "kind": "T1_U_writeback_pending"})
        for field in ("产品描述", "外包装形状", "外包装类型", "外包装图片", "分类id", "产品属性", "SPU属性", "来源url", "产地", "SKU分类", "SKU分类数量", "SKU分类单位", "净含量", "SKU分类总数量", "SKU分类总数量单位", "总净含量", "包装清单"):
            if not clean(row[field]): fatal.append({"row": row["_row"], "D": d, "kind": f"required_{field}_empty"})
        for field in ("外包装图片", "来源url"):
            if not is_url(row[field]): fatal.append({"row": row["_row"], "D": d, "kind": f"invalid_{field}_url"})
        for field in ("产品属性", "SPU属性", "包装清单"):
            try: json.loads(clean(row[field]))
            except Exception: fatal.append({"row": row["_row"], "D": d, "kind": f"invalid_{field}_json"})
        video = clean(row["视频Url"])
        if video and not is_url(video): fatal.append({"row": row["_row"], "D": d, "kind": "video_url"})

    for d, group in groups.items():
        for field in ("产品标题", "变种属性名称一", "轮播图", "产品素材图", "分类id", "产品属性", "SPU属性", "产地"):
            if len({clean(row[field]) for row in group}) != 1:
                fatal.append({"D": d, "kind": f"same_D_{field}_drift"})
        if len({clean(row["变种属性值一"]) for row in group}) != len(group):
            fatal.append({"D": d, "kind": "same_D_duplicate_G"})
    for fp, ds in fp_to_d.items():
        if len(ds) > 1: fatal.append({"D": sorted(ds), "kind": "duplicate_fingerprint", "fingerprint": fp})

    # System values are informational only before upload, but must not be treated as a routed shop identity.
    metadata.append({"kind": "store_field_not_routing_authority", "values": sorted({clean(row["所属店铺"]) for row in rows})})
    metadata.append({"kind": "post_upload_columns_present", "SKCID_values": sorted({clean(row["SKCID"]) for row in rows}), "SKUID_values": sorted({clean(row["SKUID"]) for row in rows}), "creation_count": sum(bool(clean(row["创建时间"])) for row in rows), "update_count": sum(bool(clean(row["更新时间"])) for row in rows)})
    metadata.append({"kind": "video_blank_rows", "count": sum(not bool(clean(row["视频Url"])) for row in rows)})

    report = {
        "workbook": str(BOOK), "source_for_delete_diff": str(SOURCE), "columns": len(headers), "effective_rows": len(rows), "exact_D": len(groups),
        "removed_D": sorted(REMOVED), "fatal_issue_counts": dict(Counter(item["kind"] for item in fatal)), "fatal_issues": fatal,
        "pending_writeback_counts": dict(Counter(item["kind"] for item in pending)), "pending_writeback": pending[:50],
        "metadata_confirmation": metadata, "column_stats": column_stats,
        "verdict": "STRUCTURAL_LINKS_PASS_T1_U_PENDING" if not fatal else "STRUCTURAL_LINKS_BLOCKED",
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"rows": len(rows), "D": len(groups), "fatal": dict(Counter(i["kind"] for i in fatal)), "pending": dict(Counter(i["kind"] for i in pending)), "out": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
