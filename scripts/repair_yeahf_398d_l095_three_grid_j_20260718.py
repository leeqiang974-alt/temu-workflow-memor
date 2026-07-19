"""Render local-only corrected J candidates for L095 SKU L095-01 (three-grid)."""
from __future__ import annotations

import html
import importlib.util
import json
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook


PROJECT = Path(r"C:\Users\Administrator\Documents\temu自动化")
ROOT = PROJECT / "outputs" / "yeahf_title_dedup_20260709" / "yeahf_400d_refill_20260718"
BOOK = ROOT / "YeahF_398D_剔除T1不可用L085两组_待J_T1回填.xlsx"
SOURCE = Path(r"E:\JIT制图--新店\L095\sku文件_最终抠图PNG\3\3.png")
OUT = ROOT / "j_l095_three_grid_repair_20260718"
MANIFEST = OUT / "l095_three_grid_j_repair_manifest.json"
REVIEW = OUT / "l095_three_grid_j_repair_review.html"
GENERATOR = PROJECT / "scripts" / "generate_yeahf_400d_j_candidates.py"


def uri(path: str | Path) -> str:
    return "file:///" + quote(str(path).replace("\\", "/"), safe="/:._-")


def load_generator():
    spec = importlib.util.spec_from_file_location("yeahf_j_generator", GENERATOR)
    if not spec or not spec.loader:
        raise RuntimeError("cannot load J generator")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main() -> None:
    if not SOURCE.is_file():
        raise RuntimeError(f"missing verified three-grid source: {SOURCE}")
    ws = load_workbook(BOOK, read_only=True, data_only=False).active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    ix = {name: headers.index(name) for name in headers}
    module = load_generator()
    records = []
    OUT.mkdir(parents=True, exist_ok=True)
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        d_value = str(values[ix["产品货号"]] or "").strip()
        sku = str(values[ix["SKU货号"]] or "").strip()
        if not d_value.startswith("L095") or sku != "L095-01":
            continue
        row = {
            "row": str(row_no), "D": d_value, "L0xx": "L095",
            "G": str(values[ix["变种属性值一"]] or "").strip(), "SKU": sku,
        }
        target = OUT / "images" / f"r{row_no:04d}_{d_value}_{sku}_three_grid.jpg"
        info = module.make_preview(SOURCE, row, target)
        records.append({
            "row": row_no, "D": d_value, "G": row["G"], "SKU": sku,
            "old_j": str(values[ix["预览图"]] or "").strip(),
            "confirmed_source": str(SOURCE), "candidate_local_path": str(target),
            "status": "candidate-awaiting-human-review-and-oss",
            "match_mode": "explicit-three-grid-source", "bytes": info["bytes"],
        })
    if len(records) != 12:
        raise RuntimeError(f"expected 12 L095 three-grid rows, found {len(records)}")
    cards = []
    for record in records:
        cards.append(f"""<article><h2>{html.escape(record['D'])} · 行 {record['row']}</h2><p>G: {html.escape(record['G'])} | SKU: {record['SKU']}</p><div class='images'><figure><img src='{html.escape(record['old_j'])}'><figcaption>旧 J：错误复用 2格源</figcaption></figure><figure><img src='{html.escape(uri(SOURCE))}'><figcaption>确认的 3格 SKU 源</figcaption></figure><figure><img src='{html.escape(uri(record['candidate_local_path']))}'><figcaption>新 J 候选：未上传、未回填</figcaption></figure></div></article>""")
    REVIEW.write_text(f"""<!doctype html><html lang='zh-CN'><meta charset='utf-8'><title>L095 三格 J 修复审查</title><style>body{{font:14px Arial,'Microsoft YaHei';margin:18px;background:#f4f6f8}}article{{background:#fff;border:1px solid #cbd5e1;margin:12px 0;padding:12px}}h1{{margin-bottom:4px}}h2{{font-size:16px}}.images{{display:grid;grid-template-columns:repeat(3,minmax(230px,1fr));gap:10px}}figure{{margin:0}}img{{width:100%;aspect-ratio:1;object-fit:contain;background:#eef2f5}}figcaption{{padding:6px 0;color:#475569}}@media(max-width:900px){{.images{{grid-template-columns:1fr}}}}</style><h1>L095 · 12 个 3格 J 修复候选</h1><p>仅本地生成。确认后才上传 OSS，并同步写入 J 与 SKC 属性预览图。</p>{''.join(cards)}</html>""", encoding="utf-8")
    MANIFEST.write_text(json.dumps({"workbook": str(BOOK), "source": str(SOURCE), "records": records, "review": str(REVIEW), "writeback": "not performed"}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"records": len(records), "review": str(REVIEW), "manifest": str(MANIFEST)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
