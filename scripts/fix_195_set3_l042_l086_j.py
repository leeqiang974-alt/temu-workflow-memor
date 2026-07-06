#!/usr/bin/env python3
"""Fix focused J preview issues in the 195-D set3 workbook.

Scope from 2026-07-05 user feedback:
- L042 J must be regenerated from first-level E:\\jit制图\\L042\\sku\\黑色/绿色
  whole size-chart images, as true five-cell grids.
- L086 beige/米杏色 variants must map to the wood/original source, not black.
- L086 has no white variant; white sources are forbidden.

The script copies the source workbook, uploads regenerated J images to OSS,
writes only the affected J cells, and emits focused audit/validation files.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import random
import re
import shutil
import sys
import time
import urllib.parse
import uuid
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

import oss2
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageOps

COMFYUI_BASE = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui")
if str(COMFYUI_BASE / "work") not in sys.path:
    sys.path.insert(0, str(COMFYUI_BASE / "work"))

import generate_apply_ali_tfirst_0608 as ali_t

try:
    from encode_workbook_image_urls import scan_url_safety, validate_t4
except Exception:
    scan_url_safety = None
    validate_t4 = None

SIZE = 800
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
L042_SKU_ROOT = Path(r"E:\jit制图\L042\sku")
L086_SKU_ROOT = Path(r"E:\JIT制图--新店\L086\sku文件_最终抠图PNG")
OSS_PREFIX = "temu-jit/dxxmall-0616-2/195-set3-jfix-l042-l086"
SIZE_RE = re.compile(r"(尺寸|尺码|size|cm|inch|length|height|width|宽|长|高)", re.I)
URL_RE = re.compile(r"https?://.*?(?=(?:[,;，；]?\s*https?://)|$)", re.S)


def headers(ws) -> list[str]:
    return [str(ws.cell(1, column).value or "").strip() for column in range(1, ws.max_column + 1)]


def col(header_values: list[str], name: str) -> int:
    if name not in header_values:
        raise RuntimeError(f"missing required column: {name}")
    return header_values.index(name) + 1


def split_urls(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    urls: list[str] = []
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line:
            continue
        found = [m.group(0).strip() for m in URL_RE.finditer(line)]
        urls.extend(found or [line])
    return urls


def is_size_url(url: str) -> bool:
    parts = urllib.parse.urlsplit(str(url or ""))
    decoded_path = urllib.parse.unquote(parts.path or str(url or ""))
    basename = decoded_path.rsplit("/", 1)[-1]
    return bool(SIZE_RE.search(basename) or "/carousel-ocr-size/" in decoded_path.lower())


def stable_int(text: str) -> int:
    return int(hashlib.sha256(text.encode("utf-8")).hexdigest()[:16], 16)


def save_jpeg_under(image: Image.Image, out_path: Path, max_bytes: int = 150 * 1024) -> dict[str, Any]:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    rgb = image.convert("RGB")
    best = b""
    quality_used = 90
    for quality in range(90, 42, -4):
        buf = BytesIO()
        rgb.save(buf, format="JPEG", quality=quality, optimize=True, progressive=True, subsampling=1)
        best = buf.getvalue()
        quality_used = quality
        if len(best) <= max_bytes:
            break
    out_path.write_bytes(best)
    return {"bytes": len(best), "quality": quality_used, "width": SIZE, "height": SIZE}


def color_scores(path: Path) -> dict[str, float]:
    image = ImageOps.exif_transpose(Image.open(path)).convert("RGB").resize((300, 300), Image.Resampling.LANCZOS)
    pixels = list(image.getdata())
    green_pixels = sum(1 for r, g, b in pixels if g > r + 18 and g > b + 12 and g > 55)
    title = image.crop((0, 0, 120, 46))
    title_pixels = list(title.getdata())
    title_green = sum(1 for r, g, b in title_pixels if g > r + 12 and g > b + 8 and g > 40)
    return {
        "overall_green": green_pixels / max(1, len(pixels)),
        "title_green": title_green / max(1, len(title_pixels)),
    }


def l042_variant(g_value: object, sku_value: object) -> tuple[str, str]:
    text = f"{g_value or ''} {sku_value or ''}".lower()
    if "绿" in text or "green" in text or "l042-02" in text:
        return "green", "绿色"
    return "black", "黑色"


def l042_first_level_sources(color_cn: str, reject_dir: Path) -> list[Path]:
    folder = L042_SKU_ROOT / color_cn
    if not folder.exists():
        raise FileNotFoundError(f"L042 source folder missing: {folder}")
    files = [p for p in sorted(folder.iterdir()) if p.is_file() and p.suffix.lower() in IMAGE_EXTS]
    preferred_name = "L042_green_final_cutout.png" if color_cn == "绿色" else "L042_black_final_cutout.png"
    preferred = [p for p in files if p.name == preferred_name]
    if preferred:
        return preferred
    filtered: list[Path] = []
    rejected: list[dict[str, Any]] = []
    for path in files:
        scores = color_scores(path)
        if color_cn == "绿色":
            keep = scores["overall_green"] >= 0.05 and scores["title_green"] >= 0.03
        else:
            keep = scores["overall_green"] < 0.05 and scores["title_green"] < 0.03
        if keep:
            filtered.append(path)
        else:
            rejected.append({"path": str(path), **{k: round(v, 4) for k, v in scores.items()}})
    reject_dir.mkdir(parents=True, exist_ok=True)
    (reject_dir / f"L042_rejected_{color_cn}_visual_mismatch.json").write_text(
        json.dumps(rejected, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not filtered:
        raise RuntimeError(f"no visually valid L042 {color_cn} first-level images in {folder}")
    return filtered


def l042_pick(files: list[Path], row: int, d_value: str, g_value: object, sku_value: object) -> Path:
    rng = random.Random(stable_int(f"{row}|{d_value}|{g_value or ''}|{sku_value or ''}"))
    return rng.choice(files)


def texture_bg(seed: int) -> Image.Image:
    rng = random.Random(seed)
    palettes = [
        ((237, 232, 220), (214, 223, 207)),
        ((235, 225, 211), (224, 233, 236)),
        ((232, 228, 218), (235, 215, 204)),
        ((224, 232, 221), (238, 228, 205)),
    ]
    top, bottom = palettes[seed % len(palettes)]
    bg = Image.new("RGB", (SIZE, SIZE), top)
    px = bg.load()
    for y in range(SIZE):
        t = y / (SIZE - 1)
        base = tuple(int(top[i] * (1 - t) + bottom[i] * t) for i in range(3))
        for x in range(SIZE):
            n = rng.randint(-4, 4)
            px[x, y] = tuple(max(0, min(255, c + n)) for c in base)
    bg = bg.filter(ImageFilter.GaussianBlur(0.35))
    draw = ImageDraw.Draw(bg, "RGBA")
    for x in range(-80, 900, 150):
        draw.line((x, 0, x + 260, 800), fill=(255, 255, 255, 24), width=3)
    return bg


def make_l042_five_grid(source: Path, out_path: Path, seed: int) -> dict[str, Any]:
    rng = random.Random(seed)
    source_img = ImageOps.exif_transpose(Image.open(source)).convert("RGB")
    source_img.thumbnail((252, 252), Image.Resampling.LANCZOS)
    bg = texture_bg(seed)
    cell = SIZE / 3
    placements = [
        (cell * 0.5, cell * 0.5),
        (cell * 2.5, cell * 0.5),
        (cell * 1.5, cell * 1.5),
        (cell * 0.5, cell * 2.5),
        (cell * 2.5, cell * 2.5),
    ]
    for cx, cy in placements:
        item = source_img.copy()
        scale = rng.uniform(0.98, 1.02)
        item = item.resize((int(item.width * scale), int(item.height * scale)), Image.Resampling.LANCZOS)
        x = round(cx - item.width / 2)
        y = round(cy - item.height / 2)
        layer = bg.convert("RGBA")
        shadow = Image.new("RGBA", item.size, (0, 0, 0, 0))
        shadow.putalpha(Image.new("L", item.size, 42).filter(ImageFilter.GaussianBlur(8)))
        layer.alpha_composite(shadow, (x + 7, y + 9))
        bg = layer.convert("RGB")
        bg.paste(item, (x, y))
    return save_jpeg_under(bg, out_path)


def make_generic_five_preview(source: Path, out_path: Path, seed: int) -> dict[str, Any]:
    rng = random.Random(seed)
    canvas = texture_bg(seed).convert("RGBA")
    slots = [(0, 0, 400, 270), (400, 0, 400, 270), (200, 265, 400, 270), (0, 530, 400, 270), (400, 530, 400, 270)]
    for index, (left, top, width, height) in enumerate(slots):
        item = ImageOps.exif_transpose(Image.open(source)).convert("RGBA")
        bbox = item.getbbox()
        if bbox:
            item = item.crop(bbox)
        scale = min((width * 0.74) / item.width, (height * 0.74) / item.height, 1.55)
        item = item.resize((max(1, int(item.width * scale)), max(1, int(item.height * scale))), Image.Resampling.LANCZOS)
        x = left + (width - item.width) // 2 + [-8, 8, 0, -6, 6][index]
        y = top + (height - item.height) // 2 + [5, -5, 0, -4, 4][index]
        alpha = item.getchannel("A")
        shadow = Image.new("RGBA", item.size, (0, 0, 0, 0))
        shadow.putalpha(alpha.filter(ImageFilter.GaussianBlur(9 + rng.randint(0, 4))))
        canvas.paste(shadow, (x + 8, y + 10), shadow)
        canvas.paste(item, (x, y), item)
    canvas = ImageEnhance.Contrast(ImageEnhance.Color(canvas.convert("RGB")).enhance(1.04)).enhance(1.06)
    return save_jpeg_under(canvas, out_path)


def l086_source(g_value: object, sku_value: object) -> tuple[Path, str, list[str], list[str]]:
    text = f"{g_value or ''} {sku_value or ''}".lower()
    wanted: list[str] = []
    if "白" in text or "white" in text:
        raise RuntimeError(f"L086 white source is forbidden: {g_value} / {sku_value}")
    if "黑" in text or "black" in text or "l086-00" in text:
        wanted.append("black")
        return L086_SKU_ROOT / "黑" / "黑.png", "l086_black", wanted, ["black"]
    if any(token in text for token in ["米", "杏", "原木", "木色", "胡桃", "beige", "wood", "l086-01"]):
        wanted.extend(["beige", "wood"])
        return L086_SKU_ROOT / "胡桃原木" / "原色.png", "l086_beige_mapped_to_wood", wanted, ["wood"]
    raise RuntimeError(f"L086 variant has no safe match: {g_value} / {sku_value}")


def load_manifest(path: Path) -> dict[str, Any]:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8-sig"))
    return {"uploaded": {}}


def save_manifest(path: Path, manifest: dict[str, Any]) -> None:
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def oss_bucket() -> oss2.Bucket:
    cfg = ali_t.ali.read_oss_config()
    return oss2.Bucket(oss2.Auth(cfg["access_key_id"], cfg["access_key_secret"]), f"https://{cfg['endpoint']}", cfg["bucket"])


def upload_once(bucket: oss2.Bucket, manifest_path: Path, manifest: dict[str, Any], row_key: str, image_path: Path) -> str:
    cache_key = f"{row_key}|{image_path.resolve()}"
    uploaded = manifest.setdefault("uploaded", {})
    if cache_key in uploaded:
        return uploaded[cache_key]
    object_key = f"{OSS_PREFIX}/{time.strftime('%Y%m%d')}/{uuid.uuid4().hex}_{row_key}_j_800.jpg"
    bucket.put_object_from_file(object_key, str(image_path), headers={"Content-Type": "image/jpeg"})
    endpoint = bucket.endpoint.replace("https://", "").replace("http://", "")
    url = f"https://{bucket.bucket_name}.{endpoint}/{object_key}"
    uploaded[cache_key] = url
    save_manifest(manifest_path, manifest)
    return url


def build_html(path: Path, records: list[dict[str, Any]], title: str) -> None:
    cards = []
    for r in records:
        local_uri = Path(r["local_image"]).as_uri() if r.get("local_image") and Path(r["local_image"]).exists() else ""
        source_uri = Path(r["sku_source"]).as_uri() if r.get("sku_source") and Path(r["sku_source"]).exists() else ""
        cards.append(
            "<section class='card'>"
            f"<h2>{html.escape(r['D'])} <span>row {r['row']}</span></h2>"
            f"<p><b>G:</b> {html.escape(str(r['G']))}<br><b>SKU:</b> {html.escape(str(r['SKU']))}<br>"
            f"<b>mode:</b> {html.escape(str(r['match_mode']))}<br><b>source:</b> {html.escape(str(r['sku_source']))}<br>"
            f"<b>wanted:</b> {html.escape(','.join(r.get('wanted_tokens') or []))}<br>"
            f"<b>matched:</b> {html.escape(','.join(r.get('matched_tokens') or []))}</p>"
            f"<div class='imgs'><figure><figcaption>新J图</figcaption><img loading='lazy' src='{html.escape(local_uri)}'></figure>"
            f"<figure><figcaption>源图</figcaption><img loading='lazy' src='{html.escape(source_uri)}'></figure></div>"
            f"<p class='url'>{html.escape(str(r.get('oss_url','')))}</p>"
            "</section>"
        )
    path.write_text(
        f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;background:#f5f6f8;color:#172033;margin:18px}}
.card{{background:#fff;border:1px solid #d8dee9;border-radius:8px;padding:12px;margin:10px 0}}
h1{{font-size:22px}} h2{{font-size:17px;margin:0 0 6px}} h2 span{{font-size:12px;color:#667085}}
p{{font-size:13px;word-break:break-all}} .imgs{{display:grid;grid-template-columns:260px 260px;gap:12px}}
figure{{margin:0}} figcaption{{font-size:12px;font-weight:700}} img{{width:260px;height:260px;object-fit:contain;background:#fff;border:1px solid #e5e7eb}}
.url{{color:#475467}}
</style></head><body><h1>{html.escape(title)} · {len(records)} 行</h1>{''.join(cards)}</body></html>""",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--output-name", required=True)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    output_path = args.out_dir / args.output_name
    shutil.copy2(args.source, output_path)

    manifest_path = args.out_dir / "j_upload_manifest.json"
    manifest = load_manifest(manifest_path)
    bucket = oss_bucket()

    wb = load_workbook(output_path)
    ws = wb.active
    h = headers(ws)
    c_d = col(h, "产品货号")
    c_g = col(h, "变种属性值一")
    c_sku = col(h, "SKU货号")
    c_j = col(h, "预览图")
    c_t = col(h, "轮播图")
    c_u = col(h, "产品素材图")

    local_dir = args.out_dir / "j_generated"
    reject_dir = args.out_dir / "j_rejects"
    l042_pools = {
        "black": l042_first_level_sources("黑色", reject_dir),
        "green": l042_first_level_sources("绿色", reject_dir),
    }

    records: list[dict[str, Any]] = []
    changed_cells: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, c_d).value or "").strip()
        if not d_value:
            continue
        prefix = d_value[:4]
        if prefix not in {"L042", "L086"}:
            continue
        g_value = str(ws.cell(row_idx, c_g).value or "").strip()
        sku_value = str(ws.cell(row_idx, c_sku).value or "").strip()
        old_j = str(ws.cell(row_idx, c_j).value or "").strip()
        if prefix == "L042":
            variant, color_cn = l042_variant(g_value, sku_value)
            source = l042_pick(l042_pools[variant], row_idx, d_value, g_value, sku_value)
            local = local_dir / "L042" / f"row{row_idx}_{d_value}_{variant}_skill_fivegrid.jpg"
            info = make_l042_five_grid(source, local, stable_int(f"L042|{row_idx}|{d_value}|{source}"))
            source_scores = {k: round(v, 4) for k, v in color_scores(source).items()}
            mode = f"l042_first_level_{color_cn}_true_five_grid"
            wanted = [variant]
            matched = [variant]
        else:
            source, mode, wanted, matched = l086_source(g_value, sku_value)
            if not source.exists():
                raise FileNotFoundError(f"L086 source missing: {source}")
            local = local_dir / "L086" / f"row{row_idx}_{d_value}_{mode}.jpg"
            info = make_generic_five_preview(source, local, stable_int(f"L086|{row_idx}|{d_value}|{source}"))
            source_scores = {}

        oss_url = upload_once(bucket, manifest_path, manifest, f"{d_value}_r{row_idx}", local)
        ws.cell(row_idx, c_j).value = oss_url
        changed_cells.append({"row": row_idx, "column": "J", "D": d_value, "old": old_j, "new": oss_url})
        records.append(
            {
                "row": row_idx,
                "D": d_value,
                "G": g_value,
                "SKU": sku_value,
                "old_j": old_j,
                "new_j": oss_url,
                "oss_url": oss_url,
                "local_image": str(local),
                "sku_source": str(source),
                "wanted_tokens": wanted,
                "matched_tokens": matched,
                "match_mode": mode,
                "image_info": info,
                "source_color_scores": source_scores,
                "warning": "",
            }
        )

    wb.save(output_path)
    wb.close()

    wb2 = load_workbook(output_path, read_only=True, data_only=False, keep_links=False)
    ws2 = wb2.active
    h2 = headers(ws2)
    d_col = col(h2, "产品货号")
    j_col = col(h2, "预览图")
    t_col = col(h2, "轮播图")
    u_col = col(h2, "产品素材图")
    effective_rows = 0
    by_d: set[str] = set()
    issues: dict[str, list[Any]] = defaultdict(list)
    for row_idx, row in enumerate(ws2.iter_rows(min_row=2, values_only=True), start=2):
        d_value = str(row[d_col - 1] or "").strip()
        if not d_value:
            continue
        effective_rows += 1
        by_d.add(d_value)
        j_value = str(row[j_col - 1] or "").strip()
        t_urls = split_urls(row[t_col - 1])
        u_value = str(row[u_col - 1] or "").strip()
        if not j_value:
            issues["empty_j"].append({"row": row_idx, "D": d_value})
        if len(t_urls) > 10:
            issues["t_over_10"].append({"row": row_idx, "D": d_value, "count": len(t_urls)})
        if len(t_urls) < 6:
            issues["t_below_6"].append({"row": row_idx, "D": d_value, "count": len(t_urls)})
        if t_urls and u_value != t_urls[0]:
            issues["u_not_t1"].append({"row": row_idx, "D": d_value})
        if len(t_urls) < 4 or not is_size_url(t_urls[3]):
            issues["t4_not_size"].append({"row": row_idx, "D": d_value})
    wb2.close()

    l042_bad_source = [r for r in records if r["D"].startswith("L042") and r"甲方不存在" in r["sku_source"]]
    l042_wrong_root = [
        r for r in records
        if r["D"].startswith("L042") and not str(r["sku_source"]).startswith(str(L042_SKU_ROOT))
    ]
    l086_beige_wrong = [
        r for r in records
        if r["D"].startswith("L086")
        and any(token in r["G"] for token in ["米", "杏"])
        and "胡桃原木" not in r["sku_source"]
    ]
    l086_white = [
        r for r in records
        if r["D"].startswith("L086") and ("白" in r["G"].lower() or "white" in r["G"].lower())
    ]
    if l042_bad_source:
        issues["l042_bad_source"].extend(l042_bad_source)
    if l042_wrong_root:
        issues["l042_wrong_root"].extend(l042_wrong_root)
    if l086_beige_wrong:
        issues["l086_beige_wrong"].extend(l086_beige_wrong)
    if l086_white:
        issues["l086_white"].extend(l086_white)

    url_safety = scan_url_safety(output_path) if scan_url_safety else {"skipped": True}
    t4_validation = validate_t4(output_path) if validate_t4 else {"skipped": True}
    issue_counts = {f"{name}_count": len(values) for name, values in sorted(issues.items())}
    issue_counts.setdefault("empty_j_count", 0)
    issue_counts.setdefault("t_over_10_count", 0)
    issue_counts.setdefault("t_below_6_count", 0)
    issue_counts.setdefault("u_not_t1_count", 0)
    issue_counts.setdefault("t4_not_size_count", 0)
    issue_counts.setdefault("l042_wrong_root_count", 0)
    issue_counts.setdefault("l086_beige_wrong_count", 0)
    issue_counts.setdefault("l086_white_count", 0)
    if isinstance(url_safety, dict) and "unsafe_url_occurrences_after" in url_safety:
        issue_counts["unsafe_url_occurrences_after"] = url_safety["unsafe_url_occurrences_after"]
    if isinstance(t4_validation, dict) and "t4_unreachable_count" in t4_validation:
        issue_counts["t4_unreachable_count"] = t4_validation["t4_unreachable_count"]

    validation = {
        "source": str(args.source),
        "output": str(output_path),
        "effective_rows": effective_rows,
        "unique_d": len(by_d),
        "changed_j_cells": len(changed_cells),
        "changed_prefix_counts": dict(Counter(r["D"][:4] for r in records)),
        "issue_counts": issue_counts,
        "pass_hard_checks": all(v == 0 for k, v in issue_counts.items() if not str(v).startswith("skipped")),
        "url_safety": url_safety,
        "t4_validation": t4_validation,
    }
    payload = {
        "validation": validation,
        "records": records,
        "changed_cells": changed_cells,
        "issues": {k: v[:30] for k, v in issues.items()},
    }
    (args.out_dir / "j_l042_l086_fix_records.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out_dir / "validation_report.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    build_html(args.out_dir / "j_l042_l086_fix_audit.html", records, "L042/L086 J 修复审核")

    print(json.dumps(validation, ensure_ascii=False, indent=2))
    return 0 if validation["pass_hard_checks"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
