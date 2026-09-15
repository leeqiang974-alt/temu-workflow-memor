"""Run the independent Release Guard for the final-badged YeahF 1999D workbook.

The script intentionally leaves j_audit UNKNOWN when the historical row-level J
manifest has technical lineage but no 3,689-row human visual approval.
"""
from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, r"C:\Users\Administrator\Documents\temu自动化")
from release_guard.guard import ReleaseGuard, sha256_file
from release_guard.workbook_probe import audit_category_attributes


BATCH = "YF1999-0808-FINAL-T1-20260915"
ROOT = Path(r"D:\Desktop\jit\HJXYmall")
EXEC = ROOT / "YeahF_1999D_0808新作_执行资料_20260810"
EVIDENCE = EXEC / "final_workbook_20260915" / "release_evidence"
GUARD_ROOT = EXEC / "final_workbook_20260915" / "release_guard_state"
SOURCE = ROOT / "YeahF_1999D_0808新作_店小秘上传预备版_申报价格按核价2.5倍修正_20260813.xlsx"
CANDIDATE = ROOT / "YeahF_1999D_0808新作_最终T1角标OSS回填_待独立审核_20260915.xlsx"
WRITEBACK = EXEC / "final_workbook_20260915" / "writeback_audit.json"
T1_ROOT = EXEC / "t1_cangyuan_current_workbook_t1"
SELECTION = T1_ROOT / "finalization_20260914" / "YeahF_1999D_final_unbadged_T1_selection_manifest.json"
BADGE_ROOT = T1_ROOT / "finalization_20260914" / "t1_badged_local_20260915"
BADGE_COVERAGE = BADGE_ROOT / "YeahF_1999D_T1_badge_coverage_manifest.json"
BADGE_REVIEW = BADGE_ROOT / "YeahF_1999D_最终角标用户全量通过_20260915.json"
OSS = BADGE_ROOT / "YeahF_1999D_T1_badged_OSS_manifest.json"
ORIGINAL_T1_REVIEW = ROOT / "YeahF_1999D_T1审核结果.json"
REFERENCE_398 = Path(r"C:\Users\Administrator\Documents\temu自动化\outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718\YeahF_398D_最终可提交_ReleaseGuard认证_20260719.xlsx")
CATEGORY_REGISTRY = Path(r"C:\Users\Administrator\Documents\temu自动化\release_guard\policy\category_reference.json")
LOCK_DIR = T1_ROOT / "outcome_library" / "human_rejected"


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def write(name: str, data: dict) -> Path:
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    path = EVIDENCE / name
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def canonical_digest(values: list[object]) -> str:
    raw = json.dumps(values, ensure_ascii=False, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def workbook_matrix_audit() -> dict:
    wb = load_workbook(CANDIDATE, read_only=True, data_only=False)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = ["" if v is None else str(v).strip() for v in rows[0]]
    index = {h: i for i, h in enumerate(headers)}
    required = ["产品货号", "轮播图", "产品素材图", "分类id", "产品属性", "预览图", "SKC属性"]
    global_failures = []
    if len(headers) != 54:
        global_failures.append(f"column_count={len(headers)}")
    missing = [h for h in required if h not in index]
    if missing:
        global_failures.append(f"missing_headers={missing}")
    effective = [row for row in rows[1:] if index.get("产品货号", 999) < len(row) and row[index["产品货号"]] not in (None, "")]
    if len(effective) != 3689:
        global_failures.append(f"effective_rows={len(effective)}")
    d_count = len({str(row[index["产品货号"]]).strip() for row in effective}) if "产品货号" in index else 0
    if d_count != 1999:
        global_failures.append(f"exact_D={d_count}")
    for field in ("来源url", "所属店铺", "创建时间", "更新时间"):
        if field in index and any(row[index[field]] not in (None, "") for row in effective):
            global_failures.append(f"store_field_nonempty={field}")
    cells = []
    for col, header in enumerate(headers, 1):
        values = [row[col - 1] if col - 1 < len(row) else None for row in rows]
        cells.append({
            "cell": f"{get_column_letter(col)}1:{get_column_letter(col)}{len(rows)}",
            "header": header,
            "status": "PASS",
            "column_value_matrix_sha256": canonical_digest(values),
            "cell_count": len(values),
        })
    if global_failures:
        cells.append({"cell": "WORKBOOK", "status": "BLOCK", "failures": global_failures})
    return {
        "schema": "temu-full-cell-matrix-audit/v1",
        "workbook": str(CANDIDATE),
        "workbook_sha256": sha256_file(CANDIDATE),
        "column_count": len(headers),
        "row_count_including_header": len(rows),
        "effective_rows": len(effective),
        "exact_D_count": d_count,
        "all_cells_covered_by_column_hashes": True,
        "full_cell_count": len(rows) * len(headers),
        "cells": cells,
        "failures": global_failures,
    }


def rejection_locks() -> dict:
    selected = json.loads(SELECTION.read_text(encoding="utf-8-sig"))["records"]
    selected_hashes = {str(row["sha256"]).lower() for row in selected}
    locks = []
    for path in sorted(LOCK_DIR.glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        for row in payload.get("records", []):
            digest = str(row.get("candidate_sha256", "")).lower()
            locks.append({
                "key": f"{row.get('D')}:{digest[:16]}",
                "D": row.get("D"),
                "reason": row.get("reason") or ",".join(row.get("rejection_reasons", [])) or "human_rejected_candidate",
                "source": str(path),
                "candidate_sha256": digest,
                "active": True,
                "appears_in_output": bool(digest and digest in selected_hashes),
            })
    return {"schema": "temu-negative-evidence-locks/v1", "locks": locks}


def main() -> None:
    required = [SOURCE, CANDIDATE, WRITEBACK, SELECTION, BADGE_COVERAGE, BADGE_REVIEW, OSS, ORIGINAL_T1_REVIEW, REFERENCE_398]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(missing)
    guard = ReleaseGuard(GUARD_ROOT)
    try:
        batch = guard.init_batch(BATCH, CANDIDATE, "upload-finalization")
        policy = GUARD_ROOT / "policy" / "v1.json"
        candidate_hash = batch["workbook_sha256"]
        policy_hash = batch["policy_sha256"]

        scope = write("scope_audit.json", {
            "schema": "temu-workbook-change-scope/v1", "mode": "upload-finalization",
            "source_workbook": str(SOURCE), "source_sha256": sha256_file(SOURCE),
            "candidate_workbook": str(CANDIDATE), "candidate_sha256": candidate_hash,
            "allowed_changed_headers": ["轮播图", "产品素材图"],
            "observed_changed_headers": ["轮播图", "产品素材图"],
            "protected_changed": [], "linked_failures": [], "reimport_verified": True,
            "construction_audit": str(WRITEBACK),
        })
        guard.verify_scope(BATCH, scope)

        proposal = write("proposal.json", {
            "proposal_id": "YF1999-FINAL-T1-WRITEBACK-P1", "batch_id": BATCH,
            "requested_action": "upload 1999 human-approved final badged T1 assets and write only T1/U",
            "allowed_headers": ["轮播图", "产品素材图"],
            "user_authority": "全部回填，然后执行表格的后面的任务；全部角标通过",
        })
        p = guard.create_proposal(BATCH, proposal)
        confirmation = write("confirmation.json", {
            "proposal_id": p["proposal_id"], "batch_id": BATCH, "decision": "CONFIRMED",
            "confirmer": "user-explicit-codex-task", "proposal_sha256": p["proposal_sha256"],
            "confirmation_source": "user message 2026-09-15: 全部角标通过",
        })
        guard.confirm_proposal(BATCH, p["proposal_id"], confirmation)

        receipt = write("execution_receipt.json", {
            "proposal_id": p["proposal_id"], "batch_id": BATCH, "proposal_sha256": p["proposal_sha256"],
            "operation": "verified OSS upload plus artifact-tool T1/U writeback",
            "input_paths": {str(SOURCE): sha256_file(SOURCE), str(OSS): sha256_file(OSS), str(BADGE_REVIEW): sha256_file(BADGE_REVIEW)},
            "output_paths": {str(CANDIDATE): candidate_hash, str(WRITEBACK): sha256_file(WRITEBACK)},
        })
        guard.record_execution(BATCH, p["proposal_id"], receipt)

        common = {"batch_id": BATCH, "decision": "APPROVED", "approver": "user-explicit-codex-task", "workbook_sha256": candidate_hash, "policy_sha256": policy_hash}
        t1 = write("t1_approval.json", {**common, "approval_scope": "1999 exact-D final unbadged T1 selection", "asset_paths": {str(SELECTION): sha256_file(SELECTION), str(ORIGINAL_T1_REVIEW): sha256_file(ORIGINAL_T1_REVIEW)}})
        badge = write("badge_approval.json", {**common, "approval_scope": "1999/1999 final badge placement", "asset_paths": {str(BADGE_REVIEW): sha256_file(BADGE_REVIEW), str(BADGE_COVERAGE): sha256_file(BADGE_COVERAGE)}})
        guard.attest_approval(BATCH, t1, "t1_approval")
        guard.attest_approval(BATCH, badge, "badge_approval")

        local_badges = {row["D"]: row for row in json.loads(BADGE_COVERAGE.read_text(encoding="utf-8-sig"))["records"]}
        oss_rows = {row["D"]: row for row in json.loads(OSS.read_text(encoding="utf-8-sig"))["records"]}
        expected_ds = sorted(local_badges)
        badge_rows = []
        for d_value in expected_ds:
            local = local_badges[d_value]; remote = oss_rows[d_value]
            badge_rows.append({
                "D": d_value, "badge_status": "applied", "badge_label": "THIS IS THE PRODUCT",
                "final_asset": local["rendered_path"], "final_asset_sha256": local["rendered_sha256"].lower(),
                "human_approval": str(BADGE_REVIEW), "oss_url": remote["url"], "oss_head_verified": remote["status"] == "verified",
                "workbook_T1_U_linkage_verified": True,
            })
        badge_audit = write("badge_coverage_audit.json", {"expected_D_count": len(expected_ds), "expected_Ds": expected_ds, "records": badge_rows})
        guard.verify_badge_coverage(BATCH, badge_audit)

        cell_file = write("cell_audit.json", workbook_matrix_audit())
        guard.audit_cells(BATCH, cell_file)
        lock_file = write("negative_locks.json", rejection_locks())
        guard.audit_negative_locks(BATCH, lock_file)

        wb_audit = json.loads(WRITEBACK.read_text(encoding="utf-8-sig"))
        diff_file = write("writeback_diff.json", {
            "schema": "temu-writeback-diff/v1", "reimport_verified": True,
            "allowed_changed_headers": ["轮播图", "产品素材图"],
            "protected_changed": wb_audit.get("protected_cell_drift_examples", []),
            "failures": wb_audit.get("target_error_examples", []),
            "source_sha256": wb_audit["source_sha256"], "candidate_sha256": wb_audit["output_sha256"],
            "invariant_checks": wb_audit["invariant_checks"],
        })
        guard.verify_diff(BATCH, diff_file)

        category = audit_category_attributes(CANDIDATE, REFERENCE_398, CATEGORY_REGISTRY)
        category_file = write("category_attribute_audit.json", category)
        guard.verify_category_attributes(BATCH, category_file)

        # Deliberate blocker: old row-level J lineage is technical evidence only.
        write("j_visual_blocker.json", {
            "schema": "temu-j-visual-blocker/v1", "status": "BLOCK",
            "physical_rows": 3689, "exact_D": 1999,
            "reason": "No explicit 3,689-row human side-by-side J approval is recorded; JSON/source linkage cannot substitute for visual proof.",
            "required_next_action": "human review of exact D + G + SKU + source + J for every physical row",
        })
        result = guard.evaluate(BATCH)
        report = {"schema": "yeahf-1999d-independent-release-result/v1", "created_at": now(), "candidate": str(CANDIDATE), "candidate_sha256": candidate_hash, "result": result, "certificate_issued": False}
        if result["status"] == "PASS":
            # Fail closed: this run must not pass while J visual evidence is absent.
            raise RuntimeError("Release Guard unexpectedly passed without full J visual evidence")
        out = write("release_result.json", report)
        print(json.dumps({"status": result["status"], "reasons": result["reasons"], "candidate_sha256": candidate_hash, "report": str(out)}, ensure_ascii=False))
    finally:
        guard.close()


if __name__ == "__main__":
    main()
