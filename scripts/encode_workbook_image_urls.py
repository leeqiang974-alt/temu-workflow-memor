#!/usr/bin/env python3
"""Percent-encode workbook image URLs so marketplace uploaders can display them.

Temu/Xuanxshop workbooks often contain OSS object keys with Chinese text,
spaces, and parentheses, especially OCR size-chart images such as
`..._尺寸_生成带纹理背景的尺寸图 (8).jpg`. Excel can store those strings, but
uploaders and browser previews may split or fail on raw spaces/non-ASCII.

This script copies a workbook, percent-encodes every URL path/query/fragment
inside string cells, and validates that T[4] remains a reachable size image.
"""

from __future__ import annotations

import argparse
import json
import re
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

URL_RE = re.compile(r"https?://[^\r\n\"'<>\\]\}]+")
SIZE_RE = re.compile(r"(尺寸|尺码|size|%E5%B0%BA%E5%AF%B8|%E5%B0%BA%E7%A0%81)", re.I)


def encode_url(url: str) -> str:
    trail = ""
    while url and url[-1] in ",;，；":
        trail = url[-1] + trail
        url = url[:-1]

    parts = urllib.parse.urlsplit(url)
    if not parts.scheme or not parts.netloc:
        return url + trail

    path = urllib.parse.quote(urllib.parse.unquote(parts.path), safe="/%")
    query = urllib.parse.quote(urllib.parse.unquote(parts.query), safe="=&%")
    fragment = urllib.parse.quote(urllib.parse.unquote(parts.fragment), safe="%")
    return urllib.parse.urlunsplit((parts.scheme, parts.netloc, path, query, fragment)) + trail


def encode_text(text: str) -> tuple[str, list[dict[str, str]]]:
    changes: list[dict[str, str]] = []

    def replace(match: re.Match[str]) -> str:
        original_match = match.group(0)
        old = original_match.strip()
        new = encode_url(old)
        if new != old:
            changes.append({"old": old, "new": new})
        return original_match.replace(old, new)

    return URL_RE.sub(replace, text), changes


def head(url: str) -> dict[str, Any]:
    try:
        request = urllib.request.Request(url, method="HEAD", headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(request, timeout=18) as response:
            return {
                "ok": 200 <= response.status < 400,
                "status": response.status,
                "content_type": response.headers.get("Content-Type", ""),
                "content_length": response.headers.get("Content-Length", ""),
            }
    except Exception as exc:  # noqa: BLE001 - report exact upload/display failure.
        return {"ok": False, "status": getattr(exc, "code", None), "error": str(exc)}


def validate_t4(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    worksheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    header_index = {header: index for index, header in enumerate(headers)}
    for required in ("产品货号", "轮播图"):
        if required not in header_index:
            raise ValueError(f"Missing required header: {required}")

    d_col = header_index["产品货号"]
    t_col = header_index["轮播图"]
    rows: list[dict[str, Any]] = []
    unique_t4: dict[str, dict[str, Any]] = {}
    for excel_row, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        d_value = row[d_col]
        if not d_value:
            continue
        t_lines = [
            line.strip()
            for line in str(row[t_col] or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
            if line.strip()
        ]
        t4 = t_lines[3] if len(t_lines) >= 4 else ""
        record = {
            "row": excel_row,
            "D": str(d_value),
            "t_count": len(t_lines),
            "t4": t4,
            "t4_size_clue": bool(SIZE_RE.search(t4)),
        }
        rows.append(record)
        if t4:
            info = unique_t4.setdefault(t4, {"rows": [], "Ds": set()})
            info["rows"].append(excel_row)
            info["Ds"].add(str(d_value))

    checks: dict[str, dict[str, Any]] = {}
    with ThreadPoolExecutor(max_workers=24) as executor:
        futures = {executor.submit(head, url): url for url in unique_t4}
        for future in as_completed(futures):
            checks[futures[future]] = future.result()

    missing: list[dict[str, Any]] = []
    not_size: list[dict[str, Any]] = []
    unreachable: list[dict[str, Any]] = []
    for record in rows:
        if record["t_count"] < 4 or not record["t4"]:
            missing.append(record)
            continue
        if not record["t4_size_clue"]:
            not_size.append(record)
        check = checks.get(record["t4"])
        if not check or not check.get("ok"):
            unreachable.append({**record, "check": check})

    return {
        "effective_rows": len(rows),
        "unique_D": len({record["D"] for record in rows}),
        "unique_t4_count": len(unique_t4),
        "missing_t4_count": len(missing),
        "t4_not_size_clue_count": len(not_size),
        "t4_unreachable_count": len(unreachable),
        "all_t4_display_ok": not missing and not not_size and not unreachable,
        "missing_t4_examples": missing[:20],
        "t4_not_size_examples": not_size[:20],
        "t4_unreachable_examples": unreachable[:20],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Copy a workbook and percent-encode image URLs.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.input)
    changed_cells: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "http" in cell.value:
                    new_value, changes = encode_text(cell.value)
                    if changes and new_value != cell.value:
                        changed_cells.append(
                            {
                                "sheet": worksheet.title,
                                "cell": cell.coordinate,
                                "change_count": len(changes),
                                "sample": changes[:3],
                            }
                        )
                        cell.value = new_value

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)

    validation = validate_t4(args.output)
    report = {
        "source": str(args.input),
        "output": str(args.output),
        "changed_cell_count": len(changed_cells),
        "changed_url_occurrence_count": sum(item["change_count"] for item in changed_cells),
        "changed_cells_sample": changed_cells[:30],
        **validation,
    }
    report_path = args.report or args.output.with_suffix(".image_url_encode_report.json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if validation["all_t4_display_ok"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
