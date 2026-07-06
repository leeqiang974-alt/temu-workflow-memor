#!/usr/bin/env python3
"""Write one approved 195x3 T-first set into a copied workbook.

Scope for 2026-07-06:
- source workbook is not modified in place
- use one selected set from the reviewed 195x3 passed pool
- replace T1 only when the current D already has at least 6 T URLs
- preserve the size image at T4 and cap T at 10 URLs
- set U = T1 for changed D rows
- fix L085 weight to 150
- rewrite title tracking codes for this workbook
- rebuild all J preview images, and sync SKC属性.previewImgUrls
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import random
import re
import shutil
import sys
import time
import uuid
import urllib.parse
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

import oss2
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageOps


COMFYUI_BASE = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui")
if str(COMFYUI_BASE / "work") not in sys.path:
    sys.path.insert(0, str(COMFYUI_BASE / "work"))

import generate_apply_ali_tfirst_0608 as ali_t

try:
    from encode_workbook_image_urls import encode_text, scan_url_safety, validate_t4
except Exception:  # pragma: no cover
    encode_text = None
    scan_url_safety = None
    validate_t4 = None


TARGET_SIZE = 800
TARGET_MAX_BYTES = 150 * 1024
OSS_PREFIX = "temu-jit/dxxmall-0616-2/195x3-passed-t-writeback"
J_OSS_PREFIX = "temu-jit/dxxmall-0616-2/192-set1-j-row-sku-variant"
URL_RE = re.compile(r"https?://.*?(?=(?:[,;，；]?\s*https?://)|$)", re.S)
SIZE_RE = re.compile(r"(尺寸|尺码|size|cm|inch|length|height|width|宽|长|高)", re.I)
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
FORBIDDEN_SOURCE_WORDS = ("九宫格", "9grid", "out", "output", "背景素材")
L042_SKU_ROOT = Path(r"E:\jit制图\L042\sku")
L086_SKU_ROOT = Path(r"E:\JIT制图--新店\L086\sku文件_最终抠图PNG")
FORBIDDEN_TITLE_WORDS = [
    "自动",
    "智能",
    "AI",
    "APP",
    "蓝牙",
    "wifi",
    "遥控",
    "感应",
    "电动",
    "充电",
    "电池",
    "USB",
    "儿童",
    "玩具",
    "环保",
    "有机",
    "可降解",
    "食品级",
    "最佳",
    "最好",
    "顶级",
    "久用不塌陷",
    "长久不变形",
    "亲肤",
    "高弹填充",
    "四面延伸",
]

BROKEN_L095_T4 = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "85ef73d852584a65aaf05a1c5615dac7_L095_%E5%B0%BA%E5%AF%B8_jimeng-2026-05-08-1613-"
    "%E8%BF%99%E5%BC%A0%E5%9B%BE%E7%89%87"
)

FULL_L095_T4 = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "85ef73d852584a65aaf05a1c5615dac7_L095_%E5%B0%BA%E5%AF%B8_jimeng-2026-05-08-1613-"
    "%E8%BF%99%E5%BC%A0%E5%9B%BE%E7%89%87%EF%BC%8C%E4%BF%9D%E6%8C%81%E4%BA%A7%E5%93%81"
    "%E5%92%8C%E6%A0%87%E5%B0%BA%E3%80%81%E6%95%B0%E5%AD%97%E3%80%81%E6%96%87%E5%AD%97"
    "%E4%BF%A1%E6%81%AF%E4%B8%8D%E5%81%9A.jpg"
)


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def save_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def headers(ws) -> list[str]:
    return [str(ws.cell(1, column).value or "").strip() for column in range(1, ws.max_column + 1)]


def col(header_values: list[str], name: str) -> int:
    if name not in header_values:
        raise RuntimeError(f"missing required column: {name}")
    return header_values.index(name) + 1


def split_urls(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    urls: list[str] = []
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line:
            continue
        matches = [match.group(0).strip() for match in URL_RE.finditer(line)]
        urls.extend(matches or [line])
    return urls


def dedupe(values: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            result.append(value)
            seen.add(value)
    return result


def is_size_url(url: str) -> bool:
    text = str(url or "")
    if not text:
        return False
    parts = urllib.parse.urlsplit(text)
    decoded_path = urllib.parse.unquote(parts.path or text)
    basename = decoded_path.rsplit("/", 1)[-1]
    return bool(SIZE_RE.search(basename) or "/carousel-ocr-size/" in decoded_path.lower())


def find_size_url(urls: list[str]) -> str:
    for url in urls:
        if is_size_url(url):
            return url
    return urls[3] if len(urls) >= 4 else ""


def build_t_urls(old_urls: list[str], new_t1: str) -> tuple[list[str], str, str]:
    size_url = find_size_url(old_urls)
    if len(old_urls) < 6:
        return dedupe(old_urls + [new_t1])[:10], "append_t1_because_old_t_below_6", size_url
    rest = [url for idx, url in enumerate(old_urls) if idx != 0 and url not in {new_t1, size_url}]
    if size_url:
        urls = [new_t1] + rest[:2] + [size_url] + rest[2:]
    else:
        urls = [new_t1] + rest
    return dedupe(urls)[:10], "replace_t1", size_url


def load_pool(path: Path, set_no: int) -> dict[str, dict[str, Any]]:
    records = load_json(path)
    selected: dict[str, dict[str, Any]] = {}
    for record in records:
        d_value = str(record.get("original_d") or record.get("d") or "").strip()
        if not d_value or int(record.get("set_no") or 0) != set_no:
            continue
        local_path = Path(str(record.get("local_path") or ""))
        if not local_path.exists():
            raise FileNotFoundError(f"{d_value}: missing set{set_no} image {local_path}")
        selected[d_value] = record
    return selected


def resize_compress(input_path: Path, output_path: Path) -> dict[str, Any]:
    image = Image.open(input_path).convert("RGB").resize((TARGET_SIZE, TARGET_SIZE), Image.Resampling.LANCZOS)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    from io import BytesIO

    best = b""
    best_quality = 88
    for quality in range(88, 34, -3):
        buffer = BytesIO()
        image.save(buffer, format="JPEG", quality=quality, optimize=True, progressive=True, subsampling=1)
        best = buffer.getvalue()
        best_quality = quality
        if len(best) <= TARGET_MAX_BYTES:
            break
    output_path.write_bytes(best)
    return {"width": TARGET_SIZE, "height": TARGET_SIZE, "bytes": len(best), "quality": best_quality}


def load_manifest(path: Path) -> dict[str, Any]:
    if path.exists():
        return load_json(path)
    return {"uploaded": {}}


def upload_once(bucket: oss2.Bucket, manifest_path: Path, manifest: dict[str, Any], d_value: str, set_no: int, path: Path) -> str:
    key = f"{d_value}|set{set_no}|{path.resolve()}"
    uploaded = manifest.setdefault("uploaded", {})
    if key in uploaded:
        return uploaded[key]
    object_key = f"{OSS_PREFIX}/{time.strftime('%Y%m%d')}/{uuid.uuid4().hex}_{d_value}_set{set_no}_tfirst_800.jpg"
    bucket.put_object_from_file(object_key, str(path), headers={"Content-Type": "image/jpeg"})
    endpoint = bucket.endpoint.replace("https://", "").replace("http://", "")
    url = f"https://{bucket.bucket_name}.{endpoint}/{object_key}"
    uploaded[key] = url
    save_json(manifest_path, manifest)
    return url


def token_set(text: object) -> set[str]:
    raw = str(text or "").lower()
    mapping = [
        ("black", ["黑", "black"]),
        ("green", ["绿", "green"]),
        ("white", ["白", "透明", "white"]),
        ("gray", ["灰", "grey", "gray"]),
        ("yellow", ["黄", "yellow"]),
        ("pink", ["粉", "pink"]),
        ("red", ["红", "red"]),
        ("blue", ["蓝", "blue"]),
        ("purple", ["紫", "purple"]),
        ("wood", ["木色", "原木", "胡桃", "walnut", "wood"]),
        ("beige", ["米", "杏", "beige"]),
        ("15", ["15"]),
        ("30", ["30"]),
        ("50", ["50"]),
        ("2", ["2", "双", "两", "二"]),
        ("3", ["3", "三"]),
    ]
    found: set[str] = set()
    for canonical, aliases in mapping:
        if any(alias in raw for alias in aliases):
            found.add(canonical)
    return found


def clean_source_files(folder: Path) -> list[Path]:
    files: list[Path] = []
    if not folder.exists():
        return files
    for path in folder.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTS:
            continue
        lower = str(path).lower()
        if any(word.lower() in lower for word in FORBIDDEN_SOURCE_WORDS):
            continue
        files.append(path)
    return sorted(files, key=lambda p: str(p))


def choose_j_source(prefix: str, variant: str, sku: str, roots: list[Path]) -> dict[str, Any]:
    wanted = token_set(f"{variant} {sku}")
    if prefix == "L043":
        base = roots[0] / "L043" / "sku文件_最终抠图PNG"
        color = "灰" if "gray" in wanted else "白"
        source = base / color / f"{color}.png"
        return {
            "sku_root": str(base),
            "sku_source": str(source),
            "wanted_tokens": sorted(wanted),
            "matched_tokens": sorted(wanted & token_set(str(source))),
            "match_mode": "l043_historical_white_gray_hard_rule",
            "warning": "" if source.exists() else "missing_l043_source",
        }

    folder_names = ["sku文件_最终抠图PNG", "sku文件", "sku"]
    chosen_folder = None
    candidates: list[Path] = []
    for root in roots:
        for folder_name in folder_names:
            folder = root / prefix / folder_name
            files = clean_source_files(folder)
            if files:
                chosen_folder = folder
                candidates = files
                break
        if candidates:
            break
    if not candidates or chosen_folder is None:
        return {
            "sku_root": "",
            "sku_source": "",
            "wanted_tokens": sorted(wanted),
            "matched_tokens": [],
            "match_mode": "missing_source",
            "warning": "missing_source_folder",
        }

    def score(path: Path) -> tuple[int, str]:
        got = token_set(str(path))
        score_value = len(wanted & got) * 10
        for token in {"black", "green", "white", "gray", "pink", "red", "blue", "purple", "yellow", "wood", "beige"}:
            if token in wanted and token in got:
                score_value += 40
        for token in {"15", "30", "50", "2", "3"}:
            if token in wanted and token in got:
                score_value += 8
        if wanted & {"black", "green", "white", "gray", "pink", "red", "blue", "purple", "yellow", "wood", "beige"} and not (
            wanted & got & {"black", "green", "white", "gray", "pink", "red", "blue", "purple", "yellow", "wood", "beige"}
        ):
            score_value -= 25
        return score_value, str(path)

    top = sorted(candidates, key=score, reverse=True)[0]
    top_score = score(top)[0]
    got = token_set(str(top))
    return {
        "sku_root": str(chosen_folder),
        "sku_source": str(top),
        "wanted_tokens": sorted(wanted),
        "matched_tokens": sorted(wanted & got),
        "match_score": top_score,
        "match_mode": "token_match" if top_score > 0 else "prefix_fallback",
        "warning": "" if top_score > 0 else "no_positive_token_match",
    }


def stable_int(text: str) -> int:
    return int(hashlib.sha256(text.encode("utf-8")).hexdigest()[:16], 16)


def stable_tracking_code(d_value: str, used: set[str], salt: str) -> str:
    letters = "ABCDEFGHJKLMNPQRSTUVWXYZ"
    digits = "23456789"
    nonce = 0
    while True:
        digest = hashlib.sha256(f"{salt}|{d_value}|{nonce}".encode("utf-8")).digest()
        code = letters[digest[0] % len(letters)] + digits[digest[1] % len(digits)] + letters[digest[2] % len(letters)]
        if code not in used:
            used.add(code)
            return code
        nonce += 1


def apply_tracking_code(title: object, code: str) -> str:
    text = re.sub(r"\s+[A-Z][0-9][A-Z]\s*$", "", str(title or "").strip())
    if len(text) + 4 > 80:
        text = text[:76].rstrip("，、 -")
    return f"{text} {code}".strip()


def color_scores(path: Path) -> dict[str, float]:
    image = ImageOps.exif_transpose(Image.open(path)).convert("RGB").resize((300, 300), Image.Resampling.LANCZOS)
    pixels = list(image.getdata())
    green_pixels = sum(1 for r, g, b in pixels if g > r + 18 and g > b + 12 and g > 55)
    title = image.crop((0, 0, 120, 46))
    title_pixels = list(title.getdata())
    title_green = sum(1 for r, g, b in title_pixels if g > r + 12 and g > b + 8 and g > 40)
    return {
        "overall_green": green_pixels / max(1, len(pixels)),
        "title_green": title_green / max(1, len(title_pixels)),
    }


def l042_variant(g_value: object, sku_value: object) -> tuple[str, str]:
    text = f"{g_value or ''} {sku_value or ''}".lower()
    if "绿" in text or "green" in text or "l042-02" in text:
        return "green", "绿色"
    return "black", "黑色"


def l042_first_level_sources(color_cn: str, reject_dir: Path) -> list[Path]:
    folder = L042_SKU_ROOT / color_cn
    if not folder.exists():
        raise FileNotFoundError(f"L042 source folder missing: {folder}")
    files = [p for p in sorted(folder.iterdir()) if p.is_file() and p.suffix.lower() in IMAGE_EXTS]
    filtered: list[Path] = []
    rejected: list[dict[str, Any]] = []
    for path in files:
        scores = color_scores(path)
        if color_cn == "绿色":
            keep = scores["overall_green"] >= 0.05 and scores["title_green"] >= 0.03
        else:
            keep = scores["overall_green"] < 0.05 and scores["title_green"] < 0.03
        if keep:
            filtered.append(path)
        else:
            rejected.append({"path": str(path), **{k: round(v, 4) for k, v in scores.items()}})
    reject_dir.mkdir(parents=True, exist_ok=True)
    (reject_dir / f"L042_rejected_{color_cn}_visual_mismatch.json").write_text(
        json.dumps(rejected, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not filtered:
        raise RuntimeError(f"no visually valid L042 {color_cn} first-level images in {folder}")
    return filtered


def l086_source(g_value: object, sku_value: object) -> tuple[Path, str, list[str], list[str]]:
    text = f"{g_value or ''} {sku_value or ''}".lower()
    if "白" in text or "white" in text:
        raise RuntimeError(f"L086 white source is forbidden: {g_value} / {sku_value}")
    if "黑" in text or "black" in text or "l086-00" in text:
        return L086_SKU_ROOT / "黑" / "黑.png", "l086_black", ["black"], ["black"]
    if any(token in text for token in ["米", "杏", "原木", "木色", "胡桃", "beige", "wood", "l086-01"]):
        return L086_SKU_ROOT / "胡桃原木" / "原色.png", "l086_beige_mapped_to_wood", ["beige", "wood"], ["wood"]
    raise RuntimeError(f"L086 variant has no safe match: {g_value} / {sku_value}")


def texture_bg(seed: int) -> Image.Image:
    rng = random.Random(seed)
    palettes = [
        ((237, 232, 220), (214, 223, 207)),
        ((235, 225, 211), (224, 233, 236)),
        ((232, 228, 218), (235, 215, 204)),
        ((224, 232, 221), (238, 228, 205)),
        ((230, 224, 237), (216, 225, 232)),
    ]
    top, bottom = palettes[seed % len(palettes)]
    bg = Image.new("RGB", (TARGET_SIZE, TARGET_SIZE), top)
    px = bg.load()
    for y in range(TARGET_SIZE):
        t = y / (TARGET_SIZE - 1)
        base = tuple(int(top[i] * (1 - t) + bottom[i] * t) for i in range(3))
        for x in range(TARGET_SIZE):
            n = rng.randint(-4, 4)
            px[x, y] = tuple(max(0, min(255, c + n)) for c in base)
    bg = bg.filter(ImageFilter.GaussianBlur(0.35))
    draw = ImageDraw.Draw(bg, "RGBA")
    for x in range(-80, 900, 150):
        draw.line((x, 0, x + 260, TARGET_SIZE), fill=(255, 255, 255, 24), width=3)
    return bg


def save_jpeg_under(image: Image.Image, out_path: Path, max_bytes: int = TARGET_MAX_BYTES) -> dict[str, Any]:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    rgb = image.convert("RGB")
    best = b""
    quality_used = 90
    for quality in range(90, 42, -4):
        buf = BytesIO()
        rgb.save(buf, format="JPEG", quality=quality, optimize=True, progressive=True, subsampling=1)
        best = buf.getvalue()
        quality_used = quality
        if len(best) <= max_bytes:
            break
    out_path.write_bytes(best)
    return {"bytes": len(best), "quality": quality_used, "width": TARGET_SIZE, "height": TARGET_SIZE}


def make_l042_five_grid(source: Path, out_path: Path, seed: int) -> dict[str, Any]:
    rng = random.Random(seed)
    source_img = ImageOps.exif_transpose(Image.open(source)).convert("RGB")
    source_img.thumbnail((252, 252), Image.Resampling.LANCZOS)
    bg = texture_bg(seed)
    cell = TARGET_SIZE / 3
    placements = [
        (cell * 0.5, cell * 0.5),
        (cell * 2.5, cell * 0.5),
        (cell * 1.5, cell * 1.5),
        (cell * 0.5, cell * 2.5),
        (cell * 2.5, cell * 2.5),
    ]
    for cx, cy in placements:
        item = source_img.copy()
        scale = rng.uniform(0.98, 1.02)
        item = item.resize((int(item.width * scale), int(item.height * scale)), Image.Resampling.LANCZOS)
        x = round(cx - item.width / 2)
        y = round(cy - item.height / 2)
        layer = bg.convert("RGBA")
        shadow = Image.new("RGBA", item.size, (0, 0, 0, 0))
        shadow.putalpha(Image.new("L", item.size, 42).filter(ImageFilter.GaussianBlur(8)))
        layer.alpha_composite(shadow, (x + 7, y + 9))
        bg = layer.convert("RGB")
        bg.paste(item, (x, y))
    return save_jpeg_under(bg, out_path)


def make_generic_five_preview(source: Path, out_path: Path, seed: int) -> dict[str, Any]:
    rng = random.Random(seed)
    canvas = texture_bg(seed).convert("RGBA")
    slots = [(0, 0, 400, 270), (400, 0, 400, 270), (200, 265, 400, 270), (0, 530, 400, 270), (400, 530, 400, 270)]
    for index, (left, top, width, height) in enumerate(slots):
        item = ImageOps.exif_transpose(Image.open(source)).convert("RGBA")
        bbox = item.getbbox()
        if bbox:
            item = item.crop(bbox)
        scale = min((width * 0.74) / item.width, (height * 0.74) / item.height, 1.55)
        item = item.resize((max(1, int(item.width * scale)), max(1, int(item.height * scale))), Image.Resampling.LANCZOS)
        x = left + (width - item.width) // 2 + [-8, 8, 0, -6, 6][index]
        y = top + (height - item.height) // 2 + [5, -5, 0, -4, 4][index]
        alpha = item.getchannel("A")
        shadow = Image.new("RGBA", item.size, (0, 0, 0, 0))
        shadow.putalpha(alpha.filter(ImageFilter.GaussianBlur(9 + rng.randint(0, 4))))
        canvas.paste(shadow, (x + 8, y + 10), shadow)
        canvas.paste(item, (x, y), item)
    canvas = ImageEnhance.Contrast(ImageEnhance.Color(canvas.convert("RGB")).enhance(1.04)).enhance(1.06)
    return save_jpeg_under(canvas, out_path)


def upload_j_once(bucket: oss2.Bucket, manifest_path: Path, manifest: dict[str, Any], row_key: str, image_path: Path) -> str:
    cache_key = f"{row_key}|{image_path.resolve()}"
    uploaded = manifest.setdefault("uploaded", {})
    if cache_key in uploaded:
        return uploaded[cache_key]
    object_key = f"{J_OSS_PREFIX}/{time.strftime('%Y%m%d')}/{uuid.uuid4().hex}_{row_key}_j_800.jpg"
    bucket.put_object_from_file(object_key, str(image_path), headers={"Content-Type": "image/jpeg"})
    endpoint = bucket.endpoint.replace("https://", "").replace("http://", "")
    url = f"https://{bucket.bucket_name}.{endpoint}/{object_key}"
    uploaded[cache_key] = url
    save_json(manifest_path, manifest)
    return url


def update_skc_preview(value: object, preview_url: str) -> tuple[str, bool]:
    text = str(value or "").strip()
    if not text:
        return text, False
    try:
        data = json.loads(text)
    except Exception:
        return text, False
    changed = False
    if isinstance(data, list):
        for item in data:
            if isinstance(item, dict) and item.get("previewImgUrls") != preview_url:
                item["previewImgUrls"] = preview_url
                changed = True
    elif isinstance(data, dict) and data.get("previewImgUrls") != preview_url:
        data["previewImgUrls"] = preview_url
        changed = True
    return json.dumps(data, ensure_ascii=False, separators=(",", ":")), changed


def file_url(path: str) -> str:
    try:
        return Path(path).resolve().as_uri()
    except Exception:
        return ""


def build_t_audit_html(path: Path, report: dict[str, Any]) -> None:
    sections = []
    for item in report["changes"]:
        imgs = "".join(
            f"<figure><figcaption>T{i + 1}{' 尺寸图' if i == 3 else ''}</figcaption><img loading='lazy' src='{html.escape(url)}'><div>{html.escape(url)}</div></figure>"
            for i, url in enumerate(item["final_urls"][:10])
        )
        sections.append(
            f"<section><h2>{html.escape(item['D'])}</h2><p>rows: {item['rows']} | mode: {html.escape(item['mode'])} | old_count: {item['old_count']} | final_count: {item['final_count']} | T4 size: {item['t4_is_size']}</p><p class='small'>new T1: {html.escape(item['new_t1'])}</p><div class='grid'>{imgs}</div></section>"
        )
    path.write_text(
        f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>195x3 set{report['set_no']} T writeback audit</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;background:#f5f6f8;color:#172033;margin:18px}}
.summary,section{{background:#fff;border:1px solid #d8dee9;border-radius:8px;padding:12px;margin:10px 0}}
h1{{font-size:22px}} h2{{font-size:17px}} pre{{white-space:pre-wrap;background:#f3f4f6;padding:10px;border-radius:6px}}
.small{{font-size:12px;word-break:break-all;color:#667085}} .grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:10px}} figure{{margin:0}}
figcaption{{font-size:12px;font-weight:700}} img{{width:100%;aspect-ratio:1/1;object-fit:contain;background:#fff;border:1px solid #e5e7eb}}
figure div{{font-size:10px;color:#667085;word-break:break-all;max-height:48px;overflow:auto}}
</style></head><body><div class='summary'><h1>195x3 set{report['set_no']} T 首图回填审核</h1><p>{html.escape(report['output_workbook'])}</p><pre>{html.escape(json.dumps(report['validation'], ensure_ascii=False, indent=2))}</pre></div>{''.join(sections)}</body></html>""",
        encoding="utf-8",
    )


def build_j_audit_html(path: Path, records: list[dict[str, Any]]) -> None:
    cards = []
    for record in records:
        warning = record.get("warning") or ""
        cards.append(
            "<section class='card'>"
            f"<h3>{html.escape(str(record['D']))} <span>row {record['row']}</span></h3>"
            f"<p><b>G:</b> {html.escape(str(record.get('G','')))}<br><b>SKU:</b> {html.escape(str(record.get('SKU','')))}</p>"
            f"<p><b>source:</b> {html.escape(str(record.get('sku_source','')))}<br><b>mode:</b> {html.escape(str(record.get('match_mode','')))}<br><b>wanted:</b> {html.escape(','.join(record.get('wanted_tokens') or []))}<br><b>matched:</b> {html.escape(','.join(record.get('matched_tokens') or []))}</p>"
            f"<p class='warn'>{html.escape(warning)}</p>"
            f"<div class='imgs'><figure><figcaption>当前J</figcaption><img loading='lazy' src='{html.escape(str(record.get('current_j','')))}'></figure><figure><figcaption>匹配源图</figcaption><img loading='lazy' src='{html.escape(file_url(str(record.get('sku_source',''))))}'></figure></div>"
            "</section>"
        )
    path.write_text(
        f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>J match audit</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;background:#f5f6f8;color:#172033;margin:18px}}
.card{{background:#fff;border:1px solid #d8dee9;border-radius:8px;padding:12px;margin:10px 0}}
h1{{font-size:22px}} h3{{font-size:17px;margin:0 0 6px}} h3 span{{font-size:12px;color:#667085}}
p{{font-size:13px;word-break:break-all}} .warn{{color:#b42318;font-weight:700}}
.imgs{{display:grid;grid-template-columns:180px 180px;gap:12px}} figure{{margin:0}} figcaption{{font-size:12px;font-weight:700}}
img{{width:180px;height:180px;object-fit:contain;background:#fff;border:1px solid #e5e7eb}}
</style></head><body><h1>J 图 SKU 属性匹配审核 · {len(records)} 行</h1>{''.join(cards)}</body></html>""",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--pool", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--output-name", required=True)
    parser.add_argument("--set-no", type=int, choices=(1, 2, 3), default=1)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    output_path = args.out_dir / args.output_name
    shutil.copy2(args.source, output_path)

    selected_pool = load_pool(args.pool, args.set_no)
    compressed_dir = args.out_dir / "t_first_800"
    manifest_path = args.out_dir / "t_upload_manifest.json"
    oss_config = ali_t.ali.read_oss_config()
    bucket = oss2.Bucket(
        oss2.Auth(oss_config["access_key_id"], oss_config["access_key_secret"]),
        f"https://{oss_config['endpoint']}",
        oss_config["bucket"],
    )
    manifest = load_manifest(manifest_path)

    wb = load_workbook(output_path)
    ws = wb.active
    h = headers(ws)
    col_title = col(h, "产品标题")
    col_d = col(h, "产品货号")
    col_g = col(h, "变种属性值一")
    col_sku = col(h, "SKU货号")
    col_j = col(h, "预览图")
    col_t = col(h, "轮播图")
    col_u = col(h, "产品素材图")
    col_weight = col(h, "重量")
    col_skc = col(h, "SKC属性")
    j_manifest_path = args.out_dir / "j_upload_manifest.json"
    j_manifest = load_manifest(j_manifest_path)
    j_generated_dir = args.out_dir / "j_generated"
    j_reject_dir = args.out_dir / "j_rejects"
    l042_pools = {
        "black": l042_first_level_sources("黑色", j_reject_dir),
        "green": l042_first_level_sources("绿色", j_reject_dir),
    }
    title_salt = f"0616-2-192-set{args.set_no}-title-fingerprint-20260706"

    deleted_rows: list[int] = []
    for row_idx in range(ws.max_row, 1, -1):
        if not str(ws.cell(row_idx, col_d).value or "").strip():
            deleted_rows.append(row_idx)
            ws.delete_rows(row_idx, 1)

    d_rows: dict[str, list[int]] = defaultdict(list)
    original_first_by_d: dict[str, dict[str, Any]] = {}
    snapshots: dict[str, dict[str, Any]] = {}
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        if not d_value:
            continue
        d_rows[d_value].append(row_idx)
        original_first_by_d.setdefault(
            d_value,
            {
                "title": str(ws.cell(row_idx, col_title).value or ""),
                "t_urls": split_urls(ws.cell(row_idx, col_t).value),
                "u": str(ws.cell(row_idx, col_u).value or "").strip(),
            },
        )
        snapshots[str(row_idx)] = {
            "D": d_value,
            "title": ws.cell(row_idx, col_title).value,
            "G": ws.cell(row_idx, col_g).value,
            "SKU": ws.cell(row_idx, col_sku).value,
            "J": ws.cell(row_idx, col_j).value,
            "T": ws.cell(row_idx, col_t).value,
            "U": ws.cell(row_idx, col_u).value,
            "weight": ws.cell(row_idx, col_weight).value,
            "SKC": ws.cell(row_idx, col_skc).value,
        }

    missing_pool = sorted(set(d_rows) - set(selected_pool))
    if missing_pool:
        raise RuntimeError(f"workbook D missing from 195x3 passed pool: {missing_pool[:30]}")

    title_codes: dict[str, str] = {}
    used_title_codes: set[str] = set()
    title_changes: list[dict[str, Any]] = []
    for d_value in sorted(d_rows):
        title_codes[d_value] = stable_tracking_code(d_value, used_title_codes, title_salt)
        for row_idx in d_rows[d_value]:
            old_title = ws.cell(row_idx, col_title).value
            new_title = apply_tracking_code(old_title, title_codes[d_value])
            if old_title != new_title:
                ws.cell(row_idx, col_title).value = new_title
                title_changes.append({"row": row_idx, "D": d_value, "old": old_title, "new": new_title, "code": title_codes[d_value]})

    j_roots = [Path(r"E:\JIT制图--新店"), Path(r"E:\jit制图")]
    j_write_records: list[dict[str, Any]] = []
    j_cell_changes: list[dict[str, Any]] = []
    skc_preview_changes: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        if not d_value:
            continue
        prefix = d_value[:4]
        g_value = str(ws.cell(row_idx, col_g).value or "").strip()
        sku_value = str(ws.cell(row_idx, col_sku).value or "").strip()
        old_j = str(ws.cell(row_idx, col_j).value or "").strip()
        if prefix == "L042":
            variant, color_cn = l042_variant(g_value, sku_value)
            files = l042_pools[variant]
            source = random.Random(stable_int(f"L042|{row_idx}|{d_value}|{g_value}|{sku_value}")).choice(files)
            local = j_generated_dir / "L042" / f"row{row_idx}_{d_value}_{variant}_fivegrid.jpg"
            image_info = make_l042_five_grid(source, local, stable_int(f"L042|{row_idx}|{d_value}|{source}"))
            mode = f"l042_first_level_{color_cn}_true_five_grid"
            wanted = [variant]
            matched = [variant]
            warning = ""
        elif prefix == "L086":
            source, mode, wanted, matched = l086_source(g_value, sku_value)
            if not source.exists():
                raise FileNotFoundError(f"L086 source missing: {source}")
            local = j_generated_dir / "L086" / f"row{row_idx}_{d_value}_{mode}.jpg"
            image_info = make_generic_five_preview(source, local, stable_int(f"L086|{row_idx}|{d_value}|{source}"))
            warning = ""
        else:
            match = choose_j_source(prefix, g_value, sku_value, j_roots)
            if not match.get("sku_source"):
                raise RuntimeError(f"missing J source for row {row_idx} {d_value}: {match}")
            source = Path(str(match["sku_source"]))
            if not source.exists():
                raise FileNotFoundError(f"J source missing for row {row_idx} {d_value}: {source}")
            local = j_generated_dir / prefix / f"row{row_idx}_{d_value}_{stable_int(str(source)) & 0xffff:x}_fivegrid.jpg"
            image_info = make_generic_five_preview(source, local, stable_int(f"{prefix}|{row_idx}|{d_value}|{source}"))
            mode = str(match.get("match_mode") or "")
            wanted = list(match.get("wanted_tokens") or [])
            matched = list(match.get("matched_tokens") or [])
            warning = str(match.get("warning") or "")
        j_url = upload_j_once(bucket, j_manifest_path, j_manifest, f"{d_value}_r{row_idx}", local)
        ws.cell(row_idx, col_j).value = j_url
        j_cell_changes.append({"row": row_idx, "D": d_value, "column": "预览图", "old": old_j, "new": j_url})
        old_skc = ws.cell(row_idx, col_skc).value
        new_skc, skc_changed = update_skc_preview(old_skc, j_url)
        if skc_changed:
            ws.cell(row_idx, col_skc).value = new_skc
            skc_preview_changes.append({"row": row_idx, "D": d_value, "column": "SKC属性", "old": old_skc, "new": new_skc})
        j_write_records.append(
            {
                "row": row_idx,
                "D": d_value,
                "G": g_value,
                "SKU": sku_value,
                "old_j": old_j,
                "current_j": j_url,
                "new_j": j_url,
                "oss_url": j_url,
                "local_image": str(local),
                "sku_root": str(source.parent),
                "sku_source": str(source),
                "wanted_tokens": wanted,
                "matched_tokens": matched,
                "match_mode": mode,
                "warning": warning,
                "image_info": image_info,
                "skc_preview_synced": skc_changed or (str(old_skc or "").find(j_url) >= 0),
            }
        )

    upload_records: dict[str, dict[str, Any]] = {}
    changes: list[dict[str, Any]] = []
    for d_value in sorted(d_rows):
        record = selected_pool[d_value]
        local_path = Path(str(record["local_path"]))
        compressed = compressed_dir / d_value[:4] / f"{d_value}_set{args.set_no}_tfirst_800.jpg"
        compression = resize_compress(local_path, compressed)
        new_t1 = upload_once(bucket, manifest_path, manifest, d_value, args.set_no, compressed)
        upload_records[d_value] = {
            "source_record": record,
            "compressed": str(compressed),
            "compression": compression,
            "uploaded_url": new_t1,
        }
        old_urls = original_first_by_d[d_value]["t_urls"]
        final_urls, mode, size_url = build_t_urls(old_urls, new_t1)
        for row_idx in d_rows[d_value]:
            ws.cell(row_idx, col_t).value = "\n".join(final_urls)
            ws.cell(row_idx, col_u).value = final_urls[0] if final_urls else ""
        changes.append(
            {
                "D": d_value,
                "rows": d_rows[d_value],
                "mode": mode,
                "set_no": args.set_no,
                "old_count": len(old_urls),
                "final_count": len(final_urls),
                "old_t1": old_urls[0] if old_urls else "",
                "new_t1": final_urls[0] if final_urls else "",
                "old_t4": old_urls[3] if len(old_urls) >= 4 else "",
                "new_t4": final_urls[3] if len(final_urls) >= 4 else "",
                "size_url_selected": size_url,
                "t4_is_size": bool(len(final_urls) >= 4 and is_size_url(final_urls[3])),
                "final_urls": final_urls,
                "source_candidate_id": record.get("selected_candidate_id"),
                "source_stage": record.get("selected_stage"),
                "source_id": record.get("source_id"),
            }
        )

    weight_changes: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        if d_value.startswith("L085"):
            old = ws.cell(row_idx, col_weight).value
            if old != 150:
                ws.cell(row_idx, col_weight).value = 150
                weight_changes.append({"row": row_idx, "D": d_value, "old": old, "new": 150})

    if encode_text:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "http" in cell.value:
                        cell.value = cell.value.replace(BROKEN_L095_T4, FULL_L095_T4)
                        cell.value = encode_text(cell.value)[0].replace(BROKEN_L095_T4, FULL_L095_T4)

    wb.save(output_path)
    wb.close()

    wb2 = load_workbook(output_path, read_only=True, data_only=False, keep_links=False)
    ws2 = wb2.active
    h2 = headers(ws2)
    c_title = col(h2, "产品标题")
    c_d = col(h2, "产品货号")
    c_g = col(h2, "变种属性值一")
    c_sku = col(h2, "SKU货号")
    c_j = col(h2, "预览图")
    c_t = col(h2, "轮播图")
    c_u = col(h2, "产品素材图")
    c_weight = col(h2, "重量")
    c_skc = col(h2, "SKC属性")
    by_d_titles: dict[str, set[str]] = defaultdict(set)
    by_d_t: dict[str, set[str]] = defaultdict(set)
    title_code_by_d: dict[str, str] = {}
    issues: dict[str, list[Any]] = defaultdict(list)
    effective_rows = 0
    j_records: list[dict[str, Any]] = list(j_write_records)
    j_record_by_row = {int(record["row"]): record for record in j_records}
    for row_idx, row in enumerate(ws2.iter_rows(min_row=2, values_only=True), start=2):
        d_value = str(row[c_d - 1] or "").strip()
        if not d_value:
            continue
        effective_rows += 1
        title = str(row[c_title - 1] or "")
        g_value = str(row[c_g - 1] or "")
        sku_value = str(row[c_sku - 1] or "")
        j_value = str(row[c_j - 1] or "").strip()
        urls = split_urls(row[c_t - 1])
        u_value = str(row[c_u - 1] or "").strip()
        skc_value = str(row[c_skc - 1] or "")
        by_d_titles[d_value].add(title)
        by_d_t[d_value].add("\n".join(urls))
        code_match = re.search(r"\s([A-Z][0-9][A-Z])$", title)
        if not code_match:
            issues["title_code_missing"].append({"row": row_idx, "D": d_value, "title": title})
        else:
            title_code_by_d.setdefault(d_value, code_match.group(1))
            if title_code_by_d[d_value] != code_match.group(1):
                issues["title_code_mismatch"].append({"row": row_idx, "D": d_value, "title": title})
        forbidden_title_words = [word for word in FORBIDDEN_TITLE_WORDS if word.lower() in title.lower()]
        if forbidden_title_words:
            issues["forbidden_title_words"].append({"row": row_idx, "D": d_value, "words": forbidden_title_words, "title": title})
        if not j_value:
            issues["empty_j"].append({"row": row_idx, "D": d_value})
        expected_j = j_record_by_row.get(row_idx, {}).get("oss_url")
        if expected_j and j_value != expected_j:
            issues["j_not_rebuilt_url"].append({"row": row_idx, "D": d_value, "expected": expected_j, "actual": j_value})
        if expected_j and expected_j not in skc_value:
            issues["skc_preview_not_synced_to_j"].append({"row": row_idx, "D": d_value})
        if not urls:
            issues["empty_t"].append({"row": row_idx, "D": d_value})
        if len(urls) > 10:
            issues["t_over_10"].append({"row": row_idx, "D": d_value, "count": len(urls)})
        if len(urls) < 6:
            issues["t_below_6"].append({"row": row_idx, "D": d_value, "count": len(urls)})
        if urls and u_value != urls[0]:
            issues["u_not_t1"].append({"row": row_idx, "D": d_value})
        if len(urls) < 4 or not is_size_url(urls[3]):
            issues["t4_not_size"].append({"row": row_idx, "D": d_value, "T4": urls[3] if len(urls) >= 4 else ""})
        expected_t1 = upload_records[d_value]["uploaded_url"]
        if urls and urls[0] != expected_t1:
            issues["t1_not_selected_set"].append({"row": row_idx, "D": d_value})
        if d_value.startswith("L085") and row[c_weight - 1] != 150:
            issues["l085_weight_not_150"].append({"row": row_idx, "D": d_value, "weight": row[c_weight - 1]})
        if d_value.startswith("L085"):
            residual = []
            for idx, value in enumerate(row, start=1):
                if value is not None and "1500" in str(value):
                    residual.append({"column": h2[idx - 1], "value": str(value)[:300]})
            if residual:
                issues["l085_1500_residual"].append({"row": row_idx, "D": d_value, "residual": residual})
        if d_value.startswith("L058") and row[c_weight - 1] != 3500:
            issues["l058_weight_changed"].append({"row": row_idx, "D": d_value, "weight": row[c_weight - 1]})
    wb2.close()

    for d_value, titles in by_d_titles.items():
        if len(titles) > 1:
            issues["same_d_title_mismatch"].append({"D": d_value})
    for d_value, t_values in by_d_t.items():
        if len(t_values) > 1:
            issues["same_d_t_mismatch"].append({"D": d_value})
    duplicate_title_codes = [code for code, count in Counter(title_code_by_d.values()).items() if count > 1]
    if duplicate_title_codes:
        issues["duplicate_title_code"].extend({"code": code} for code in duplicate_title_codes)

    j_missing_source = [r for r in j_records if not r.get("sku_source")]
    j_missing_file = [r for r in j_records if r.get("sku_source") and not Path(str(r["sku_source"])).exists()]
    j_warning = [r for r in j_records if r.get("warning")]
    j_forbidden_sources = [
        r for r in j_records if any(word.lower() in str(r.get("sku_source", "")).lower() for word in FORBIDDEN_SOURCE_WORDS)
    ]
    issue_counts = {name + "_count": len(items) for name, items in sorted(issues.items())}
    for required in [
        "empty_j_count",
        "empty_t_count",
        "t_over_10_count",
        "t_below_6_count",
        "u_not_t1_count",
        "t4_not_size_count",
        "t1_not_selected_set_count",
        "same_d_title_mismatch_count",
        "same_d_t_mismatch_count",
        "title_code_missing_count",
        "title_code_mismatch_count",
        "duplicate_title_code_count",
        "forbidden_title_words_count",
        "j_not_rebuilt_url_count",
        "skc_preview_not_synced_to_j_count",
        "l085_1500_residual_count",
        "l085_weight_not_150_count",
        "l058_weight_changed_count",
    ]:
        issue_counts.setdefault(required, 0)
    issue_counts.update(
        {
            "j_missing_source_count": len(j_missing_source),
            "j_missing_file_count": len(j_missing_file),
            "j_warning_count": len(j_warning),
            "j_forbidden_source_count": len(j_forbidden_sources),
        }
    )

    url_safety = scan_url_safety(output_path) if scan_url_safety else {"skipped": True}
    t4_validation = validate_t4(output_path) if validate_t4 else {"skipped": True}
    if isinstance(url_safety, dict) and "unsafe_url_occurrences_after" in url_safety:
        issue_counts["unsafe_url_occurrences_after"] = url_safety["unsafe_url_occurrences_after"]
    if isinstance(t4_validation, dict) and "t4_unreachable_count" in t4_validation:
        issue_counts["t4_unreachable_count"] = t4_validation["t4_unreachable_count"]

    hard_check_exempt = {"j_warning_count"}
    pass_hard_checks = all(value == 0 for key, value in issue_counts.items() if key not in hard_check_exempt and isinstance(value, int))
    validation = {
        "source": str(args.source),
        "output": str(output_path),
        "set_no": args.set_no,
        "effective_rows": effective_rows,
        "unique_d": len(by_d_t),
        "deleted_blank_d_rows": sorted(deleted_rows),
        "changed_d_count": len(changes),
        "title_changes_count": len(title_changes),
        "j_changes_count": len(j_cell_changes),
        "skc_preview_changes_count": len(skc_preview_changes),
        "weight_changes": weight_changes,
        "issue_counts": issue_counts,
        "pass_hard_checks": pass_hard_checks,
        "examples": {key: value[:30] for key, value in issues.items()},
        "j_audit_summary": {
            "records": len(j_records),
            "missing_source": len(j_missing_source),
            "missing_file": len(j_missing_file),
            "warnings": len(j_warning),
            "forbidden_sources": len(j_forbidden_sources),
            "note": "J was rebuilt for every effective row and SKC属性.previewImgUrls was synced to the rebuilt J URL.",
        },
        "title_rewrite_summary": {
            "changed_cells": len(title_changes),
            "unique_codes": len(set(title_code_by_d.values())),
            "salt": title_salt,
        },
        "t_count_dist": dict(Counter(item["final_count"] for item in changes)),
        "mode_counts": dict(Counter(item["mode"] for item in changes)),
        "source_stage_counts": dict(Counter(item["source_stage"] for item in changes)),
        "url_safety": url_safety,
        "t4_validation": t4_validation,
    }
    report = {
        "source_workbook": str(args.source),
        "output_workbook": str(output_path),
        "pool": str(args.pool),
        "set_no": args.set_no,
        "changes": changes,
        "upload_records": upload_records,
        "title_changes": title_changes,
        "j_cell_changes": j_cell_changes,
        "skc_preview_changes": skc_preview_changes,
        "weight_changes": weight_changes,
        "validation": validation,
    }
    diff = {
        "allowed_columns": ["产品标题", "预览图", "SKC属性", "轮播图", "产品素材图", "重量"],
        "deleted_blank_d_rows": sorted(deleted_rows),
        "changed_cells": [
            {"row": item["row"], "D": item["D"], "column": "产品标题", "old": item["old"], "new": item["new"]}
            for item in title_changes
        ]
        + [
            {"row": item["row"], "D": item["D"], "column": "预览图", "old": item["old"], "new": item["new"]}
            for item in j_cell_changes
        ]
        + [
            {"row": item["row"], "D": item["D"], "column": "SKC属性", "old": item["old"], "new": item["new"]}
            for item in skc_preview_changes
        ]
        + [
            {"row": row, "D": item["D"], "column": "轮播图", "old": snapshots[str(row)]["T"], "new": "\n".join(item["final_urls"])}
            for item in changes
            for row in item["rows"]
        ]
        + [
            {"row": row, "D": item["D"], "column": "产品素材图", "old": snapshots[str(row)]["U"], "new": item["final_urls"][0]}
            for item in changes
            for row in item["rows"]
        ]
        + [
            {"row": item["row"], "D": item["D"], "column": "重量", "old": item["old"], "new": item["new"]}
            for item in weight_changes
        ],
    }

    save_json(args.out_dir / "writeback_report.json", report)
    save_json(args.out_dir / "validation_report.json", validation)
    save_json(args.out_dir / "j_match_audit.json", {"records": j_records})
    save_json(args.out_dir / "cell_diff_summary.json", diff)
    build_t_audit_html(args.out_dir / "t_writeback_audit.html", report)
    build_j_audit_html(args.out_dir / "j_match_audit.html", j_records)
    print(json.dumps(validation, ensure_ascii=False, indent=2))
    return 0 if pass_hard_checks else 2


if __name__ == "__main__":
    raise SystemExit(main())
