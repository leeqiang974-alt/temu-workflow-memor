"""Append the certified YeahF 398-D workbook to the global fingerprint registry."""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import time
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(r"C:\Users\Administrator\Documents\temu自动化\outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718\YeahF_398D_最终可提交_ReleaseGuard认证_20260719.xlsx")
CERTIFICATE = Path(r"C:\Users\Administrator\Documents\temu自动化\outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718\release_guard_state\batches\YF398-20260719-SORTED-JT1-V1\release_certificate.json")
REGISTRY = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui\work\temu_workbook_registry.json")
ENTRY_ID = "yeahf-398d-certified-20260719"
MIXED_FP_RE = re.compile(r"(?=.*[A-Z])(?=.*\d)[A-Z0-9]{4}$")
LEGACY_FP_RE = re.compile(r"\d{4}$")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def groups(path: Path, *, require_mixed: bool = False) -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    d_col = headers.index("产品货号")
    title_col = headers.index("产品标题")
    result = {}
    for row in rows:
        d_value = str(row[d_col] or "").strip().upper()
        title = str(row[title_col] or "").strip()
        if not d_value:
            continue
        fingerprint = title[-4:].upper()
        valid = bool(MIXED_FP_RE.fullmatch(fingerprint)) if require_mixed else bool(MIXED_FP_RE.fullmatch(fingerprint) or LEGACY_FP_RE.fullmatch(fingerprint))
        if not valid:
            raise RuntimeError(f"invalid fingerprint: {d_value} / {fingerprint}")
        old = result.get(d_value)
        current = {"title": title, "fingerprint": fingerprint}
        if old and old != current:
            raise RuntimeError(f"same-D title drift: {d_value}")
        result[d_value] = current
    workbook.close()
    return result


def main() -> None:
    cert = json.loads(CERTIFICATE.read_text(encoding="utf-8"))
    workbook_hash = sha256(WORKBOOK)
    if cert.get("status") != "CERTIFIED" or cert.get("workbook_sha256") != workbook_hash:
        raise RuntimeError("certificate does not match workbook")
    candidate = groups(WORKBOOK, require_mixed=True)
    if len(candidate) != 398:
        raise RuntimeError(f"expected 398 D, got {len(candidate)}")
    local = defaultdict(list)
    for d_value, item in candidate.items():
        local[item["fingerprint"]].append(d_value)
    duplicate = {key: values for key, values in local.items() if len(values) > 1}
    if duplicate:
        raise RuntimeError(f"candidate duplicate fingerprints: {duplicate}")

    registry = json.loads(REGISTRY.read_text(encoding="utf-8"))
    workbooks = registry.setdefault("workbooks", [])
    existing_path = str(WORKBOOK).lower()
    existing_entry = next((item for item in workbooks if item.get("id") == ENTRY_ID), None)
    existing_fingerprints = defaultdict(list)
    for record in workbooks:
        if not record.get("enabled_for_fingerprint_lookup"):
            continue
        if record.get("id") == ENTRY_ID:
            continue
        path = Path(str(record.get("path") or ""))
        if path.resolve() == WORKBOOK.resolve():
            continue
        if not path.is_file():
            raise RuntimeError(f"enabled registry workbook missing: {path}")
        for d_value, item in groups(path).items():
            existing_fingerprints[item["fingerprint"]].append({"D": d_value, "path": str(path)})
    collisions = {item["fingerprint"]: existing_fingerprints[item["fingerprint"]] for item in candidate.values() if item["fingerprint"] in existing_fingerprints}
    if collisions:
        raise RuntimeError(f"global fingerprint collision count={len(collisions)}")

    entry = {
        "id": ENTRY_ID,
        "store_metadata": "YeahF",
        "path": str(WORKBOOK),
        "status": "certified_final_submission",
        "enabled_for_fingerprint_lookup": True,
        "approved_at": "2026-07-19",
        "batch_id": "YF398-20260719-SORTED-JT1-V1",
        "certificate_path": str(CERTIFICATE),
        "certificate_sha256": sha256(CERTIFICATE),
        "workbook_sha256": workbook_hash,
        "exact_D_count": 398,
        "note": "Release Guard certified YeahF 398-D workbook. Append-only global trailing fingerprint lookup; store is metadata only.",
    }
    if existing_entry == entry:
        pass
    else:
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        backup = REGISTRY.with_name(f"{REGISTRY.stem}.backup_{timestamp}{REGISTRY.suffix}")
        shutil.copy2(REGISTRY, backup)
        if existing_entry:
            if not str(existing_entry.get("path") or "").endswith("YeahF_398D_guarded_candidate_output.xlsx"):
                raise RuntimeError("existing registry id points to an unexpected workbook")
            workbooks[workbooks.index(existing_entry)] = entry
        else:
            workbooks.append(entry)
        registry["updated_at"] = "2026-07-19"
        temp = REGISTRY.with_suffix(REGISTRY.suffix + ".tmp")
        temp.write_text(json.dumps(registry, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temp, REGISTRY)
    check = json.loads(REGISTRY.read_text(encoding="utf-8"))
    matches = [item for item in check.get("workbooks", []) if item.get("id") == ENTRY_ID and item.get("enabled_for_fingerprint_lookup")]
    if len(matches) != 1 or matches[0].get("workbook_sha256") != workbook_hash:
        raise RuntimeError("registry re-read verification failed")
    first_d = sorted(candidate)[0]
    print(json.dumps({
        "status": "registered",
        "registry": str(REGISTRY),
        "entry_id": ENTRY_ID,
        "workbook_sha256": workbook_hash,
        "exact_D": len(candidate),
        "global_collision_count": 0,
        "probe": {"D": first_d, "fingerprint": candidate[first_d]["fingerprint"], "title": candidate[first_d]["title"]},
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
