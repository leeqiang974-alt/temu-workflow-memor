#!/usr/bin/env python3
"""Write back 197x3 set3 T images into the current 195-D workbook and audit J.

This is intentionally scoped to the 2026-07-05 request:
- delete rows whose 产品货号/D is empty
- keep existing titles unchanged
- replace T1 with the already uploaded 197x3 set3 URL
- set U = T1
- preserve a valid size image at T4
- cap T at 10 URLs
- do not rebuild J, only build a row-level J source-match audit
"""

from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import urllib.parse
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from encode_workbook_image_urls import encode_text, scan_url_safety, validate_t4
except Exception:  # pragma: no cover - local script fallback
    encode_text = None
    scan_url_safety = None
    validate_t4 = None


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
FORBIDDEN_SOURCE_WORDS = ("九宫格", "9grid", "out", "output", "背景素材")
SIZE_RE = re.compile(r"(尺寸|尺码|size|cm|inch|length|height|width|宽|长|高)", re.I)
URL_RE = re.compile(r"https?://.*?(?=(?:[,;，；]?\s*https?://)|$)", re.S)

BROKEN_L095_T4 = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "85ef73d852584a65aaf05a1c5615dac7_L095_%E5%B0%BA%E5%AF%B8_jimeng-2026-05-08-1613-"
    "%E8%BF%99%E5%BC%A0%E5%9B%BE%E7%89%87"
)

FULL_L095_T4 = (
    "https://ozonshanghai.oss-cn-shanghai.aliyuncs.com/temu-jit/carousel-ocr-size/20260605/"
    "85ef73d852584a65aaf05a1c5615dac7_L095_%E5%B0%BA%E5%AF%B8_jimeng-2026-05-08-1613-"
    "%E8%BF%99%E5%BC%A0%E5%9B%BE%E7%89%87%EF%BC%8C%E4%BF%9D%E6%8C%81%E4%BA%A7%E5%93%81"
    "%E5%92%8C%E6%A0%87%E5%B0%BA%E3%80%81%E6%95%B0%E5%AD%97%E3%80%81%E6%96%87%E5%AD%97"
    "%E4%BF%A1%E6%81%AF%E4%B8%8D%E5%81%9A.jpg"
)


def headers(ws) -> list[str]:
    return [str(ws.cell(1, column).value or "").strip() for column in range(1, ws.max_column + 1)]


def col(header_values: list[str], name: str) -> int:
    if name not in header_values:
        raise RuntimeError(f"missing required column: {name}")
    return header_values.index(name) + 1


def split_urls(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    urls: list[str] = []
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line:
            continue
        matches = [m.group(0).strip() for m in URL_RE.finditer(line)]
        urls.extend(matches or [line])
    return urls


def dedupe(values: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            out.append(value)
            seen.add(value)
    return out


def is_size_url(url: str) -> bool:
    """Detect size charts from the object name, avoiding broad parent folder names."""
    text = str(url or "")
    if not text:
        return False
    parts = urllib.parse.urlsplit(text)
    decoded_path = urllib.parse.unquote(parts.path or text)
    basename = decoded_path.rsplit("/", 1)[-1]
    if SIZE_RE.search(basename):
        return True
    # Historical OCR size chart keys live under this dedicated folder. Keep this
    # narrow so folder names like "...尺寸T4最终校验..." do not misclassify T images.
    return "/carousel-ocr-size/" in decoded_path.lower()


def token_set(text: object) -> set[str]:
    raw = str(text or "").lower()
    mapping = [
        ("black", ["黑", "black"]),
        ("green", ["绿", "green"]),
        ("white", ["白", "透明", "white"]),
        ("gray", ["灰", "grey", "gray"]),
        ("yellow", ["黄", "yellow"]),
        ("pink", ["粉", "pink"]),
        ("red", ["红", "red"]),
        ("blue", ["蓝", "blue"]),
        ("purple", ["紫", "purple"]),
        ("wood", ["木色", "原木", "胡桃", "walnut", "wood"]),
        ("beige", ["米", "杏", "beige"]),
        ("15", ["15"]),
        ("30", ["30"]),
        ("50", ["50"]),
        ("2", ["2", "双", "两", "二"]),
        ("3", ["3", "三"]),
    ]
    found: set[str] = set()
    for canonical, aliases in mapping:
        if any(alias in raw for alias in aliases):
            found.add(canonical)
    return found


def load_set3_manifest(path: Path) -> dict[str, str]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    uploaded = data.get("uploaded", data)
    result: dict[str, str] = {}
    for key, url in uploaded.items():
        parts = str(key).split("|")
        if len(parts) >= 2 and parts[1] == "set3":
            result[parts[0]] = str(url)
    if not result:
        raise RuntimeError(f"no set3 URLs found in {path}")
    return result


def find_size_url(urls: list[str]) -> str:
    for url in urls:
        if is_size_url(url):
            return url
    return urls[3] if len(urls) >= 4 else ""


def build_t_urls(old_urls: list[str], new_t1: str) -> tuple[list[str], str]:
    size_url = find_size_url(old_urls)
    rest = [url for idx, url in enumerate(old_urls) if idx != 0 and url not in {new_t1, size_url}]
    if size_url:
        urls = [new_t1] + rest[:2] + [size_url] + rest[2:]
    else:
        urls = [new_t1] + rest
    return dedupe(urls)[:10], size_url


def clean_source_files(folder: Path) -> list[Path]:
    files: list[Path] = []
    if not folder.exists():
        return files
    for path in folder.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTS:
            continue
        lower = str(path).lower()
        if any(word.lower() in lower for word in FORBIDDEN_SOURCE_WORDS):
            continue
        files.append(path)
    return sorted(files, key=lambda p: str(p))


def choose_j_source(prefix: str, variant: str, sku: str, roots: list[Path]) -> dict[str, Any]:
    wanted = token_set(f"{variant} {sku}")
    if prefix == "L043":
        base = roots[0] / "L043" / "sku文件_最终抠图PNG"
        color = "灰" if "gray" in wanted else "白"
        source = base / color / f"{color}.png"
        return {
            "sku_root": str(base),
            "sku_source": str(source),
            "wanted_tokens": sorted(wanted),
            "matched_tokens": sorted(wanted & token_set(str(source))),
            "match_mode": "l043_historical_white_gray_hard_rule",
            "warning": "" if source.exists() else "missing_l043_source",
        }

    folder_names = ["sku文件_最终抠图PNG", "sku文件", "sku"]
    chosen_folder = None
    candidates: list[Path] = []
    for root in roots:
        for folder_name in folder_names:
            folder = root / prefix / folder_name
            files = clean_source_files(folder)
            if files:
                chosen_folder = folder
                candidates = files
                break
        if candidates:
            break
    if not candidates or chosen_folder is None:
        return {
            "sku_root": "",
            "sku_source": "",
            "wanted_tokens": sorted(wanted),
            "matched_tokens": [],
            "match_mode": "missing_source",
            "warning": "missing_source_folder",
        }

    def score(path: Path) -> tuple[int, str]:
        got = token_set(str(path))
        score_value = len(wanted & got) * 10
        for token in {"black", "green", "white", "gray", "pink", "red", "blue", "purple", "yellow", "wood", "beige"}:
            if token in wanted and token in got:
                score_value += 40
        for token in {"15", "30", "50", "2", "3"}:
            if token in wanted and token in got:
                score_value += 8
        if wanted & {"black", "green", "white", "gray", "pink", "red", "blue", "purple", "yellow", "wood", "beige"} and not (
            wanted & got & {"black", "green", "white", "gray", "pink", "red", "blue", "purple", "yellow", "wood", "beige"}
        ):
            score_value -= 25
        return score_value, str(path)

    scored = sorted(candidates, key=score, reverse=True)
    top = scored[0]
    top_score = score(top)[0]
    got = token_set(str(top))
    return {
        "sku_root": str(chosen_folder),
        "sku_source": str(top),
        "wanted_tokens": sorted(wanted),
        "matched_tokens": sorted(wanted & got),
        "match_score": top_score,
        "match_mode": "token_match" if top_score > 0 else "prefix_fallback",
        "warning": "" if top_score > 0 else "no_positive_token_match",
    }


def build_j_audit_html(path: Path, records: list[dict[str, Any]], title: str) -> None:
    cards = []
    for record in records:
        warning = record.get("warning") or ""
        cards.append(
            "<section class='card'>"
            f"<h3>{html.escape(str(record['D']))} <span>row {record['row']}</span></h3>"
            f"<p><b>G:</b> {html.escape(str(record.get('G','')))}<br><b>SKU:</b> {html.escape(str(record.get('SKU','')))}</p>"
            f"<p><b>source:</b> {html.escape(str(record.get('sku_source','')))}<br>"
            f"<b>mode:</b> {html.escape(str(record.get('match_mode','')))}<br>"
            f"<b>wanted:</b> {html.escape(','.join(record.get('wanted_tokens') or []))}<br>"
            f"<b>matched:</b> {html.escape(','.join(record.get('matched_tokens') or []))}</p>"
            f"<p class='warn'>{html.escape(warning)}</p>"
            f"<div class='imgs'><figure><figcaption>当前J</figcaption><img loading='lazy' src='{html.escape(str(record.get('current_j','')))}'></figure>"
            f"<figure><figcaption>匹配源图</figcaption><img loading='lazy' src='{html.escape(Path(str(record.get('sku_source',''))).as_uri() if record.get('sku_source') and Path(str(record.get('sku_source'))).exists() else '')}'></figure></div>"
            "</section>"
        )
    path.write_text(
        f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;background:#f5f6f8;color:#172033;margin:18px}}
.card{{background:#fff;border:1px solid #d8dee9;border-radius:8px;padding:12px;margin:10px 0}}
h1{{font-size:22px}} h3{{font-size:17px;margin:0 0 6px}} h3 span{{font-size:12px;color:#667085}}
p{{font-size:13px;word-break:break-all}} .warn{{color:#b42318;font-weight:700}}
.imgs{{display:grid;grid-template-columns:180px 180px;gap:12px}} figure{{margin:0}} figcaption{{font-size:12px;font-weight:700}}
img{{width:180px;height:180px;object-fit:contain;background:#fff;border:1px solid #e5e7eb}}
</style></head><body><h1>{html.escape(title)} · {len(records)} 行</h1>{''.join(cards)}</body></html>""",
        encoding="utf-8",
    )


def build_t_audit_html(path: Path, changes: list[dict[str, Any]], workbook_path: Path, validation: dict[str, Any]) -> None:
    sections = []
    for item in changes:
        imgs = "".join(
            f"<figure><figcaption>T{i+1}{' 尺寸图' if i == 3 else ''}</figcaption><img loading='lazy' src='{html.escape(url)}'><div>{html.escape(url)}</div></figure>"
            for i, url in enumerate(item["final_urls"][:10])
        )
        sections.append(
            f"<section><h2>{html.escape(item['D'])}</h2><p>rows: {item['rows']} | old_count: {item['old_count']} | final_count: {item['final_count']} | T4 size: {item['t4_is_size']}</p><div class='grid'>{imgs}</div></section>"
        )
    path.write_text(
        f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>set3 T 回填审核</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;background:#f5f6f8;color:#172033;margin:18px}}
.summary,section{{background:#fff;border:1px solid #d8dee9;border-radius:8px;padding:12px;margin:10px 0}}
h1{{font-size:22px}} h2{{font-size:17px}} pre{{white-space:pre-wrap;background:#f3f4f6;padding:10px;border-radius:6px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(150px,1fr));gap:10px}} figure{{margin:0}}
figcaption{{font-size:12px;font-weight:700}} img{{width:100%;aspect-ratio:1/1;object-fit:contain;background:#fff;border:1px solid #e5e7eb}}
figure div{{font-size:10px;color:#667085;word-break:break-all;max-height:48px;overflow:auto}}
</style></head><body><div class='summary'><h1>197 第三套 T 首图回填审核</h1><p>{html.escape(str(workbook_path))}</p><pre>{html.escape(json.dumps(validation, ensure_ascii=False, indent=2))}</pre></div>{''.join(sections)}</body></html>""",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--set3-manifest", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--output-name", required=True)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    output_path = args.out_dir / args.output_name
    shutil.copy2(args.source, output_path)

    set3_by_d = load_set3_manifest(args.set3_manifest)
    wb = load_workbook(output_path)
    ws = wb.active
    h = headers(ws)
    col_title = col(h, "产品标题")
    col_d = col(h, "产品货号")
    col_g = col(h, "变种属性值一")
    col_sku = col(h, "SKU货号")
    col_j = col(h, "预览图")
    col_t = col(h, "轮播图")
    col_u = col(h, "产品素材图")

    deleted_rows: list[int] = []
    for row_idx in range(ws.max_row, 1, -1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        if not d_value:
            deleted_rows.append(row_idx)
            ws.delete_rows(row_idx, 1)

    d_rows: dict[str, list[int]] = defaultdict(list)
    first_row_by_d: dict[str, int] = {}
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        if not d_value:
            continue
        d_rows[d_value].append(row_idx)
        first_row_by_d.setdefault(d_value, row_idx)

    missing_set3 = sorted(set(d_rows) - set(set3_by_d))
    if missing_set3:
        raise RuntimeError(f"source D missing set3 URL: {missing_set3[:30]}")

    changes: list[dict[str, Any]] = []
    old_titles = {d: str(ws.cell(rows[0], col_title).value or "") for d, rows in d_rows.items()}
    for d_value, rows in sorted(d_rows.items()):
        first_row = rows[0]
        old_urls = split_urls(ws.cell(first_row, col_t).value)
        new_t1 = set3_by_d[d_value]
        final_urls, size_url = build_t_urls(old_urls, new_t1)
        for row_idx in rows:
            ws.cell(row_idx, col_t).value = "\n".join(final_urls)
            ws.cell(row_idx, col_u).value = final_urls[0] if final_urls else ""
        changes.append(
            {
                "D": d_value,
                "rows": rows,
                "old_count": len(old_urls),
                "final_count": len(final_urls),
                "old_t1": old_urls[0] if old_urls else "",
                "new_t1": final_urls[0] if final_urls else "",
                "old_t4": old_urls[3] if len(old_urls) >= 4 else "",
                "new_t4": final_urls[3] if len(final_urls) >= 4 else "",
                "size_url_selected": size_url,
                "t4_is_size": bool(len(final_urls) >= 4 and SIZE_RE.search(final_urls[3])),
                "final_urls": final_urls,
            }
        )

    j_roots = [Path(r"E:\JIT制图--新店"), Path(r"E:\jit制图")]
    j_records: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row_idx, col_d).value or "").strip()
        if not d_value:
            continue
        g_value = str(ws.cell(row_idx, col_g).value or "").strip()
        sku_value = str(ws.cell(row_idx, col_sku).value or "").strip()
        match = choose_j_source(d_value[:4], g_value, sku_value, j_roots)
        j_records.append(
            {
                "row": row_idx,
                "D": d_value,
                "G": g_value,
                "SKU": sku_value,
                "current_j": str(ws.cell(row_idx, col_j).value or "").strip(),
                **match,
            }
        )

    if encode_text:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "http" in cell.value:
                        cell.value = cell.value.replace(BROKEN_L095_T4, FULL_L095_T4)
                        cell.value = encode_text(cell.value)[0].replace(BROKEN_L095_T4, FULL_L095_T4)

    wb.save(output_path)
    wb.close()

    # Re-open read-only for validation after save.
    wb2 = load_workbook(output_path, read_only=True, data_only=False, keep_links=False)
    ws2 = wb2.active
    h2 = headers(ws2)
    c_title = col(h2, "产品标题")
    c_d = col(h2, "产品货号")
    c_j = col(h2, "预览图")
    c_t = col(h2, "轮播图")
    c_u = col(h2, "产品素材图")
    by_d_titles: dict[str, set[str]] = defaultdict(set)
    by_d_t: dict[str, set[str]] = defaultdict(set)
    issue_examples: dict[str, list[Any]] = defaultdict(list)
    effective_rows = 0
    for row_idx, row in enumerate(ws2.iter_rows(min_row=2, values_only=True), start=2):
        d_value = str(row[c_d - 1] or "").strip()
        if not d_value:
            continue
        effective_rows += 1
        title = str(row[c_title - 1] or "")
        j_value = str(row[c_j - 1] or "").strip()
        urls = split_urls(row[c_t - 1])
        u_value = str(row[c_u - 1] or "").strip()
        by_d_titles[d_value].add(title)
        by_d_t[d_value].add("\n".join(urls))
        if not j_value:
            issue_examples["empty_j"].append({"row": row_idx, "D": d_value})
        if not urls:
            issue_examples["empty_t"].append({"row": row_idx, "D": d_value})
        if len(urls) > 10:
            issue_examples["t_over_10"].append({"row": row_idx, "D": d_value, "count": len(urls)})
        if len(urls) < 6:
            issue_examples["t_below_6"].append({"row": row_idx, "D": d_value, "count": len(urls)})
        if urls and u_value != urls[0]:
            issue_examples["u_not_t1"].append({"row": row_idx, "D": d_value})
        if len(urls) < 4 or not is_size_url(urls[3]):
            issue_examples["t4_not_size"].append({"row": row_idx, "D": d_value, "T4": urls[3] if len(urls) >= 4 else ""})
        if urls and urls[0] != set3_by_d.get(d_value):
            issue_examples["t1_not_set3"].append({"row": row_idx, "D": d_value})
    wb2.close()

    for d_value, titles in by_d_titles.items():
        if len(titles) > 1:
            issue_examples["same_d_title_mismatch"].append({"D": d_value})
        if old_titles.get(d_value) not in titles:
            issue_examples["title_changed"].append({"D": d_value})
    for d_value, t_values in by_d_t.items():
        if len(t_values) > 1:
            issue_examples["same_d_t_mismatch"].append({"D": d_value})

    j_missing_source = [r for r in j_records if not r.get("sku_source")]
    j_missing_file = [r for r in j_records if r.get("sku_source") and not Path(str(r["sku_source"])).exists()]
    j_warning = [r for r in j_records if r.get("warning")]
    forbidden_sources = [
        r for r in j_records if any(word.lower() in str(r.get("sku_source", "")).lower() for word in FORBIDDEN_SOURCE_WORDS)
    ]

    issue_counts = {name + "_count": len(items) for name, items in sorted(issue_examples.items())}
    for required in [
        "empty_j_count",
        "empty_t_count",
        "t_over_10_count",
        "t_below_6_count",
        "u_not_t1_count",
        "t4_not_size_count",
        "t1_not_set3_count",
        "same_d_title_mismatch_count",
        "same_d_t_mismatch_count",
        "title_changed_count",
    ]:
        issue_counts.setdefault(required, 0)
    issue_counts.update(
        {
            "j_missing_source_count": len(j_missing_source),
            "j_missing_file_count": len(j_missing_file),
            "j_warning_count": len(j_warning),
            "j_forbidden_source_count": len(forbidden_sources),
        }
    )

    url_safety = scan_url_safety(output_path) if scan_url_safety else {"skipped": True}
    t4_validation = validate_t4(output_path) if validate_t4 else {"skipped": True}
    if isinstance(url_safety, dict) and "unsafe_url_occurrences_after" in url_safety:
        issue_counts["unsafe_url_occurrences_after"] = url_safety["unsafe_url_occurrences_after"]
    if isinstance(t4_validation, dict) and "t4_unreachable_count" in t4_validation:
        issue_counts["t4_unreachable_count"] = t4_validation["t4_unreachable_count"]

    validation = {
        "source": str(args.source),
        "output": str(output_path),
        "effective_rows": effective_rows,
        "unique_d": len(by_d_t),
        "deleted_rows": sorted(deleted_rows),
        "set3_manifest": str(args.set3_manifest),
        "issue_counts": issue_counts,
        "pass_hard_checks": all(
            v == 0
            for k, v in issue_counts.items()
            if k not in {"j_warning_count"} and not str(v).startswith("skipped")
        ),
        "examples": {k: v[:30] for k, v in issue_examples.items()},
        "j_audit_summary": {
            "records": len(j_records),
            "missing_source": len(j_missing_source),
            "missing_file": len(j_missing_file),
            "warnings": len(j_warning),
            "forbidden_sources": len(forbidden_sources),
        },
        "t_count_dist": dict(Counter(item["final_count"] for item in changes)),
        "url_safety": url_safety,
        "t4_validation": t4_validation,
    }

    writeback_report = {
        "source": str(args.source),
        "output": str(output_path),
        "changes": changes,
        "validation": validation,
    }
    j_audit = {"source": str(args.source), "output": str(output_path), "records": j_records}

    (args.out_dir / "writeback_report.json").write_text(json.dumps(writeback_report, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out_dir / "j_match_audit.json").write_text(json.dumps(j_audit, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out_dir / "validation_report.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    build_t_audit_html(args.out_dir / "t_writeback_audit.html", changes, output_path, validation)
    build_j_audit_html(args.out_dir / "j_match_audit.html", j_records, "J 图 SKU 属性匹配审核")

    print(json.dumps(validation, ensure_ascii=False, indent=2))
    return 0 if validation["pass_hard_checks"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
