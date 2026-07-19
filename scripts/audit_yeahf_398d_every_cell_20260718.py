"""Read-only, every-cell audit for all 54 columns of the active YeahF draft."""

from __future__ import annotations

import html
import json
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Administrator\Documents\temu自动化")
OUT = ROOT / "outputs" / "yeahf_title_dedup_20260709" / "yeahf_400d_refill_20260718"
WORKBOOK = Path(os.environ.get("YEAHF_EVERY_CELL_WORKBOOK", str(OUT / "YeahF_398D_剔除T1不可用L085两组_待J_T1回填.xlsx")))
ASSETS = Path(r"D:\temu素材库\T首图已付费库\YeahF_400D_20260716_rebuild\YeahF_400D_resolved_assets_oss_urls_20260717.json")
LOCKS = OUT / "j_historical_rejection_locks_20260718.json"
REPORT = OUT / "YeahF_398D_每格全量审计_20260718.json"
HTML = OUT / "YeahF_398D_每格全量审查_20260718.html"
L095_REPAIR = ROOT / "outputs" / "yeahf_title_dedup_20260709" / "yeahf_400d_refill_20260718" / "YeahF_398D_L095三格J修正_写回报告_20260718.json"


def s(value: object) -> str:
    return "" if value is None else str(value).strip()


def url(value: str) -> bool:
    return value.startswith(("https://", "http://"))


def json_list(value: str) -> tuple[bool, list[dict]]:
    if not value:
        return True, []
    try:
        parsed = json.loads(value)
        return isinstance(parsed, list), [item for item in parsed if isinstance(item, dict)] if isinstance(parsed, list) else []
    except Exception:
        return False, []


def fp(title: str) -> str:
    match = re.search(r"(?:^|\s)([A-Z0-9]{4})\s*$", title)
    return match.group(1) if match else ""


def cell(status: str, rule: str, value: str) -> dict:
    return {"status": status, "rule": rule, "value": value}


def main() -> None:
    asset_items = json.loads(ASSETS.read_text(encoding="utf-8"))["j"]
    asset_index = {(item["D"], item["G"], item["SKU"]): item for item in asset_items}
    # Approved repair assets supersede the historical 2-grid L095 records.
    if L095_REPAIR.is_file():
        for item in json.loads(L095_REPAIR.read_text(encoding="utf-8"))["records"]:
            asset_index[(item["D"], item["G"], item["SKU"])] = {
                "D": item["D"], "G": item["G"], "SKU": item["SKU"],
                "source_path": item["source"], "j_value": item["new_j"],
            }
    locked_d = {item["D"] for item in json.loads(LOCKS.read_text(encoding="utf-8"))["locks"]}
    ws = load_workbook(WORKBOOK, read_only=True, data_only=False).active
    headers = [s(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert len(headers) == 54, f"Expected 54 columns, got {len(headers)}"
    groups: dict[str, list[dict]] = defaultdict(list)
    all_rows: list[dict] = []
    special_allowed_blank_video = {"L048", "L087"}

    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        vals = {headers[index]: s(value) for index, value in enumerate(values)}
        d_value = vals["产品货号"]
        if not d_value:
            continue
        g_value, sku = vals["变种属性值一"], vals["SKU货号"]
        model = f"{d_value}-1"
        asset = asset_index.get((d_value, g_value, sku))
        source_path = s(asset.get("source_path")) if asset else ""
        cells: dict[str, dict] = {}
        cells["产品标题"] = cell("PASS" if fp(vals["产品标题"]) else "BLOCK", "末尾必须是四位指纹；同 D 必须相同", vals["产品标题"])
        cells["英文标题"] = cell("EMPTY_OK" if not vals["英文标题"] else "BLOCK", "YeahF 当前规则：英文标题必须删除/为空", vals["英文标题"])
        cells["产品描述"] = cell("CHECK" if vals["产品描述"] else "BLOCK", "保留原生描述；非空，语义由人工核对", vals["产品描述"])
        cells["产品货号"] = cell("PASS" if re.fullmatch(r"L\d{3}0709\d{2}", d_value) else "BLOCK", "D 必须为 L0xx0709aa；不得重用已剔除号码", d_value)
        cells["变种名称"] = cell("PASS" if vals["变种名称"] == f"{g_value}，{model}" else "BLOCK", "E = G，I", vals["变种名称"])
        cells["变种属性名称一"] = cell("PASS" if vals["变种属性名称一"] else "BLOCK", "F 必须非空，并与 AD 第一属性名一致", vals["变种属性名称一"])
        cells["变种属性值一"] = cell("PASS" if g_value else "BLOCK", "G 必须非空，并与 AD 第一属性值一致", g_value)
        cells["变种属性名称二"] = cell("PASS" if vals["变种属性名称二"] == "型号" else "BLOCK", "H 固定为 型号", vals["变种属性名称二"])
        cells["变种属性值二"] = cell("PASS" if vals["变种属性值二"] == model else "BLOCK", "I 固定为 D-1；同 D 不递增", vals["变种属性值二"])
        j_status = "PASS"
        j_rule = "精确 D+G+SKU J 资产血缘 + 长期 URL"
        if not url(vals["预览图"]):
            j_status, j_rule = "BLOCK", "J 必须是长期 http(s) URL"
        elif not asset:
            j_status, j_rule = "BLOCK", "缺少精确 D+G+SKU 的 J 资产血缘"
        elif s(asset.get("j_value")) != vals["预览图"]:
            j_status, j_rule = "BLOCK", "当前 J 不等于精确资产血缘 URL"
        elif d_value in locked_d:
            j_status, j_rule = "BLOCK", "用户确认的历史错误 D 禁止复用 J 证据"
        elif d_value.startswith("L095") and sku == "L095-00" and not source_path.endswith(r"\2\2.png"):
            j_status, j_rule = "BLOCK", "L095-00 必须取 2格源 ...\\2\\2.png"
        elif d_value.startswith("L095") and sku == "L095-01" and not source_path.endswith(r"\3\3.png"):
            j_status, j_rule = "BLOCK", "L095-01 必须取 3格源 ...\\3\\3.png"
        cells["预览图"] = cell(j_status, j_rule + (f"；源：{source_path}" if source_path else ""), vals["预览图"])
        try:
            cells["申报价格"] = cell("PASS" if float(vals["申报价格"]) > 0 else "BLOCK", "必须为正数；价格来源另由核价表核对", vals["申报价格"])
        except ValueError:
            cells["申报价格"] = cell("BLOCK", "必须为正数", vals["申报价格"])
        cells["SKU货号"] = cell("PASS" if sku and sku.startswith(d_value[:4]) else "BLOCK", "L 必须非空且同 L0xx", sku)
        for name in ("长", "宽", "高"):
            try:
                cells[name] = cell("PASS" if float(vals[name]) > 0 else "BLOCK", "必须为正数；需与 SKU 尺寸来源一致", vals[name])
            except ValueError:
                cells[name] = cell("BLOCK", "必须为正数", vals[name])
        expected_weight = "3500" if d_value.startswith("L058") else "150"
        cells["重量"] = cell("PASS" if vals["重量"] == expected_weight else "BLOCK", f"{d_value[:4]} 规则重量 = {expected_weight}", vals["重量"])
        for name in ("识别码类型", "识别码"):
            pair = "识别码" if name == "识别码类型" else "识别码类型"
            cells[name] = cell("EMPTY_OK" if not vals[name] and not vals[pair] else ("PASS" if vals[name] and vals[pair] else "BLOCK"), "识别码类型与识别码必须成对，或同时为空", vals[name])
        cells["站外产品链接"] = cell("EMPTY_OK" if not vals["站外产品链接"] else ("PASS" if url(vals["站外产品链接"]) else "BLOCK"), "可空；填写时必须为 URL", vals["站外产品链接"])
        t_urls = [part.strip() for part in re.split(r"[\r\n,]+", vals["轮播图"]) if part.strip()]
        t_ok = 4 <= len(t_urls) <= 10 and all(url(item) for item in t_urls) and len(t_urls) >= 4 and "carousel-ocr-size" in t_urls[3]
        cells["轮播图"] = cell("PASS" if t_ok else "BLOCK", "T=4-10 个 URL；T4 必须是同 L0xx 尺寸图", vals["轮播图"])
        cells["产品素材图"] = cell("PASS" if t_urls and vals["产品素材图"] == t_urls[0] else "BLOCK", "U = T1", vals["产品素材图"])
        for name in ("外包装形状", "外包装类型", "建议零售价(建议零售价币种)", "分类id", "来源url", "产地", "敏感属性", "备注", "SKU分类", "SKU分类单位", "净含量单位", "混合套装类型", "SKU分类总数量单位", "总净含量单位", "独立包装", "包装清单", "生命周期", "所属店铺", "SPUID", "SKCID", "SKUID", "创建时间", "更新时间"):
            value = vals[name]
            cells[name] = cell("CHECK" if value else "EMPTY_OK", "保留/系统/业务字段；需与来源或上传阶段人工核对", value)
        cells["外包装图片"] = cell("EMPTY_OK" if not vals["外包装图片"] else ("PASS" if url(vals["外包装图片"]) else "BLOCK"), "可空；填写时必须为 URL", vals["外包装图片"])
        for name in ("SKU分类数量", "净含量", "SKU分类总数量", "总净含量"):
            value = vals[name]
            if not value:
                cells[name] = cell("EMPTY_OK", "可空；如填写必须为非负数且与相邻单位联动", value)
            else:
                try:
                    cells[name] = cell("PASS" if float(value) >= 0 else "BLOCK", "必须为非负数，并与相邻单位联动", value)
                except ValueError:
                    cells[name] = cell("BLOCK", "如填写必须为数值", value)
        ok, _ = json_list(vals["产品属性"])
        cells["产品属性"] = cell("PASS" if ok else "BLOCK", "AA 必须为空或 JSON 数组", vals["产品属性"])
        ok, _ = json_list(vals["SPU属性"])
        cells["SPU属性"] = cell("PASS" if ok else "BLOCK", "AB 必须为空或 JSON 数组", vals["SPU属性"])
        skc_ok, skc = json_list(vals["SKC属性"])
        skc_valid = skc_ok and bool(skc) and s(skc[0].get("extCode")) == d_value and s(skc[0].get("previewImgUrls")) == vals["预览图"]
        cells["SKC属性"] = cell("PASS" if skc_valid else "BLOCK", "AC JSON：extCode=D 且 previewImgUrls=J；此项不证明视觉正确", vals["SKC属性"])
        sku_ok, sku_attrs = json_list(vals["SKU属性"])
        sku_valid = sku_ok and len(sku_attrs) == 2 and s(sku_attrs[0].get("parentSpecName")) == vals["变种属性名称一"] and s(sku_attrs[0].get("specName")) == g_value and s(sku_attrs[1].get("parentSpecName")) == "型号" and s(sku_attrs[1].get("specName")) == model
        cells["SKU属性"] = cell("PASS" if sku_valid else "BLOCK", "AD JSON 必须镜像 F/G/H/I", vals["SKU属性"])
        video = vals["视频Url"]
        if video:
            cells["视频Url"] = cell("PASS" if url(video) and "temu" not in video.lower() else "BLOCK", "视频必须是非 Temu 的 http(s) 代理/OSS URL", video)
        else:
            cells["视频Url"] = cell("EMPTY_OK" if d_value[:4] in special_allowed_blank_video else "BLOCK", "仅 L048/L087 已知无同组素材可空；其余必须补同 L0xx 视频", video)
        # Never allow a hidden historical extCode into any cell of a pre-upload file.
        for name, value in vals.items():
            if re.search(r"L\d{3}0605\d+", value):
                cells[name] = cell("BLOCK", "预上传文件禁止遗留 L0xx0605xx", value)
        record = {"row": row_number, "D": d_value, "G": g_value, "SKU": sku, "title": vals["产品标题"], "cells": cells}
        groups[d_value].append(record)
        all_rows.append(record)

    # Same-D invariants are written onto the responsible cells, rather than hidden in a summary.
    for d_value, rows in groups.items():
        for field in ("产品标题", "轮播图", "产品素材图"):
            if len({row["cells"][field]["value"] for row in rows}) != 1:
                for row in rows:
                    row["cells"][field]["status"] = "BLOCK"
                    row["cells"][field]["rule"] += "；同 D 不一致"
        if len({row["G"] for row in rows}) != len(rows):
            for row in rows:
                row["cells"]["变种属性值一"]["status"] = "BLOCK"
                row["cells"]["变种属性值一"]["rule"] += "；同 D 的 G 重复"
        if len({row["cells"]["预览图"]["value"] for row in rows}) != len(rows):
            for row in rows:
                row["cells"]["预览图"]["status"] = "BLOCK"
                row["cells"]["预览图"]["rule"] += "；同 D 不同变体共用 J"

    statuses = Counter(c["status"] for row in all_rows for c in row["cells"].values())
    blocked_rows = [row for row in all_rows if any(c["status"] == "BLOCK" for c in row["cells"].values())]
    report = {
        "workbook": str(WORKBOOK), "scope": "all effective rows x all 54 columns", "rows": len(all_rows), "columns": len(headers),
        "cell_status_counts": dict(statuses), "blocked_rows": len(blocked_rows), "blocked_d": sorted({row["D"] for row in blocked_rows}),
        "records": all_rows,
        "meaning": {"PASS": "rule passed", "CHECK": "field retained; needs human/source semantic check", "EMPTY_OK": "empty allowed by contract", "BLOCK": "release blocker"},
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    render(headers, groups, report)
    print(json.dumps({"report": str(REPORT), "html": str(HTML), "rows": len(all_rows), "columns": len(headers), "cell_status_counts": dict(statuses), "blocked_rows": len(blocked_rows), "blocked_d": report["blocked_d"]}, ensure_ascii=False, indent=2))


def render(headers: list[str], groups: dict[str, list[dict]], report: dict) -> None:
    group_html = []
    for d_value, rows in groups.items():
        hard = any(any(c["status"] == "BLOCK" for c in row["cells"].values()) for row in rows)
        rendered_rows = []
        for row in rows:
            cells = []
            for header in headers:
                data = row["cells"][header]
                value = html.escape(data["value"], quote=True)
                rule = html.escape(data["rule"], quote=True)
                shown = value if len(value) <= 120 else value[:117] + "..."
                cells.append(f"<td class='{data['status']}' title='{rule}'><b>{data['status']}</b><span>{shown or '∅'}</span><em>{rule}</em></td>")
            rendered_rows.append(f"<details open><summary>行 {row['row']} · G={html.escape(row['G'])} · SKU={html.escape(row['SKU'])}</summary><div class='scroll'><table><thead><tr>{''.join(f'<th>{html.escape(h)}</th>' for h in headers)}</tr></thead><tbody><tr>{''.join(cells)}</tr></tbody></table></div></details>")
        group_html.append(f"<section class='group {'bad' if hard else ''}' data-search='{html.escape((d_value+' '+ ' '.join(r['G']+' '+r['SKU'] for r in rows)).lower())}'><h2>{html.escape(d_value)} <small>{html.escape(rows[0]['title'][-28:])}</small></h2>{''.join(rendered_rows)}</section>")
    c = report["cell_status_counts"]
    doc = f"""<!doctype html><html lang='zh-CN'><meta charset='utf-8'><title>YeahF 398D 每格全量审查</title><style>body{{margin:0;background:#f3f5f7;color:#1a2631;font:13px Arial,'Microsoft YaHei',sans-serif}}.top{{position:sticky;top:0;z-index:9;padding:13px 17px;background:#fff;border-bottom:2px solid #c8d5df}}h1{{margin:0 0 5px;font-size:20px}}p{{color:#526473}}input{{width:min(700px,95vw);padding:8px;border:1px solid #adbcca;border-radius:4px}}main{{max-width:2400px;margin:auto;padding:12px}}.note{{padding:10px;background:#fff1da;border:1px solid #e4bd78;border-radius:5px;line-height:1.5}}.stats{{display:flex;gap:8px;flex-wrap:wrap;margin:11px 0}}.stat{{background:#fff;padding:6px 9px;border:1px solid #cbd6df;border-radius:4px}}.group{{background:#fff;margin:11px 0;padding:10px;border:1px solid #cfd9e0;border-radius:5px}}.group.bad{{border:2px solid #d45c5c;background:#fff8f8}}h2{{font-size:16px;margin:0 0 7px;color:#29455e}}h2 small{{font-weight:normal;color:#72818e}}details{{border-top:1px solid #e0e7ec;padding:7px 0}}summary{{cursor:pointer;font-weight:bold;color:#40596e}}.scroll{{overflow:auto;margin-top:7px}}table{{border-collapse:collapse;min-width:6200px;width:100%;table-layout:fixed}}th{{position:sticky;top:0;background:#eef3f6;color:#263f53}}th,td{{border:1px solid #d7e1e7;padding:5px;vertical-align:top;width:145px;word-break:break-word}}td b{{display:block;font-size:11px;margin-bottom:3px}}td span{{display:block;max-height:44px;overflow:hidden}}td em{{display:block;font-style:normal;font-size:10px;color:#607180;margin-top:4px;max-height:31px;overflow:hidden}}.PASS{{background:#eff9f1}}.CHECK{{background:#fff8e7}}.EMPTY_OK{{background:#f2f5f7}}.BLOCK{{background:#ffe7e7;color:#8e2525}}.hidden{{display:none}}</style><body><div class='top'><h1>YeahF 398D · 每格全量审查</h1><p>逐行、逐列、逐格显示原值与规则。黄色是“需人工语义比对”，不等于通过；红色是硬阻断。</p><input id='q' placeholder='搜索 D / G / SKU'></div><main><div class='note'><b>范围：</b>{report['rows']} 行 × {report['columns']} 列 = {report['rows'] * report['columns']} 个格子。此页不修改 Excel；它把每一格的结构/联动规则展开，J 的视觉正确性仍必须在人审中确认。</div><div class='stats'><span class='stat'>PASS {c.get('PASS',0)}</span><span class='stat'>CHECK {c.get('CHECK',0)}</span><span class='stat'>EMPTY_OK {c.get('EMPTY_OK',0)}</span><span class='stat'>BLOCK {c.get('BLOCK',0)}</span><span class='stat'>阻断行 {report['blocked_rows']}</span></div>{''.join(group_html)}</main><script>q.oninput=()=>{{let v=q.value.toLowerCase();document.querySelectorAll('.group').forEach(x=>x.classList.toggle('hidden',!x.dataset.search.includes(v)))}}</script></body></html>"""
    HTML.write_text(doc, encoding="utf-8")


if __name__ == "__main__":
    main()
