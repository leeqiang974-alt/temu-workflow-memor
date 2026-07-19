"""Render the current 398-D T1 badge candidates locally and build a review page.

This script is deliberately local-only: it does not upload images and does not
change a workbook.  Every rendered candidate is bound to its source hashes.
"""
from __future__ import annotations

import hashlib
import html
import json
import shutil
import sys
from functools import lru_cache
from pathlib import Path
from urllib.parse import quote

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont, ImageOps, ImageStat, UnidentifiedImageError
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Administrator\Documents\temu自动化")
OUT = ROOT / r"outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718"
WORKBOOK = OUT / "YeahF_398D_按L0xx顺位归组_待J_T1回填_20260719.xlsx"
TARGET = OUT / "t1_badged_local_review_20260719"
IMAGES = TARGET / "images"
BASE_IMAGES = TARGET / "base_images"
MANIFEST = TARGET / "YeahF_398D_T1_角标本地候选_manifest_20260719.json"
REVIEW = TARGET / "YeahF_398D_T1_角标全局审查_20260719.html"
FONT = Path(r"C:\Windows\Fonts\arialbd.ttf")
SKU_ROOT = Path(r"E:\jit制图")

sys.path.insert(0, str(ROOT / "scripts"))
from build_yeahf_400d_refill_20260718_tj_audit import candidate_index  # noqa: E402


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}
BLOCKED_TOKENS = ("九宫格", "9grid", "grid", "拼图", "collage", "output", "out", "背景素材")
OPERATIONAL_DIRS = {"新", "新制作", "新建文件夹", "export", "exports", "result", "results"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def review_url(path: Path, d_value: str) -> str:
    """Return a same-origin URL, hydrating assets outside the review server root."""
    resolved = path.resolve()
    try:
        relative = resolved.relative_to(OUT.resolve())
    except ValueError:
        suffix = resolved.suffix.lower() if resolved.suffix.lower() in IMAGE_SUFFIXES else ".png"
        hydrated = BASE_IMAGES / d_value[:4] / f"{d_value}{suffix}"
        hydrated.parent.mkdir(parents=True, exist_ok=True)
        if not hydrated.is_file() or sha256(hydrated) != sha256(resolved):
            shutil.copy2(resolved, hydrated)
        relative = hydrated.relative_to(OUT)
    return "/" + quote(relative.as_posix(), safe="/")


def has_real_background(path: Path) -> bool:
    try:
        with Image.open(path) as image:
            rgba = image.convert("RGBA")
            rgba.thumbnail((160, 160))
            pixels = list(rgba.getdata())
    except (OSError, UnidentifiedImageError):
        return False
    if not pixels:
        return False
    opaque = [pixel for pixel in pixels if pixel[3] > 245]
    if len(opaque) < len(pixels) * 0.96:
        return False
    nonwhite = sum(1 for red, green, blue, _ in opaque if min(red, green, blue) < 225)
    return nonwhite / max(1, len(opaque)) >= 0.30


def allowed_file(path: Path) -> bool:
    lower = str(path).lower()
    return path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES and not any(token in lower for token in BLOCKED_TOKENS)


@lru_cache(maxsize=None)
def inset_pool(prefix: str) -> tuple[Path, ...]:
    """Resolve one variant level only; use sku-level files if no variant qualifies."""
    folder = SKU_ROOT / prefix / "sku"
    if not folder.is_dir():
        return ()
    variant_files: list[Path] = []
    for directory in folder.iterdir():
        if not directory.is_dir():
            continue
        name = directory.name.strip().lower()
        if name in OPERATIONAL_DIRS or any(token in name for token in BLOCKED_TOKENS):
            continue
        variant_files.extend(path for path in directory.iterdir() if allowed_file(path) and has_real_background(path))
    candidates = variant_files
    if not candidates:
        candidates = [path for path in folder.iterdir() if allowed_file(path) and has_real_background(path)]
    return tuple(sorted(candidates, key=lambda path: str(path).lower()))


def choose_inset(d_value: str) -> Path:
    pool = inset_pool(d_value[:4])
    if not pool:
        raise LookupError(f"{d_value[:4]} has no qualified sku inset")
    position = int(hashlib.sha256(d_value.encode("utf-8")).hexdigest(), 16) % len(pool)
    return pool[position]


def corner_score(image: Image.Image, corner: str, box_size: int, margin: int) -> float:
    width, height = image.size
    x = margin if "left" in corner else width - margin - box_size
    y = margin if "top" in corner else height - margin - box_size
    pad = max(3, box_size // 8)
    crop = image.crop((max(0, x - pad), max(0, y - pad), min(width, x + box_size + pad), min(height, y + box_size + pad)))
    gray = crop.convert("L").resize((96, 96))
    edges = gray.filter(ImageFilter.FIND_EDGES)
    variance = ImageStat.Stat(gray).var[0]
    edge_mean = ImageStat.Stat(edges).mean[0]
    # Busy/high-contrast regions are more likely to contain product, people or text.
    return variance + edge_mean * 12


def choose_corner(image: Image.Image) -> str:
    size = round(min(image.size) * 0.20)
    margin = round(min(image.size) * 0.04)
    corners = ("top-left", "top-right", "bottom-left", "bottom-right")
    return min(corners, key=lambda corner: corner_score(image, corner, size, margin))


def render(d_value: str, base_path: Path, inset_path: Path) -> tuple[Path, str]:
    target = IMAGES / d_value[:4] / f"{d_value}.png"
    target.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(base_path) as source:
        canvas = source.convert("RGBA")
    corner = choose_corner(canvas.convert("RGB"))
    width, height = canvas.size
    size = round(min(width, height) * 0.20)
    margin = round(min(width, height) * 0.04)
    with Image.open(inset_path) as source:
        inset = ImageOps.fit(source.convert("RGBA"), (size, size), method=Image.Resampling.LANCZOS)
    mask = Image.new("L", (size, size), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, size - 1, size - 1), fill=255)
    x = margin if "left" in corner else width - margin - size
    y = margin if "top" in corner else height - margin - size
    layer = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    layer.paste(inset, (x, y), mask)
    draw = ImageDraw.Draw(layer)
    draw.ellipse((x, y, x + size - 1, y + size - 1), outline="white", width=max(3, size // 75))
    # User-reviewed correction: the former ~10 px label was too small at the
    # 20% circular inset size. Increase it one restrained step while keeping it
    # attached to the lower rim and inside the circle width.
    font = ImageFont.truetype(str(FONT), max(12, size // 13))
    label = "THIS IS THE PRODUCT"
    box = draw.textbbox((0, 0), label, font=font, stroke_width=1)
    text_width, text_height = box[2] - box[0], box[3] - box[1]
    text_x = x + (size - text_width) // 2
    text_y = y + size - text_height // 2
    draw.text((text_x, text_y), label, font=font, fill="white", stroke_width=max(1, size // 180), stroke_fill="#101010")
    Image.alpha_composite(canvas, layer).convert("RGB").save(target, "PNG", optimize=True)
    return target, corner


def exact_d_values() -> list[str]:
    ws = load_workbook(WORKBOOK, read_only=True, data_only=False).active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    d_index = headers.index("产品货号")
    return sorted({str(row[d_index] or "").strip() for row in ws.iter_rows(min_row=2, values_only=True) if row[d_index]})


def build_review(records: list[dict]) -> None:
    cards = []
    for record in records:
        d_value = record["D"]
        cards.append(f'''<article class="card" data-search="{html.escape(d_value)} {d_value[:4]}" data-d="{html.escape(d_value)}">
<div class="pics"><figure><img loading="lazy" src="{html.escape(review_url(Path(record['base_path']), d_value))}"><figcaption>无角标 T1</figcaption></figure><figure><img loading="lazy" src="{html.escape(review_url(Path(record['rendered_path']), d_value))}"><figcaption>本地角标候选</figcaption></figure></div>
<h2>{html.escape(d_value)} <small>{html.escape(record['provider'])}</small></h2><p>角标源：{html.escape(record['inset_path'])}<br>位置：{html.escape(record['corner'])}</p>
<div class="decision"><select><option value="pending">待审</option><option value="approve">通过</option><option value="redo">重做</option><option value="reject">丢弃D</option></select><textarea placeholder="审核意见"></textarea></div></article>''')
    css = """*{box-sizing:border-box}body{margin:0;background:#eef2f5;font:13px Arial,'Microsoft YaHei';color:#1b2733}.top{position:sticky;top:0;z-index:4;background:#fff;border-bottom:2px solid #6d28d9;padding:12px 16px}.top h1{font-size:19px;margin:0 0 5px}.top p{margin:4px 0;color:#526474}.controls{display:flex;gap:8px;flex-wrap:wrap}input,button,select,textarea{font:inherit;padding:7px;border:1px solid #b7c5d1;border-radius:4px;background:#fff}input{width:min(480px,80vw)}main{padding:14px;display:grid;grid-template-columns:repeat(auto-fill,minmax(410px,1fr));gap:11px}.card{background:#fff;border:1px solid #ccd7e0;border-radius:6px;padding:9px}.pics{display:grid;grid-template-columns:1fr 1fr;gap:7px}figure{margin:0}img{display:block;width:100%;aspect-ratio:1;object-fit:contain;background:#e8edf1;cursor:zoom-in}figcaption{font-size:11px;color:#62717e;padding:3px 0}h2{font-size:14px;margin:5px 0}h2 small{font-weight:normal;color:#687887}p{font-size:10px;line-height:1.35;color:#5d6d78;overflow-wrap:anywhere}.decision{display:grid;grid-template-columns:95px 1fr;gap:6px}textarea{height:42px}.hidden{display:none}#modal{display:none;position:fixed;inset:0;z-index:9;background:#000d;place-items:center;padding:20px}#modal.open{display:grid}#modal img{max-width:94vw;max-height:90vh;width:auto;height:auto}"""
    script = """const key='yeahf-398d-t1-badge-review-user-redo-v3-20260719';const state=JSON.parse(localStorage.getItem(key)||'{}');const cards=[...document.querySelectorAll('.card')];const refresh=()=>{summary.textContent='已决定 '+Object.keys(state).length+'/398'};cards.forEach(c=>{const d=c.dataset.d,s=c.querySelector('select'),t=c.querySelector('textarea');if(state[d]){s.value=state[d].decision;t.value=state[d].note||''}const save=()=>{state[d]={decision:s.value,note:t.value};localStorage.setItem(key,JSON.stringify(state));refresh()};s.onchange=save;t.oninput=save;c.querySelectorAll('img').forEach(i=>i.onclick=()=>{modal.classList.add('open');modal.querySelector('img').src=i.src})});refresh();q.oninput=()=>cards.forEach(c=>c.classList.toggle('hidden',!c.dataset.search.includes(q.value.trim().toUpperCase())));approveAll.onclick=()=>{if(!confirm('确认已全局查看398张角标候选，将全部设为通过？'))return;cards.forEach(c=>{const d=c.dataset.d,s=c.querySelector('select');s.value='approve';state[d]={decision:'approve',note:'全局视觉审查后批量通过'}});localStorage.setItem(key,JSON.stringify(state));refresh()};modal.onclick=()=>modal.classList.remove('open');exportBtn.onclick=()=>{const blob=new Blob([JSON.stringify({batch:key,exported_at:new Date().toISOString(),decisions:state},null,2)],{type:'application/json'}),a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='yeahf_398d_t1_badge_decisions_user_redo_v3_20260719.json';a.click();URL.revokeObjectURL(a.href)}"""
    REVIEW.write_text(f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>YeahF 398D T1 角标全局审查</title><style>{css}</style><body><section class="top"><h1>YeahF 398D · T1 角标全局审查（398 张）</h1><p>左边无角标，右边本地候选。尚未 OSS，尚未回填。重点检查圆标是否遮住人物、产品、手部或任务动作。</p><div class="controls"><input id="q" placeholder="搜索 D / L0xx"><button id="approveAll">全局查看后全部通过</button><button id="exportBtn">导出决定 JSON</button><span id="summary"></span></div></section><main>{''.join(cards)}</main><div id="modal"><img></div><script>{script}</script></body></html>''', encoding="utf-8")


def main() -> None:
    candidates = candidate_index()
    records = []
    missing = []
    for position, d_value in enumerate(exact_d_values(), 1):
        candidate = candidates.get(d_value)
        if not candidate:
            missing.append({"D": d_value, "reason": "missing T1 candidate"})
            continue
        base_path = Path(candidate["local_path"])
        try:
            inset_path = choose_inset(d_value)
        except LookupError as exc:
            missing.append({"D": d_value, "reason": str(exc)})
            continue
        rendered, corner = render(d_value, base_path, inset_path)
        records.append({
            "D": d_value,
            "provider": candidate.get("provider_label", ""),
            "base_path": str(base_path),
            "base_sha256": sha256(base_path),
            "inset_path": str(inset_path),
            "inset_sha256": sha256(inset_path),
            "corner": corner,
            "rendered_path": str(rendered),
            "rendered_sha256": sha256(rendered),
            "status": "local-review-only",
        })
        if position % 25 == 0:
            print(json.dumps({"processed": position, "total": 398}, ensure_ascii=False), flush=True)
    payload = {
        "workbook": str(WORKBOOK),
        "workbook_sha256": sha256(WORKBOOK),
        "expected_exact_d": 398,
        "rendered": len(records),
        "missing": missing,
        "upload_performed": False,
        "writeback_performed": False,
        "records": records,
    }
    TARGET.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    build_review(records)
    print(json.dumps({"rendered": len(records), "missing": missing, "manifest": str(MANIFEST), "review": str(REVIEW)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
