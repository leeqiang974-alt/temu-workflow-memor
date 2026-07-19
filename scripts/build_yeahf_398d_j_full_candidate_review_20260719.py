"""Build a same-D grouped, full 702-row J candidate review for the 398-D workbook.

The workbook is read-only. L095 three-grid repairs remain local candidates until
the exported human decision file approves them.
"""
from __future__ import annotations

import csv
import hashlib
import html
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Administrator\Documents\temu自动化\outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260718")
WORKBOOK = ROOT / "YeahF_398D_按L0xx顺位归组_待J_T1回填_20260719.xlsx"
REFILL = ROOT / "refill_execution_report_20260718.json"
OLD_EVIDENCE = Path(r"C:\Users\Administrator\Documents\temu自动化\outputs\yeahf_title_dedup_20260709\yeahf_400d_refill_20260716\j_strict_source_preflight_20260716\j_row_source_resolved_20260717.csv")
L095_REPORT = ROOT / "YeahF_398D_L095三格J修正_写回报告_20260718.json"
WHITE_REPAIR = ROOT / r"j_white_visibility_repair_20260719\YeahF_398D_J_白色变体156行修复_manifest_20260719.json"
OUT = ROOT / "j_full_candidate_review_20260719"
MANIFEST = OUT / "YeahF_398D_J_702行候选_manifest_20260719.json"
REVIEW = OUT / "YeahF_398D_J_按D全局审查_20260719.html"


def file_url(value: str | Path) -> str:
    path = Path(value)
    return path.resolve().as_uri() if path.is_file() else str(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_rows() -> tuple[list[dict], dict[str, list[dict]]]:
    ws = load_workbook(WORKBOOK, read_only=True, data_only=False).active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {name: headers.index(name) for name in headers}
    rows, groups = [], defaultdict(list)
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        d_value = str(row[index["产品货号"]] or "").strip()
        if not d_value:
            continue
        record = {
            "row": row_number,
            "D": d_value,
            "L0xx": d_value[:4],
            "G": str(row[index["变种属性值一"]] or "").strip(),
            "SKU": str(row[index["SKU货号"]] or "").strip(),
            "current_J": str(row[index["预览图"]] or "").strip(),
        }
        rows.append(record)
        groups[d_value].append(record)
    return rows, groups


def build_candidates(rows: list[dict], groups: dict[str, list[dict]]) -> list[dict]:
    old = list(csv.DictReader(OLD_EVIDENCE.open(encoding="utf-8-sig", newline="")))
    direct = {(r["D"], r["G"], r["SKU"]): r for r in old if r.get("status") == "ready-for-visual-audit"}
    refill = json.loads(REFILL.read_text(encoding="utf-8"))
    appended = {item["target_d"]: item["source_d"] for item in refill["appended_exact_d"]}
    l095 = {(item["D"], item["G"], item["SKU"]): item for item in json.loads(L095_REPORT.read_text(encoding="utf-8"))["records"]}
    white_repairs = {}
    if WHITE_REPAIR.is_file():
        white_repairs = {(item["D"], item["G"], item["SKU"]): item for item in json.loads(WHITE_REPAIR.read_text(encoding="utf-8"))["records"]}
    result = []
    for row in rows:
        candidate = {**row, "status": "blocked", "match_mode": "", "source_D": row["D"], "source_G": row["G"], "source_SKU": row["SKU"], "source_path": "", "candidate_J": row["current_J"], "candidate_kind": "current-workbook-url", "warning": ""}
        identity = (row["D"], row["G"], row["SKU"])
        candidate["decision_key"] = "|".join(identity)
        if identity in white_repairs:
            repair = white_repairs[identity]
            local = Path(repair["local_candidate"])
            candidate.update({
                "status": "human-review-required",
                "match_mode": "white-variant-authoritative-alpha-two-color-macaron-repair",
                "source_path": repair["source_path"],
                "candidate_J": str(local),
                "candidate_kind": "local-white-visibility-repair",
                "candidate_sha256": repair["candidate_sha256"],
                "warning": "旧 J 曾错误删除白色产品像素；本地候选保留原 alpha 并增强边界，需重新审批。",
            })
        elif identity in l095:
            repair = l095[identity]
            local = Path(repair["local"])
            candidate.update({
                "status": "human-review-required",
                "match_mode": "L095-01-exact-three-grid-repair",
                "source_path": repair["source"],
                "candidate_J": str(local),
                "candidate_kind": "local-unapproved-candidate",
                "candidate_sha256": sha256(local),
                "warning": "L095-01 三格修正候选；未获用户通过前不回填。",
            })
        elif (row["D"], row["G"], row["SKU"]) in direct:
            evidence = direct[(row["D"], row["G"], row["SKU"])]
            candidate.update({"status": "evidence-ready", "match_mode": evidence["match_mode"], "source_path": evidence["sku_source"]})
        elif row["D"] in appended:
            source_d = appended[row["D"]]
            position = groups[row["D"]].index(row)
            source_rows = groups.get(source_d, [])
            if position < len(source_rows):
                source_row = source_rows[position]
                evidence = direct.get((source_d, source_row["G"], source_row["SKU"]))
                same_claim = (row["G"], row["SKU"]) == (source_row["G"], source_row["SKU"])
                same_j = row["current_J"] == source_row["current_J"]
                if evidence and same_claim and same_j:
                    candidate.update({
                        "status": "human-review-required",
                        "match_mode": "append-exact-source-row-G-SKU-J-equality",
                        "source_D": source_d,
                        "source_G": source_row["G"],
                        "source_SKU": source_row["SKU"],
                        "source_path": evidence["sku_source"],
                        "warning": "新增 D：G、SKU、J 与有证据的来源 D 同序行完全相等，仍需人工确认图像。",
                    })
                else:
                    candidate["warning"] = f"新增 D 来源校验失败：source={source_d}, same_claim={same_claim}, same_j={same_j}, evidence={bool(evidence)}"
            else:
                candidate["warning"] = f"新增 D 来源 {source_d} 行数不足。"
        else:
            candidate["warning"] = "没有精确 D+G+SKU 正向来源。"
        result.append(candidate)
    return result


def image_tag(src: str, label: str) -> str:
    if not src:
        return f"<div class='none'>{html.escape(label)}缺失</div>"
    return f"<button class='zoom' data-src='{html.escape(file_url(src), quote=True)}' data-label='{html.escape(label, quote=True)}'><img loading='lazy' src='{html.escape(file_url(src), quote=True)}'></button>"


def build_review(records: list[dict]) -> None:
    groups = defaultdict(list)
    for record in records:
        groups[record["D"]].append(record)
    sections = []
    ordered_groups = sorted(
        groups.items(),
        key=lambda pair: (
            not any(item["status"] in {"human-review-required", "blocked"} for item in pair[1]),
            pair[0],
        ),
    )
    for d_value, items in ordered_groups:
        cards = []
        blocked = any(item["status"] == "blocked" for item in items)
        for item in items:
            row_key = item["decision_key"]
            cls = "row blocked" if item["status"] == "blocked" else ("row review" if item["status"] == "human-review-required" else "row")
            cards.append(f'''<article class="{cls}" data-key="{row_key}"><header><b>行 {item['row']} · {html.escape(item['SKU'])}</b><span>{html.escape(item['status'])}</span></header>
<div class="pics"><figure>{image_tag(item['source_path'], '正向 SKU 证据')}<figcaption>源 {html.escape(item['source_D'])} · {html.escape(item['source_G'])}</figcaption></figure><figure>{image_tag(item['candidate_J'], '候选 J')}<figcaption>{html.escape(item['candidate_kind'])}</figcaption></figure></div>
<p>G：{html.escape(item['G'])}<br>SKU：{html.escape(item['SKU'])}<br>匹配：{html.escape(item['match_mode'])}<br>{html.escape(item['warning'])}</p>
<div class="decision"><select><option value="pending">待审</option><option value="approve">通过</option><option value="redo">重做</option><option value="reject">丢弃D</option></select><textarea placeholder="审核意见"></textarea></div></article>'''.replace('data-key="'+row_key+'"', 'data-key="'+row_key+'" data-status="'+item['status']+'"'))
        search_text = d_value + " " + d_value[:4] + " " + " ".join(x["G"] + " " + x["SKU"] + " " + x["status"] for x in items)
        sections.append(f'''<section class="group{' blocked' if blocked else ''}" data-search="{html.escape(search_text.lower())}"><h2>{html.escape(d_value)} <small>{len(items)} 行同 D 并排</small></h2><div class="rows">{''.join(cards)}</div></section>''')
    css = """*{box-sizing:border-box}body{margin:0;background:#eef2f5;color:#1b2733;font:13px Arial,'Microsoft YaHei'}.top{position:sticky;top:0;z-index:4;background:#fff;border-bottom:2px solid #6d28d9;padding:12px 16px}.top h1{font-size:19px;margin:0 0 5px}.top p{margin:4px 0;color:#526474}.controls{display:flex;gap:8px;flex-wrap:wrap}input,button,select,textarea{font:inherit;padding:7px;border:1px solid #b7c5d1;border-radius:4px;background:#fff}input{width:min(520px,80vw)}main{padding:14px;max-width:2200px;margin:auto}.group{background:#fff;border:1px solid #cbd7e0;border-radius:6px;margin:0 0 12px;padding:9px}.group>h2{font-size:15px;margin:0 0 7px}.group>h2 small{font-weight:normal;color:#657583}.rows{display:grid;grid-template-columns:repeat(auto-fit,minmax(390px,1fr));gap:8px}.row{border:1px solid #d6e0e7;border-radius:5px;padding:7px}.row.review{border-color:#d5a13c;background:#fffaf0}.row.blocked{border:2px solid #d34d4d;background:#fff5f5}.row header{display:flex;justify-content:space-between;gap:8px}.pics{display:grid;grid-template-columns:1fr 1fr;gap:6px;margin-top:5px}figure{margin:0}img{display:block;width:100%;aspect-ratio:1;object-fit:contain;background:#e8edf1;cursor:zoom-in}figcaption{font-size:10px;color:#687784;padding:3px 0}.none{display:grid;place-items:center;aspect-ratio:1;background:#f2f4f6;color:#a33}p{font-size:10px;line-height:1.4;color:#596a77}.decision{display:grid;grid-template-columns:90px 1fr;gap:5px}textarea{height:40px}.hidden{display:none}#modal{display:none;position:fixed;inset:0;z-index:9;background:#000d;place-items:center;padding:20px}#modal.open{display:grid}#modal img{max-width:94vw;max-height:90vh;width:auto;height:auto}"""
    script = """const key='yeahf-398d-j-702-review-sorted-v1-20260719';const state=JSON.parse(localStorage.getItem(key)||'{}');const rows=[...document.querySelectorAll('.row')],groups=[...document.querySelectorAll('.group')];const refresh=()=>summary.textContent='已决定 '+Object.keys(state).length+'/702';rows.forEach(c=>{const k=c.dataset.key,s=c.querySelector('select'),t=c.querySelector('textarea');if(state[k]){s.value=state[k].decision;t.value=state[k].note||''}const save=()=>{state[k]={decision:s.value,note:t.value};localStorage.setItem(key,JSON.stringify(state));refresh()};s.onchange=save;t.oninput=save;c.querySelectorAll('.zoom').forEach(b=>b.onclick=()=>{modal.classList.add('open');modal.querySelector('img').src=b.dataset.src})});refresh();q.oninput=()=>groups.forEach(c=>c.classList.toggle('hidden',!c.dataset.search.includes(q.value.trim().toLowerCase())));onlyRisk.onclick=()=>groups.forEach(g=>g.classList.toggle('hidden',!g.querySelector('.row[data-status="human-review-required"],.row[data-status="blocked"]')));showAll.onclick=()=>groups.forEach(g=>g.classList.remove('hidden'));approveEvidence.onclick=()=>{const count=rows.filter(c=>c.dataset.status==='evidence-ready').length;if(!confirm('确认已查看旧证据行，将 '+count+' 行 evidence-ready 设为通过？'))return;rows.filter(c=>c.dataset.status==='evidence-ready').forEach(c=>{const k=c.dataset.key,s=c.querySelector('select');s.value='approve';state[k]={decision:'approve',note:'已有精确行级证据，全局审查后批量通过'}});localStorage.setItem(key,JSON.stringify(state));refresh()};modal.onclick=()=>modal.classList.remove('open');exportBtn.onclick=()=>{const a=document.createElement('a'),blob=new Blob([JSON.stringify({batch:key,exported_at:new Date().toISOString(),decisions:state},null,2)],{type:'application/json'});a.href=URL.createObjectURL(blob);a.download='yeahf_398d_j_decisions_sorted_v1_20260719.json';a.click();URL.revokeObjectURL(a.href)}"""
    ready = sum(item["status"] == "evidence-ready" for item in records)
    review = sum(item["status"] == "human-review-required" for item in records)
    blocked = sum(item["status"] == "blocked" for item in records)
    REVIEW.write_text(f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>YeahF 398D J 按D全局审查</title><style>{css}</style><body><section class="top"><h1>YeahF 398D · J 702 行按 D 全局审查</h1><p>同 D 全部并排。已有证据 {ready} 行，需人工确认 {review} 行，阻断 {blocked} 行。白色变体修复和 L095 三格只显示本地候选，未批准前不回填。</p><div class="controls"><input id="q" placeholder="搜索 D / L0xx / G / SKU"><button id="onlyRisk">只看{review}行待确认</button><button id="showAll">显示全部</button><button id="approveEvidence">通过{ready}行已有证据</button><button id="exportBtn">导出决定 JSON</button><span id="summary"></span></div></section><main>{''.join(sections)}</main><div id="modal"><img></div><script>{script}</script></body></html>''', encoding="utf-8")


def main() -> None:
    rows, groups = load_rows()
    records = build_candidates(rows, groups)
    OUT.mkdir(parents=True, exist_ok=True)
    payload = {
        "workbook": str(WORKBOOK),
        "workbook_sha256": sha256(WORKBOOK),
        "effective_rows": len(records),
        "exact_d": len(groups),
        "status_counts": {status: sum(item["status"] == status for item in records) for status in sorted({item["status"] for item in records})},
        "upload_performed": False,
        "writeback_performed": False,
        "records": records,
    }
    MANIFEST.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    build_review(records)
    print(json.dumps({"manifest": str(MANIFEST), "review": str(REVIEW), "counts": payload["status_counts"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
