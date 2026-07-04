from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import time
import uuid
from collections import Counter, defaultdict
from pathlib import Path

import oss2
from openpyxl import load_workbook
from PIL import Image


COMFYUI_BASE = Path(r"C:\Users\Administrator\Documents\Codex\2026-06-08\comfyui")
if str(COMFYUI_BASE / "work") not in sys.path:
    sys.path.insert(0, str(COMFYUI_BASE / "work"))

import generate_apply_ali_tfirst_0608 as ali_t


TARGET_SIZE = 800
TARGET_MAX_BYTES = 150 * 1024
OSS_PREFIX = "temu-jit/dxxmall-0616-2/197x3-passed-t-writeback"
URL_RE = re.compile(r"https?://.*?(?=(?:[,;，；]?\s*https?://)|$)", re.S)
SIZE_RE = re.compile(r"(尺寸|尺码|size|cm|inch|length|height|width|宽|长|高)", re.I)


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def save_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def headers(ws) -> list[str]:
    return [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]


def col(header_values: list[str], name: str) -> int:
    if name not in header_values:
        raise RuntimeError(f"missing required column: {name}")
    return header_values.index(name) + 1


def split_urls(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    urls: list[str] = []
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line:
            continue
        matches = [match.group(0).strip() for match in URL_RE.finditer(line)]
        if matches:
            urls.extend(matches)
        else:
            urls.append(line)
    return urls


def dedupe(values: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = str(value or "").strip()
        if value and value not in seen:
            result.append(value)
            seen.add(value)
    return result


def resize_compress(input_path: Path, output_path: Path) -> dict:
    image = Image.open(input_path).convert("RGB").resize((TARGET_SIZE, TARGET_SIZE), Image.Resampling.LANCZOS)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    from io import BytesIO

    best = b""
    best_quality = 88
    for quality in range(88, 34, -3):
        buffer = BytesIO()
        image.save(buffer, format="JPEG", quality=quality, optimize=True, progressive=True, subsampling=1)
        best = buffer.getvalue()
        best_quality = quality
        if len(best) <= TARGET_MAX_BYTES:
            break
    output_path.write_bytes(best)
    return {"width": TARGET_SIZE, "height": TARGET_SIZE, "bytes": len(best), "quality": best_quality}


def load_manifest(path: Path) -> dict:
    if path.exists():
        return load_json(path)
    return {"uploaded": {}}


def upload_once(bucket: oss2.Bucket, manifest_path: Path, manifest: dict, d_value: str, set_no: int, path: Path) -> str:
    key = f"{d_value}|set{set_no}|{path.resolve()}"
    uploaded = manifest.setdefault("uploaded", {})
    if key in uploaded:
        return uploaded[key]
    object_key = f"{OSS_PREFIX}/{time.strftime('%Y%m%d')}/{uuid.uuid4().hex}_{d_value}_set{set_no}_tfirst_800.jpg"
    bucket.put_object_from_file(object_key, str(path), headers={"Content-Type": "image/jpeg"})
    endpoint = bucket.endpoint.replace("https://", "").replace("http://", "")
    url = f"https://{bucket.bucket_name}.{endpoint}/{object_key}"
    uploaded[key] = url
    save_json(manifest_path, manifest)
    return url


def load_pool(path: Path) -> dict[str, list[dict]]:
    records = load_json(path)
    by_d: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        d_value = str(record.get("original_d") or record.get("d") or "").strip()
        if not d_value:
            continue
        local_path = Path(str(record.get("local_path") or ""))
        if not local_path.exists():
            raise FileNotFoundError(f"{d_value}: missing candidate image {local_path}")
        record["set_no"] = int(record.get("set_no") or record.get("slot") or 0)
        by_d[d_value].append(record)
    for d_value, items in by_d.items():
        items.sort(key=lambda item: item["set_no"])
        if len(items) != 3:
            raise RuntimeError(f"{d_value}: expected 3 candidates, got {len(items)}")
    return by_d


def chosen_order(d_value: str, prefix_index: int) -> list[int]:
    first = (prefix_index % 3) + 1
    return [first] + [value for value in (1, 2, 3) if value != first]


def fixed_chosen_order(first: int) -> list[int]:
    if first not in (1, 2, 3):
        raise ValueError(f"fixed first set must be 1, 2, or 3; got {first}")
    return [first] + [value for value in (1, 2, 3) if value != first]


def build_urls(old_urls: list[str], candidate_urls: list[str], replace_first: bool) -> tuple[list[str], str]:
    size_url = old_urls[3] if len(old_urls) >= 4 else ""
    if replace_first:
        new_first = candidate_urls[0]
        rest = [url for idx, url in enumerate(old_urls) if idx != 0 and url != new_first and url != size_url]
        if size_url:
            urls = [new_first] + rest[:2] + [size_url] + rest[2:]
        else:
            urls = [new_first] + rest
        mode = "replace_t1"
    else:
        urls = list(old_urls)
        for candidate in candidate_urls:
            if len(urls) >= 6:
                break
            urls.append(candidate)
        mode = "append_until_t6"
    return dedupe(urls)[:10], mode


def rotate_secondary_slots(final_by_d: dict[str, list[str]]) -> list[dict]:
    changes: list[dict] = []
    by_prefix: dict[str, list[str]] = defaultdict(list)
    for d_value in final_by_d:
        by_prefix[d_value[:4]].append(d_value)
    for prefix, d_values in by_prefix.items():
        d_values = sorted(d_values)
        if len(d_values) < 2:
            continue
        for slot in (1, 2, 4):
            candidates = []
            for d_value in d_values:
                urls = final_by_d[d_value]
                if len(urls) <= slot:
                    continue
                value = urls[slot]
                if value == urls[0] or (len(urls) >= 4 and value == urls[3]) or SIZE_RE.search(value):
                    continue
                candidates.append((d_value, value))
            if len(candidates) < 2:
                continue
            rotated = [value for _, value in candidates[-1:]] + [value for _, value in candidates[:-1]]
            for (d_value, old_value), new_value in zip(candidates, rotated):
                urls = final_by_d[d_value]
                if new_value in urls or old_value == new_value:
                    continue
                urls[slot] = new_value
                changes.append({"prefix": prefix, "d": d_value, "slot": slot + 1, "old": old_value, "new": new_value})
    return changes


def write_audit_html(path: Path, report: dict) -> None:
    rows = []
    for item in report["changes"]:
        imgs = "".join(
            f'<div><div class="cap">T{i + 1}</div><img src="{html.escape(url)}"></div>'
            for i, url in enumerate(item["final_urls"][:6])
        )
        rows.append(
            "<section>"
            f"<h3>{html.escape(item['d'])} <span>{html.escape(item['mode'])}</span></h3>"
            f"<p>row(s): {html.escape(str(item['rows']))} | old_count: {item['old_count']} | final_count: {item['final_count']} | "
            f"chosen_set: {item['chosen_set']} | T4 preserved: {item['t4_preserved']}</p>"
            f"<p class='small'>old T1: {html.escape(item.get('old_t1') or '')}</p>"
            f"<p class='small'>new T1: {html.escape(item.get('new_t1') or '')}</p>"
            f"<div class='grid'>{imgs}</div>"
            "</section>"
        )
    body = "\n".join(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>196 T writeback audit</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:24px;background:#f6f7fb;color:#1f2937}}
h1{{font-size:22px}} section{{background:white;border:1px solid #d8dee9;border-radius:8px;margin:14px 0;padding:14px}}
h3{{margin:0 0 8px;font-size:18px}} h3 span{{font-size:13px;color:#2563eb;margin-left:8px}}
.small{{font-size:12px;color:#4b5563;word-break:break-all}} .grid{{display:grid;grid-template-columns:repeat(6,140px);gap:10px;align-items:start}}
img{{width:140px;height:140px;object-fit:cover;border:1px solid #e5e7eb;background:white}} .cap{{font-size:12px;font-weight:700;margin-bottom:4px}}
</style></head><body>
<h1>196 表 T 首图回填审核</h1>
<p>source: {html.escape(report['source_workbook'])}</p>
<p>output: {html.escape(report['output_workbook'])}</p>
<p>unique D: {report['unique_d']} | replaced T1: {report['replace_t1_count']} | append-only: {report['append_only_count']} | secondary rotations: {len(report['secondary_rotations'])}</p>
{body}
</body></html>""",
        encoding="utf-8",
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--pool", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--output-name", default="")
    parser.add_argument(
        "--fixed-first-set",
        type=int,
        choices=(1, 2, 3),
        default=0,
        help="Use the same passed candidate set as T1 for every D instead of rotating set1/set2/set3.",
    )
    parser.add_argument(
        "--rotate-secondary",
        action="store_true",
        help="Rotate non-first/non-fourth old T images within each L0xx prefix. Disabled by default to avoid bringing back rejected legacy images.",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    pool_path = Path(args.pool)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    output_path = out_dir / (args.output_name or f"{workbook_path.stem}_已回填197x3通过T首图_20260703.xlsx")
    manifest_path = out_dir / "t_upload_manifest.json"
    compressed_dir = out_dir / "t_first_800"

    pool = load_pool(pool_path)
    wb = load_workbook(workbook_path)
    ws = wb.active
    h = headers(ws)
    col_d = col(h, "产品货号")
    col_t = col(h, "轮播图")
    col_u = col(h, "产品素材图")

    d_rows: dict[str, list[int]] = defaultdict(list)
    original_by_d: dict[str, list[str]] = {}
    for row in range(2, ws.max_row + 1):
        d_value = str(ws.cell(row, col_d).value or "").strip()
        if not d_value:
            continue
        d_rows[d_value].append(row)
        original_by_d.setdefault(d_value, split_urls(ws.cell(row, col_t).value))

    missing_pool = sorted(set(d_rows) - set(pool))
    if missing_pool:
        raise RuntimeError(f"workbook D missing from 197x3 passed pool: {missing_pool[:20]}")

    oss_config = ali_t.ali.read_oss_config()
    bucket = oss2.Bucket(
        oss2.Auth(oss_config["access_key_id"], oss_config["access_key_secret"]),
        f"https://{oss_config['endpoint']}",
        oss_config["bucket"],
    )
    manifest = load_manifest(manifest_path)

    prefix_order: dict[str, int] = defaultdict(int)
    uploaded: dict[tuple[str, int], str] = {}
    compressed_report: dict[str, dict] = {}
    changes: list[dict] = []
    final_by_d: dict[str, list[str]] = {}

    for d_value in sorted(d_rows):
        prefix = d_value[:4]
        index = prefix_order[prefix]
        prefix_order[prefix] += 1
        order = fixed_chosen_order(args.fixed_first_set) if args.fixed_first_set else chosen_order(d_value, index)
        records_by_set = {int(record["set_no"]): record for record in pool[d_value]}
        old_urls = original_by_d[d_value]
        replace_first = len(old_urls) >= 6
        upload_order = order if not replace_first else [order[0]]
        candidate_urls: list[str] = []
        for set_no in upload_order:
            record = records_by_set[set_no]
            local_path = Path(record["local_path"])
            compressed = compressed_dir / prefix / f"{d_value}_set{set_no}_tfirst_800.jpg"
            compressed_report[f"{d_value}|set{set_no}"] = resize_compress(local_path, compressed)
            uploaded[(d_value, set_no)] = upload_once(bucket, manifest_path, manifest, d_value, set_no, compressed)
            candidate_urls.append(uploaded[(d_value, set_no)])

        final_urls, mode = build_urls(old_urls, candidate_urls, replace_first)
        final_by_d[d_value] = final_urls
        changes.append(
            {
                "d": d_value,
                "rows": d_rows[d_value],
                "mode": mode,
                "chosen_set": order[0],
                "old_count": len(old_urls),
                "final_count": len(final_urls),
                "old_t1": old_urls[0] if old_urls else "",
                "new_t1": final_urls[0] if final_urls else "",
                "old_t4": old_urls[3] if len(old_urls) >= 4 else "",
                "new_t4": final_urls[3] if len(final_urls) >= 4 else "",
                "t4_preserved": bool(len(old_urls) >= 4 and len(final_urls) >= 4 and old_urls[3] == final_urls[3]),
                "candidate_urls": candidate_urls,
                "final_urls": final_urls,
            }
        )

    secondary_rotations = rotate_secondary_slots(final_by_d) if args.rotate_secondary else []
    for item in changes:
        urls = final_by_d[item["d"]]
        item["final_urls"] = urls
        item["final_count"] = len(urls)
        item["new_t1"] = urls[0] if urls else ""
        item["new_t4"] = urls[3] if len(urls) >= 4 else ""
        item["t4_preserved"] = bool(item["old_t4"] and len(urls) >= 4 and item["old_t4"] == urls[3])

    for d_value, rows in d_rows.items():
        urls = final_by_d[d_value]
        for row in rows:
            ws.cell(row, col_t).value = "\n".join(urls)
            if len(original_by_d[d_value]) >= 6:
                ws.cell(row, col_u).value = urls[0]

    wb.save(output_path)
    wb.close()

    validation_issues = []
    for item in changes:
        urls = final_by_d[item["d"]]
        if len(urls) < 6:
            validation_issues.append({"d": item["d"], "issue": "T count below 6", "count": len(urls)})
        if len(urls) > 10:
            validation_issues.append({"d": item["d"], "issue": "T count above 10", "count": len(urls)})
        old_t4 = item["old_t4"]
        if old_t4 and (len(urls) < 4 or urls[3] != old_t4):
            validation_issues.append({"d": item["d"], "issue": "T4 not preserved", "old_t4": old_t4, "new_t4": urls[3] if len(urls) >= 4 else ""})

    report = {
        "source_workbook": str(workbook_path),
        "pool": str(pool_path),
        "output_workbook": str(output_path),
        "unique_d": len(d_rows),
        "effective_rows": sum(len(rows) for rows in d_rows.values()),
        "pool_extra_d": sorted(set(pool) - set(d_rows)),
        "replace_t1_count": sum(1 for item in changes if item["mode"] == "replace_t1"),
        "append_only_count": sum(1 for item in changes if item["mode"] == "append_until_t6"),
        "t_count_dist": dict(Counter(len(final_by_d[d]) for d in final_by_d)),
        "chosen_set_dist": dict(Counter(item["chosen_set"] for item in changes)),
        "fixed_first_set": args.fixed_first_set or None,
        "secondary_rotations": secondary_rotations,
        "compression": {
            "count": len(compressed_report),
            "max_bytes": max(item["bytes"] for item in compressed_report.values()),
            "over_target": {k: v for k, v in compressed_report.items() if v["bytes"] > TARGET_MAX_BYTES},
        },
        "validation_issues": validation_issues,
        "changes": changes,
    }
    save_json(out_dir / "writeback_report.json", report)
    write_audit_html(out_dir / "writeback_t_audit.html", report)
    print(json.dumps({k: v for k, v in report.items() if k != "changes"}, ensure_ascii=False, indent=2))
    if validation_issues:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
