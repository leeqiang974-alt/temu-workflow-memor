import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook


BAD_SUBSTRINGS = [
    "l058-extra-fixed-under145k",
    "L058_extra_fixed_800_under145k",
    "保持产品和标尺、数字、文字信息不做.jpg",
]


def split_t(value):
    if value is None:
        return []
    return [part.strip() for part in re.split(r"[\n\r]+", str(value).strip()) if part.strip()]


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    report = Path(sys.argv[3])

    shutil.copy2(src, dst)
    wb = load_workbook(dst)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    d_col = headers.index("产品货号") + 1
    t_col = headers.index("轮播图") + 1

    changes = []
    for row_idx in range(2, ws.max_row + 1):
        d_value = ws.cell(row_idx, d_col).value
        if d_value is None or str(d_value).strip() == "":
            continue
        old_t = ws.cell(row_idx, t_col).value
        parts = split_t(old_t)
        if not parts:
            continue
        kept = []
        removed = []
        for part in parts:
            if any(bad in part for bad in BAD_SUBSTRINGS):
                removed.append(part)
            else:
                kept.append(part)
        if removed:
            ws.cell(row_idx, t_col).value = "\n".join(kept)
            changes.append(
                {
                    "row": row_idx,
                    "D": str(d_value).strip(),
                    "removed_count": len(removed),
                    "removed": removed,
                    "new_t_count": len(kept),
                }
            )

    wb.save(dst)
    payload = {"source": str(src), "output": str(dst), "changes": changes, "changed_rows": len(changes)}
    report.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
