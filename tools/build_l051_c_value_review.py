import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_final_199_writeback_20260702_fix_feedback_20260702\0616-2_197_最终回传_按新197D清单从199同步删除_L091两D_硬校验修复_20260702.xlsx"
)
OUT_DIR = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_final_199_writeback_20260702_fix_feedback_20260702\l051060505_c_value_review"
)
D_VALUE = "L051060505"


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col_d = headers.index("产品货号") + 1
    col_c = 3
    col_c_name = headers[col_c - 1]
    record = None
    for row in ws.iter_rows(min_row=2, values_only=False):
        d = row[col_d - 1].value
        if d and str(d).strip() == D_VALUE:
            record = {
                "row": row[0].row,
                "D": D_VALUE,
                "column": "C",
                "header": col_c_name,
                "value": row[col_c - 1].value or "",
            }
            break
    if not record:
        raise SystemExit(f"{D_VALUE} not found")

    value = str(record["value"])
    img_urls = re.findall(r'<img\s+src="([^"]+)"\s*/?>', value, flags=re.I)
    record["images"] = [{"index": i + 1, "url": url} for i, url in enumerate(img_urls)]
    (OUT_DIR / "c_value_record.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")

    cards = []
    for item in record["images"]:
        idx = item["index"]
        url = item["url"]
        cards.append(
            f"""
            <section class="card" data-index="{idx}" data-url="{html.escape(url)}">
              <div class="top">
                <label><input type="checkbox" class="del" data-index="{idx}" data-url="{html.escape(url)}"> 删除这张</label>
                <span>第 {idx} 张</span>
              </div>
              <img src="{html.escape(url)}" loading="lazy">
              <textarea class="note" data-index="{idx}" placeholder="删除原因/备注"></textarea>
              <code>{html.escape(url)}</code>
            </section>
            """
        )

    page = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{D_VALUE} C值产品描述审核</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:0;background:#f5f5f5;color:#222}}
header{{position:sticky;top:0;background:#fff;border-bottom:1px solid #ddd;padding:14px 18px;z-index:2}}
h1{{font-size:20px;margin:0 0 8px}}
.meta{{font-size:13px;color:#555;line-height:1.5}}
.actions{{display:flex;gap:10px;margin-top:10px;flex-wrap:wrap}}
button{{border:0;background:#2563eb;color:#fff;border-radius:6px;padding:8px 12px;cursor:pointer}}
button.secondary{{background:#0f766e}}
main{{padding:18px;display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:16px}}
.card{{background:#fff;border:1px solid #ddd;border-radius:8px;padding:12px;display:flex;flex-direction:column;gap:10px}}
.top{{display:flex;justify-content:space-between;align-items:center;font-weight:700}}
.card img{{width:100%;height:auto;border:1px solid #eee;background:#fff}}
textarea{{min-height:54px;resize:vertical;border:1px solid #ccc;border-radius:6px;padding:8px;font-family:inherit}}
code{{font-size:12px;white-space:pre-wrap;word-break:break-all;background:#f7f7f7;padding:8px;border-radius:6px}}
details{{margin:0 18px 18px;background:#fff;border:1px solid #ddd;border-radius:8px;padding:12px}}
pre{{white-space:pre-wrap;word-break:break-all;max-height:260px;overflow:auto;background:#111;color:#eee;padding:12px;border-radius:6px}}
</style>
</head>
<body>
<header>
  <h1>{D_VALUE} / C列 {html.escape(str(col_c_name))}</h1>
  <div class="meta">来源表：{html.escape(str(WORKBOOK))}<br>Excel 行：{record['row']}，共 {len(img_urls)} 张图片。勾选要删除的图片后点导出反馈。</div>
  <div class="actions">
    <button onclick="exportFeedback()">导出删除反馈 JSON</button>
    <button class="secondary" onclick="copyFeedback()">复制删除反馈</button>
  </div>
</header>
<main>
{''.join(cards)}
</main>
<details open>
  <summary>原始 C 值</summary>
  <pre>{html.escape(value)}</pre>
</details>
<script>
function collect(){{
  const deletes = [];
  document.querySelectorAll('.del').forEach(cb => {{
    if (cb.checked) {{
      const idx = Number(cb.dataset.index);
      const note = document.querySelector('.note[data-index="'+idx+'"]').value || '';
      deletes.push({{index: idx, url: cb.dataset.url, note}});
    }}
  }});
  return {{
    workbook: {json.dumps(str(WORKBOOK), ensure_ascii=False)},
    D: {json.dumps(D_VALUE)},
    row: {record['row']},
    column: 'C',
    header: {json.dumps(str(col_c_name), ensure_ascii=False)},
    decision: 'delete_selected_images_from_c_value',
    deletes,
    ts: new Date().toISOString()
  }};
}}
function exportFeedback(){{
  const payload = collect();
  const blob = new Blob([JSON.stringify(payload,null,2)], {{type:'application/json'}});
  const a = document.createElement('a');
  a.href = URL.createObjectURL(blob);
  a.download = 'L051060505_C_value_delete_feedback.json';
  a.click();
}}
async function copyFeedback(){{
  await navigator.clipboard.writeText(JSON.stringify(collect(),null,2));
  alert('已复制');
}}
</script>
</body>
</html>"""
    (OUT_DIR / "index.html").write_text(page, encoding="utf-8")
    print(OUT_DIR / "index.html")
    print(json.dumps(record, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
