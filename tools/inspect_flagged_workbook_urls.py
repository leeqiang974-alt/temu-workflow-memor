import sys
from pathlib import Path

from openpyxl import load_workbook


def main():
    workbook = Path(sys.argv[1])
    wb = load_workbook(workbook, read_only=True, data_only=False)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    wanted = ["产品货号", "轮播图", "产品素材图", "预览图"]
    idx = {name: headers.index(name) + 1 for name in wanted if name in headers}
    needles = ["l058-extra-fixed-under145k", "L058_extra_fixed_800_under145k", "redo"]

    for row_idx in range(2, ws.max_row + 1):
        d_value = ws.cell(row_idx, idx["产品货号"]).value
        if not d_value:
            continue
        values = {name: str(ws.cell(row_idx, col).value or "") for name, col in idx.items()}
        combined = "\n".join(values.values())
        if any(needle in combined for needle in needles):
            print(f"\nROW {row_idx} D {str(d_value).strip()}")
            for name in ["预览图", "轮播图", "产品素材图"]:
                print(f"{name}={values.get(name, '')[:1500]}")


if __name__ == "__main__":
    main()
