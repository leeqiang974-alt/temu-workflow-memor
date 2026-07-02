from __future__ import annotations

import html
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFilter


WORKBOOK = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_final_199_writeback_20260702_fix_feedback_20260702"
    r"\0616-2_199_最终回传_已应用反馈_L042J重做_T清理_20260702.xlsx"
)
SKU_ROOT = Path(r"E:\jit制图\L042\sku")
OUT_ROOT = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_final_199_writeback_20260702_fix_feedback_20260702"
    r"\l042_sizechart_j_redo"
)
SERVED_ROOT = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_seedream_fallback_9_20260702"
    r"\l042_j_review_sizechart_redo"
)


def classify_variant(g_value: object, sku_value: object) -> str:
    text = f"{g_value or ''} {sku_value or ''}".lower()
    return "green" if any(token in text for token in ["绿", "green"]) else "black"


def choose_sizechart(color_cn: str) -> Path:
    folder = SKU_ROOT / color_cn
    files = [p for p in sorted(folder.iterdir()) if p.is_file()]
    if not files:
        raise FileNotFoundError(folder)
    # Prefer a direct first-level file with readable complete text/nails.
    def score(path: Path) -> tuple[int, int, str]:
        name = path.name
        wh_bonus = 0 if "_WH_" in name else 1
        return (wh_bonus, -path.stat().st_size, name)

    return sorted(files, key=score)[0]


def normalize_square(source: Path, out_path: Path) -> None:
    img = Image.open(source).convert("RGB")
    img.thumbnail((800, 800), Image.LANCZOS)
    canvas = Image.new("RGB", (800, 800), "white")
    canvas.paste(img, ((800 - img.width) // 2, (800 - img.height) // 2))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out_path, "JPEG", quality=94, optimize=True)


def make_five_grid(source: Path, out_path: Path) -> None:
    source_img = Image.open(source).convert("RGB")
    source_img.thumbnail((310, 310), Image.LANCZOS)
    bg = Image.new("RGB", (800, 800), (236, 232, 222))
    bg = bg.filter(ImageFilter.GaussianBlur(0.2))
    draw = ImageDraw.Draw(bg, "RGBA")
    for x in range(-80, 900, 150):
        draw.line((x, 0, x + 260, 800), fill=(255, 255, 255, 28), width=3)
    positions = [(190, 180), (610, 180), (400, 400), (190, 620), (610, 620)]
    for cx, cy in positions:
        shadow = Image.new("RGBA", source_img.size, (0, 0, 0, 0))
        shadow.putalpha(Image.new("L", source_img.size, 42).filter(ImageFilter.GaussianBlur(8)))
        layer = bg.convert("RGBA")
        x, y = cx - source_img.width // 2, cy - source_img.height // 2
        layer.alpha_composite(shadow, (x + 7, y + 9))
        bg = layer.convert("RGB")
        bg.paste(source_img, (x, y))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    bg.save(out_path, "JPEG", quality=92, optimize=True)


def workbook_l042_rows() -> list[dict]:
    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb.active
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
    rows = []
    for row in range(2, ws.max_row + 1):
        d_value = ws.cell(row, headers["产品货号"]).value
        if not d_value or not str(d_value).startswith("L042"):
            continue
        rows.append(
            {
                "row": row,
                "d": str(d_value),
                "g": ws.cell(row, headers.get("变种属性值一")).value if headers.get("变种属性值一") else "",
                "sku": ws.cell(row, headers.get("SKU货号")).value if headers.get("SKU货号") else "",
            }
        )
    return rows


def build_review(rows: list[dict], records: dict[str, dict]) -> None:
    SERVED_ROOT.mkdir(parents=True, exist_ok=True)
    for path in OUT_ROOT.iterdir():
        if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png", ".json"}:
            shutil.copy2(path, SERVED_ROOT / path.name)

    cards = []
    for row in rows:
        variant = classify_variant(row["g"], row["sku"])
        rec = records[variant]
        cards.append(
            f"<article><h3>row {row['row']} / {html.escape(row['d'])} / {variant}</h3>"
            f"<p>G: {html.escape(str(row['g']))} | SKU: {html.escape(str(row['sku']))}</p>"
            f"<p>source: {html.escape(rec['source'])}</p>"
            f"<div class='pair'><section><b>整张 SKU 尺寸图</b><img src='./{Path(rec['single']).name}'></section>"
            f"<section><b>五宫格 SKU 尺寸图</b><img src='./{Path(rec['grid']).name}'></section></div>"
            f"</article>"
        )

    page = f"""<!doctype html>
<meta charset="utf-8">
<title>L042 J sizechart redo</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:0;background:#f4f0e7;color:#1f2328}}
header{{position:sticky;top:0;background:#fff;padding:14px 18px;border-bottom:1px solid #ddd;z-index:2}}
main{{display:grid;grid-template-columns:1fr;gap:16px;padding:16px;max-width:1180px;margin:0 auto}}
article{{background:#fff;border:1px solid #ddd;border-radius:8px;padding:12px}}
h3{{font-size:15px;margin:0 0 6px}} p{{font-size:12px;margin:4px 0;word-break:break-all}}
.pair{{display:grid;grid-template-columns:1fr 1fr;gap:12px}} section{{font-size:13px;font-weight:600}}
img{{display:block;width:100%;height:auto;border:1px solid #ddd;background:#eee;margin-top:6px}}
@media(max-width:760px){{.pair{{grid-template-columns:1fr}}}}
</style>
<header><strong>L042 J sizechart redo</strong> · {len(rows)} rows · 使用黑/绿一级文件整张 SKU 尺寸图，保留文字、钉子、尺寸信息</header>
<main>{''.join(cards)}</main>"""
    (SERVED_ROOT / "index.html").write_text(page, encoding="utf-8")


def main() -> None:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    records = {}
    for variant, color_cn in [("black", "黑色"), ("green", "绿色")]:
        source = choose_sizechart(color_cn)
        single = OUT_ROOT / f"L042_{variant}_sizechart_single.jpg"
        grid = OUT_ROOT / f"L042_{variant}_sizechart_fivegrid.jpg"
        normalize_square(source, single)
        make_five_grid(source, grid)
        records[variant] = {
            "variant": variant,
            "source": str(source),
            "single": str(single),
            "grid": str(grid),
        }
    rows = workbook_l042_rows()
    report = {"workbook": str(WORKBOOK), "rows": rows, "records": records, "review": str(SERVED_ROOT / "index.html")}
    (OUT_ROOT / "l042_sizechart_j_redo_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    build_review(rows, records)
    print(json.dumps({"rows": len(rows), "review": str(SERVED_ROOT / "index.html")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
