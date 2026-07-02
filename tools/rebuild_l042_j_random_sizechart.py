from __future__ import annotations

import hashlib
import html
import json
import random
import shutil
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFilter
import numpy as np


WORKBOOK = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_final_199_writeback_20260702_fix_feedback_20260702"
    r"\0616-2_199_最终回传_已应用反馈_L042J重做_T清理_20260702.xlsx"
)
SKU_ROOT = Path(r"E:\jit制图\L042\sku")
OUT_ROOT = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_final_199_writeback_20260702_fix_feedback_20260702"
    r"\l042_random_sizechart_j"
)
SERVED_ROOT = Path(
    r"D:\Desktop\jit\DXXmall\outputs\store_newskill_seedream_fallback_9_20260702"
    r"\l042_j_review_random_sizechart"
)


def variant_key(g_value: object, sku_value: object) -> tuple[str, str]:
    text = f"{g_value or ''} {sku_value or ''}".lower()
    if any(token in text for token in ["绿", "green"]):
        return "green", "绿色"
    return "black", "黑色"


def color_scores(path: Path) -> dict[str, float]:
    image = Image.open(path).convert("RGB").resize((400, 400), Image.LANCZOS)
    arr = np.array(image).astype(int)
    r, g, b = arr[:, :, 0], arr[:, :, 1], arr[:, :, 2]
    overall_green = float(((g > r + 18) & (g > b + 12) & (g > 55)).mean())
    title = arr[:48, :130, :]
    tr, tg, tb = title[:, :, 0], title[:, :, 1], title[:, :, 2]
    title_green = float(((tg > tr + 12) & (tg > tb + 8) & (tg > 40)).mean())
    return {"overall_green": overall_green, "title_green": title_green}


def green_score(path: Path) -> float:
    return color_scores(path)["overall_green"]


def first_level_sources(color_cn: str) -> list[Path]:
    folder = SKU_ROOT / color_cn
    files = [p for p in sorted(folder.iterdir()) if p.is_file()]
    if not files:
        raise FileNotFoundError(f"No first-level files in {folder}")
    filtered: list[Path] = []
    rejected: list[dict] = []
    for path in files:
        scores = color_scores(path)
        score = scores["overall_green"]
        title_score = scores["title_green"]
        if color_cn == "绿色":
            keep = score >= 0.05 and title_score >= 0.03
        else:
            keep = score < 0.05 and title_score < 0.03
        if keep:
            filtered.append(path)
        else:
            rejected.append(
                {
                    "path": str(path),
                    "overall_green": round(score, 4),
                    "title_green": round(title_score, 4),
                }
            )
    if not filtered:
        raise RuntimeError(f"No visually valid {color_cn} first-level files in {folder}")
    rejected_path = OUT_ROOT / f"L042_rejected_{color_cn}_visual_mismatch.json"
    rejected_path.parent.mkdir(parents=True, exist_ok=True)
    rejected_path.write_text(json.dumps(rejected, ensure_ascii=False, indent=2), encoding="utf-8")
    return filtered


def deterministic_pick(files: list[Path], row: int, d_value: str, g_value: object, sku_value: object) -> Path:
    key = f"{row}|{d_value}|{g_value or ''}|{sku_value or ''}".encode("utf-8")
    seed = int(hashlib.sha256(key).hexdigest()[:12], 16)
    rng = random.Random(seed)
    return rng.choice(files)


def textured_background(seed: int) -> Image.Image:
    rng = random.Random(seed)
    palettes = [
        ((237, 232, 220), (214, 223, 207)),
        ((235, 225, 211), (224, 233, 236)),
        ((232, 228, 218), (235, 215, 204)),
        ((224, 232, 221), (238, 228, 205)),
    ]
    top, bottom = palettes[seed % len(palettes)]
    bg = Image.new("RGB", (800, 800), top)
    px = bg.load()
    for y in range(800):
        t = y / 799
        base = tuple(int(top[i] * (1 - t) + bottom[i] * t) for i in range(3))
        for x in range(800):
            n = rng.randint(-4, 4)
            px[x, y] = tuple(max(0, min(255, c + n)) for c in base)
    bg = bg.filter(ImageFilter.GaussianBlur(0.35))
    draw = ImageDraw.Draw(bg, "RGBA")
    for x in range(-80, 900, 150):
        draw.line((x, 0, x + 260, 800), fill=(255, 255, 255, 24), width=3)
    return bg


def make_five_grid(source: Path, out_path: Path, seed: int) -> None:
    rng = random.Random(seed)
    source_img = Image.open(source).convert("RGB")
    cell = 800 / 3
    source_img.thumbnail((252, 252), Image.LANCZOS)
    bg = textured_background(seed)
    # True 五宫格: divide the canvas into a 3x3 grid and fill five cells:
    # top-left, top-right, center, bottom-left, bottom-right.
    placements = [
        (cell * 0.5, cell * 0.5),
        (cell * 2.5, cell * 0.5),
        (cell * 1.5, cell * 1.5),
        (cell * 0.5, cell * 2.5),
        (cell * 2.5, cell * 2.5),
    ]
    for i, (cx, cy) in enumerate(placements):
        item = source_img.copy()
        scale = rng.uniform(0.98, 1.02)
        item = item.resize((int(item.width * scale), int(item.height * scale)), Image.LANCZOS)
        x = round(cx - item.width / 2)
        y = round(cy - item.height / 2)
        layer = bg.convert("RGBA")
        shadow = Image.new("RGBA", item.size, (0, 0, 0, 0))
        shadow.putalpha(Image.new("L", item.size, 42).filter(ImageFilter.GaussianBlur(8)))
        layer.alpha_composite(shadow, (x + 7, y + 9))
        bg = layer.convert("RGB")
        bg.paste(item, (x, y))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    bg.save(out_path, "JPEG", quality=92, optimize=True)


def l042_rows() -> tuple[list[dict], dict[str, int]]:
    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb.active
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
    rows: list[dict] = []
    for row in range(2, ws.max_row + 1):
        d_value = ws.cell(row, headers["产品货号"]).value
        if not d_value or not str(d_value).startswith("L042"):
            continue
        g_value = ws.cell(row, headers.get("变种属性值一")).value if headers.get("变种属性值一") else ""
        sku_value = ws.cell(row, headers.get("SKU货号")).value if headers.get("SKU货号") else ""
        variant, color_cn = variant_key(g_value, sku_value)
        rows.append(
            {
                "row": row,
                "d": str(d_value),
                "g": g_value,
                "sku": sku_value,
                "variant": variant,
                "color_cn": color_cn,
            }
        )
    return rows, headers


def build_review(records: list[dict]) -> None:
    SERVED_ROOT.mkdir(parents=True, exist_ok=True)
    for path in OUT_ROOT.iterdir():
        if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png", ".json"}:
            shutil.copy2(path, SERVED_ROOT / path.name)

    cards = []
    for record in records:
        cards.append(
            f"<article><h3>row {record['row']} / {html.escape(record['d'])} / {html.escape(record['variant'])}</h3>"
            f"<p>G: {html.escape(str(record['g']))} | SKU: {html.escape(str(record['sku']))}</p>"
            f"<p>source: {html.escape(record['source'])}</p>"
            f"<img src='./{html.escape(Path(record['j_path']).name)}' loading='lazy'></article>"
        )
    page = f"""<!doctype html>
<meta charset="utf-8">
<title>L042 random sizechart five-grid</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:0;background:#f4f0e7;color:#222}}
header{{position:sticky;top:0;background:#fff;padding:14px 18px;border-bottom:1px solid #ddd;z-index:2}}
main{{display:grid;grid-template-columns:repeat(auto-fill,minmax(330px,1fr));gap:14px;padding:16px}}
article{{background:#fff;border:1px solid #ddd;border-radius:8px;padding:12px}}
h3{{font-size:15px;margin:0 0 6px}}p{{font-size:12px;margin:4px 0;word-break:break-all}}
img{{display:block;width:100%;height:auto;border:1px solid #ddd;background:#eee;margin-top:8px}}
</style>
<header><strong>L042 random sizechart five-grid</strong> · {len(records)} rows · 按属性匹配黑/绿，随机使用一级文件夹图片，不进二级目录</header>
<main>{''.join(cards)}</main>"""
    (SERVED_ROOT / "index.html").write_text(page, encoding="utf-8")


def main() -> None:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    rows, _headers = l042_rows()
    pools = {"black": first_level_sources("黑色"), "green": first_level_sources("绿色")}
    records = []
    for row in rows:
        pool = pools[row["variant"]]
        source = deterministic_pick(pool, row["row"], row["d"], row["g"], row["sku"])
        seed = int(hashlib.sha256(f"{row['row']}|{source}".encode("utf-8")).hexdigest()[:12], 16)
        j_path = OUT_ROOT / f"L042_row{row['row']}_{row['variant']}_random_fivegrid.jpg"
        make_five_grid(source, j_path, seed)
        records.append(
            {
                **row,
                "source": str(source),
                "source_color_scores": {
                    key: round(value, 4) for key, value in color_scores(source).items()
                },
                "j_path": str(j_path),
            }
        )
    report = {
        "workbook": str(WORKBOOK),
        "source_rule": "first-level E:\\jit制图\\L042\\sku\\黑色 and 绿色 only; no nested folders",
        "records": records,
        "review": "http://127.0.0.1:8765/outputs/store_newskill_seedream_fallback_9_20260702/l042_j_review_random_sizechart/index.html",
    }
    (OUT_ROOT / "l042_random_sizechart_j_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    build_review(records)
    print(json.dumps({"rows": len(records), "review": report["review"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
