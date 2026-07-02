import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def split_urls(value):
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    urls = []
    for part in re.split(r"[\n\r]+", text):
        part = part.strip()
        if not part:
            continue
        found = re.findall(r"https?://[^\s;，,]+", part)
        urls.extend(found or [part])
    return urls


def main():
    workbook = Path(sys.argv[1])
    report_path = Path(sys.argv[2]) if len(sys.argv) > 2 else workbook.with_suffix(".audit.json")

    wb = load_workbook(workbook, read_only=True, data_only=False)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    required = ["产品标题", "产品货号", "变种属性值一", "预览图", "SKU货号", "轮播图", "产品素材图"]
    cols = {name: (headers.index(name) + 1 if name in headers else None) for name in required}

    rows = []
    for row_offset, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d_col = cols["产品货号"]
        d_value = values[d_col - 1] if d_col else None
        if d_value is None or str(d_value).strip() == "":
            continue
        rec = {"row": row_offset}
        for name, col_idx in cols.items():
            rec[name] = values[col_idx - 1] if col_idx else None
        rec["D"] = str(rec["产品货号"]).strip()
        rows.append(rec)

    by_d = defaultdict(list)
    for rec in rows:
        by_d[rec["D"]].append(rec)

    forbidden_title = [
        "自动",
        "智能",
        "AI",
        "APP",
        "蓝牙",
        "wifi",
        "遥控",
        "感应",
        "电动",
        "充电",
        "电池",
        "USB",
        "儿童",
        "玩具",
        "环保",
        "有机",
        "可降解",
        "食品级",
        "最佳",
        "最好",
        "顶级",
        "久用不塌陷",
        "长久不变形",
        "亲肤",
        "高弹填充",
        "四面延伸",
    ]
    locked_keywords = [
        "不要",
        "死刑",
        "delete",
        "deleted",
        "reject",
        "rejected",
        "wrong",
        "l058-extra-fixed-under145k",
        "L058_extra_fixed_800_under145k",
    ]
    process_keywords = ["过程", "不入库", "review", "candidate", "temp", "tmp", "fallback_9_review"]
    warning_keywords = ["redo"]
    size_re = re.compile(r"(尺寸|尺码|size|Size|SIZE|measure|dimension|chart)", re.I)
    code_re = re.compile(r"[A-Z][0-9][A-Z]")

    issues = []
    warnings = []
    examples = defaultdict(list)
    warning_examples = defaultdict(list)

    def add(kind, rec_or_d, detail):
        issues.append(kind)
        if len(examples[kind]) >= 20:
            return
        if isinstance(rec_or_d, dict):
            examples[kind].append({"row": rec_or_d.get("row"), "D": rec_or_d.get("D"), "detail": detail})
        else:
            examples[kind].append({"D": rec_or_d, "detail": detail})

    def warn(kind, rec_or_d, detail):
        warnings.append(kind)
        if len(warning_examples[kind]) >= 20:
            return
        if isinstance(rec_or_d, dict):
            warning_examples[kind].append({"row": rec_or_d.get("row"), "D": rec_or_d.get("D"), "detail": detail})
        else:
            warning_examples[kind].append({"D": rec_or_d, "detail": detail})

    title_codes = {}
    for d_value, group in by_d.items():
        titles = [str(item.get("产品标题") or "").strip() for item in group]
        t_values = [str(item.get("轮播图") or "").strip() for item in group]
        u_values = [str(item.get("产品素材图") or "").strip() for item in group]
        if len(set(titles)) > 1:
            add("same_d_title_mismatch", d_value, list(dict.fromkeys(titles))[:3])
        if len(set(t_values)) > 1:
            add("same_d_t_mismatch", d_value, [item["row"] for item in group[:5]])
        if len(set(u_values)) > 1:
            add("same_d_u_mismatch", d_value, [item["row"] for item in group[:5]])

        title = titles[0] if titles else ""
        codes = code_re.findall(title)
        title_codes[d_value] = codes[-1] if codes else None
        if not codes:
            add("title_missing_fingerprint", group[0], title[:120])
        for word in forbidden_title:
            if word in title:
                add("title_forbidden_word", group[0], {"word": word, "title": title[:120]})
                break

    code_to_d = defaultdict(list)
    for d_value, code in title_codes.items():
        if code:
            code_to_d[code].append(d_value)
    for code, d_values in code_to_d.items():
        if len(d_values) > 1:
            add("title_fingerprint_duplicate_across_d", d_values[0], {"code": code, "ds": d_values[:10]})

    for rec in rows:
        j_value = str(rec.get("预览图") or "").strip()
        t_value = str(rec.get("轮播图") or "").strip()
        u_value = str(rec.get("产品素材图") or "").strip()
        if not j_value:
            add("empty_j", rec, "")
        if not t_value:
            add("empty_t", rec, "")
            continue

        urls = split_urls(t_value)
        if len(urls) > 10:
            add("t_over_10", rec, len(urls))
        if len(set(urls)) != len(urls):
            add("t_duplicate_url", rec, {"count": len(urls), "unique": len(set(urls))})
        if urls:
            if u_value != urls[0]:
                add("u_not_equal_t1", rec, {"u": u_value[:160], "t1": urls[0][:160]})
            if len(urls) < 4:
                add("t_less_than_4_for_size_slot", rec, len(urls))
            elif not size_re.search(urls[3]):
                add("t4_not_size_by_url_clue", rec, urls[3][:220])

        all_text = "\n".join([j_value, t_value, u_value])
        for keyword in locked_keywords:
            if keyword in all_text:
                add("locked_or_deleted_keyword_returned", rec, keyword)
                break
        for keyword in process_keywords:
            if keyword.lower() in all_text.lower():
                add("process_temp_keyword_in_urls", rec, keyword)
                break
        for keyword in warning_keywords:
            if keyword.lower() in all_text.lower():
                warn("approved_redo_keyword_in_urls", rec, keyword)
                break

    issue_counts = dict(Counter(issues))
    warning_counts = dict(Counter(warnings))
    summary = {
        "workbook": str(workbook),
        "sheet": ws.title,
        "max_row": ws.max_row,
        "effective_rows": len(rows),
        "unique_d": len(by_d),
        "contains_L091060503": "L091060503" in by_d,
        "contains_L091060507": "L091060507" in by_d,
        "required_columns": cols,
        "issue_counts": issue_counts,
        "warning_counts": warning_counts,
        "examples": dict(examples),
        "warning_examples": dict(warning_examples),
    }
    summary["blocking_issue_total"] = sum(issue_counts.values())
    summary["pass_hard_checks"] = (
        summary["blocking_issue_total"] == 0
        and len(rows) == 323
        and len(by_d) == 197
        and not summary["contains_L091060503"]
        and not summary["contains_L091060507"]
    )

    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"REPORT={report_path}")


if __name__ == "__main__":
    main()
